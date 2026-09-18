#!/usr/bin/env python3
"""Excel 两阶段脱敏工具：先扫描生成可编辑方案，确认后再执行。"""

from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


PLAN_SHEET = "脱敏方案"
SUMMARY_SHEET = "扫描摘要"
INSTRUCTION_SHEET = "使用说明"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
PLAN_HEADERS = [
    "启用", "复核状态", "源文件", "工作表", "表头行", "列序号", "字段名称",
    "识别类型", "风险等级", "识别依据", "脱敏前样式（实际样例）", "处理方式",
    "脱敏后样式（可修改）", "预计处理行数", "备注",
]

METHODS = ["稳定替换", "固定替换", "保留末四位", "删除", "保留原值", "人工确认"]
DEFAULT_RULES_PATH = Path(__file__).with_name("default_rules.json")


@dataclass
class Candidate:
    source_file: str
    sheet: str
    header_row: int
    column: int
    field_name: str
    risk_type: str
    risk_level: str
    reason: str
    samples: list[str]
    method: str
    template: str
    estimated_rows: int
    enabled: str


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def discover_files(source: Path) -> tuple[Path, list[Path]]:
    source = source.expanduser().resolve()
    if source.is_file():
        if source.suffix.lower() not in SUPPORTED_SUFFIXES:
            raise ValueError("仅支持 .xlsx 和 .xlsm 文件")
        return source.parent, [source]
    if not source.is_dir():
        raise FileNotFoundError(f"找不到输入路径：{source}")
    files = sorted(
        p for p in source.rglob("*")
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$")
    )
    if not files:
        raise ValueError("输入目录中没有可处理的 .xlsx 或 .xlsm 文件")
    return source, files


def load_rules(path: Path) -> dict[str, Any]:
    path = path.expanduser().resolve()
    try:
        config = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"规则文件不是有效 JSON：{path}（{exc}）") from exc
    if not isinstance(config.get("rules"), list) or not config["rules"]:
        raise ValueError("规则文件必须包含非空的 rules 数组")
    compiled = dict(config)
    compiled.update({"source": str(path), "header_exclusions": [], "rules": []})
    for pattern in config.get("header_exclusions", []):
        compiled["header_exclusions"].append(re.compile(pattern, re.I))
    for index, item in enumerate(config["rules"], start=1):
        required = ["label", "risk_level", "method", "template"]
        missing = [name for name in required if not normalize_text(item.get(name))]
        if missing:
            raise ValueError(f"规则第 {index} 项缺少：{', '.join(missing)}")
        rule = dict(item)
        rule["header_patterns_compiled"] = [re.compile(pattern, re.I) for pattern in item.get("header_patterns", [])]
        rule["content_patterns_compiled"] = [re.compile(pattern, re.I) for pattern in item.get("content_patterns", [])]
        variants = []
        for variant in item.get("template_variants", []):
            variants.append({
                "patterns": [re.compile(pattern, re.I) for pattern in variant.get("header_patterns", [])],
                "template": variant.get("template", item["template"]),
            })
        rule["template_variants_compiled"] = variants
        compiled["rules"].append(rule)
    return compiled


def choose_header_row(ws, rules: dict[str, Any], scan_rows: int = 30) -> int:
    best_row, best_score = 1, -1.0
    upper = min(ws.max_row, scan_rows)
    for row_no in range(1, upper + 1):
        values = [normalize_text(ws.cell(row_no, col).value) for col in range(1, ws.max_column + 1)]
        nonempty = [value for value in values if value]
        if not nonempty:
            continue
        keyword_hits = sum(
            any(pattern.search(value) for rule in rules["rules"] for pattern in rule["header_patterns_compiled"])
            for value in nonempty
        )
        text_count = sum(not value.replace(".", "", 1).isdigit() for value in nonempty)
        score = keyword_hits * 12 + len(nonempty) + text_count * 0.25
        if score > best_score:
            best_row, best_score = row_no, score
    return best_row


def combined_header(ws, header_row: int, column: int) -> str:
    parts = []
    for row_no in range(max(1, header_row - 2), header_row + 1):
        value = normalize_text(ws.cell(row_no, column).value)
        if value and value not in parts:
            parts.append(value)
    return " / ".join(parts) or f"第{column}列"


def unique_samples(ws, column: int, start_row: int, limit: int, sample_rows: int) -> tuple[list[str], int]:
    samples: list[str] = []
    seen = set()
    populated = 0
    end_row = min(ws.max_row, start_row + sample_rows - 1)
    for row_no in range(start_row, end_row + 1):
        cell = ws.cell(row_no, column)
        if cell.data_type == "f":
            continue
        value = normalize_text(cell.value)
        if not value:
            continue
        populated += 1
        if value not in seen and len(samples) < limit:
            samples.append(value)
            seen.add(value)
    return samples, populated


def classify_column(header: str, samples: list[str], rules: dict[str, Any]) -> tuple[str, str, str, str, str, str] | None:
    excluded = any(pattern.search(header) for pattern in rules["header_exclusions"])
    if not excluded:
        for rule in rules["rules"]:
            matched = next((pattern for pattern in rule["header_patterns_compiled"] if pattern.search(header)), None)
            if not matched:
                continue
            template = rule["template"]
            for variant in rule["template_variants_compiled"]:
                if any(pattern.search(header) for pattern in variant["patterns"]):
                    template = variant["template"]
                    break
            enabled = normalize_text(rule.get("default_enabled", "Y")).upper()
            return rule["label"], rule["risk_level"], f"字段名命中外部规则：{matched.pattern}", rule["method"], template, enabled

    checked = [sample for sample in samples if len(sample) <= int(rules.get("max_content_length", 100))]
    for rule in rules["rules"]:
        for pattern in rule["content_patterns_compiled"]:
            hits = sum(bool(pattern.search(sample)) for sample in checked)
            required_hits = min(len(checked), int(rule.get("content_min_hits", 1)))
            if checked and hits >= max(1, required_hits):
                enabled = normalize_text(rule.get("default_enabled", "Y")).upper()
                return rule["label"], rule["risk_level"], f"样例内容命中外部规则 {hits} 项", rule["method"], rule["template"], enabled
    return None


def scan_workbook(path: Path, root: Path, sample_rows: int, rules: dict[str, Any]) -> tuple[list[Candidate], list[dict[str, Any]]]:
    keep_vba = path.suffix.lower() == ".xlsm"
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
    candidates: list[Candidate] = []
    summaries: list[dict[str, Any]] = []
    relative = str(path.relative_to(root)) if path.is_relative_to(root) else path.name
    for ws in wb.worksheets:
        header_row = choose_header_row(ws, rules)
        sheet_candidates = 0
        for column in range(1, ws.max_column + 1):
            header = combined_header(ws, header_row, column)
            samples, populated = unique_samples(ws, column, header_row + 1, 3, sample_rows)
            if not samples:
                continue
            classification = classify_column(header, samples, rules)
            if not classification:
                continue
            risk_type, risk_level, reason, method, template, enabled = classification
            candidates.append(Candidate(
                source_file=relative,
                sheet=ws.title,
                header_row=header_row,
                column=column,
                field_name=header,
                risk_type=risk_type,
                risk_level=risk_level,
                reason=reason,
                samples=samples,
                method=method,
                template=template,
                estimated_rows=populated,
                enabled=enabled,
            ))
            sheet_candidates += 1
        summaries.append({
            "源文件": relative,
            "工作表": ws.title,
            "识别表头行": header_row,
            "数据行数（约）": max(0, ws.max_row - header_row),
            "字段列数": ws.max_column,
            "候选脱敏字段数": sheet_candidates,
            "说明": "候选项需人工复核后才能执行",
        })
    wb.close()
    return candidates, summaries


def style_table(ws, max_row: int, max_col: int, freeze: str = "A2") -> None:
    purple = "4F246A"
    light = "F2ECF5"
    orange = "E2A33A"
    thin = Side(style="thin", color="DDD4E1")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=purple)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=light)
    if max_row >= 2:
        ws.conditional_formatting.add(
            f"A2:A{max_row}",
            FormulaRule(formula=["A2=\"Y\""], fill=PatternFill("solid", fgColor="E2F0D9")),
        )
        ws.conditional_formatting.add(
            f"B2:B{max_row}",
            FormulaRule(formula=["B2=\"待确认\""], fill=PatternFill("solid", fgColor="FFF2CC")),
        )
    ws.row_dimensions[1].height = 34
    widths = [10, 12, 28, 20, 10, 10, 24, 16, 10, 32, 46, 15, 34, 15, 30]
    for index in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(index)].width = widths[index - 1] if index <= len(widths) else 18
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.tabColor = orange


def create_plan(
    output: Path,
    candidates: list[Candidate],
    summaries: list[dict[str, Any]],
    source: Path,
    rules_source: str,
) -> None:
    wb = Workbook()
    guide = wb.active
    guide.title = INSTRUCTION_SHEET
    guide.sheet_view.showGridLines = False
    guide["A1"] = "Excel 数据脱敏方案｜使用说明"
    guide["A1"].font = Font(size=18, bold=True, color="4F246A")
    guide.merge_cells("A1:F1")
    guide["A3"] = "操作顺序"
    guide["A3"].font = Font(bold=True, color="FFFFFF")
    guide["A3"].fill = PatternFill("solid", fgColor="4F246A")
    instructions = [
        "1. 本文件由 scan 命令生成，扫描过程不会修改源 Excel。",
        "2. 在“脱敏方案”中检查真实样例、识别类型、处理方式和脱敏后样式。",
        "3. 不需要处理的字段把“启用”改为 N；需要处理的字段保持 Y。",
        "4. 可修改“处理方式”和“脱敏后样式”。稳定替换可使用 {序号:03d}，掩码可使用 {末4位}，保留原值可使用 {原值}。",
        "5. 将所有启用项的“复核状态”改为“已确认”。未确认或仍为“人工确认”的启用项会阻止执行。",
        "6. 执行 apply 命令后，程序在新目录生成脱敏文件，不覆盖源文件。",
        "7. 脱敏方案包含真实样例，属于受控资料；不要与普通学员包一起分发。",
    ]
    for idx, text in enumerate(instructions, start=4):
        guide.cell(idx, 1, text)
        guide.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=6)
        guide.cell(idx, 1).alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[idx].height = 30
    guide["A13"] = "输入位置"
    guide["B13"] = str(source)
    guide["A14"] = "生成时间"
    guide["B14"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    guide["A15"] = "识别规则"
    guide["B15"] = rules_source
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 90

    summary_ws = wb.create_sheet(SUMMARY_SHEET)
    summary_headers = ["源文件", "工作表", "识别表头行", "数据行数（约）", "字段列数", "候选脱敏字段数", "说明"]
    summary_ws.append(summary_headers)
    for item in summaries:
        summary_ws.append([item[header] for header in summary_headers])
    style_table(summary_ws, max(2, summary_ws.max_row), len(summary_headers))
    for col, width in enumerate([32, 22, 14, 16, 12, 18, 30], start=1):
        summary_ws.column_dimensions[get_column_letter(col)].width = width

    plan_ws = wb.create_sheet(PLAN_SHEET)
    plan_ws.append(PLAN_HEADERS)
    for item in candidates:
        plan_ws.append([
            item.enabled, "待确认", item.source_file, item.sheet, item.header_row, item.column,
            item.field_name, item.risk_type, item.risk_level, item.reason,
            "\n".join(item.samples), item.method, item.template, item.estimated_rows, "",
        ])
    style_table(plan_ws, max(2, plan_ws.max_row), len(PLAN_HEADERS))
    if plan_ws.max_row >= 2:
        yes_no = DataValidation(type="list", formula1='"Y,N"', allow_blank=False)
        review = DataValidation(type="list", formula1='"待确认,已确认"', allow_blank=False)
        methods = DataValidation(type="list", formula1='"' + ",".join(METHODS) + '"', allow_blank=False)
        plan_ws.add_data_validation(yes_no)
        plan_ws.add_data_validation(review)
        plan_ws.add_data_validation(methods)
        yes_no.add(f"A2:A{plan_ws.max_row}")
        review.add(f"B2:B{plan_ws.max_row}")
        methods.add(f"L2:L{plan_ws.max_row}")
        for row_no in range(2, plan_ws.max_row + 1):
            plan_ws.row_dimensions[row_no].height = 52
            plan_ws.cell(row_no, 1).font = Font(color="0000FF")
            plan_ws.cell(row_no, 2).font = Font(color="0000FF")
            plan_ws.cell(row_no, 12).font = Font(color="0000FF")
            plan_ws.cell(row_no, 13).font = Font(color="0000FF")

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def load_plan(plan_path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(plan_path, data_only=False)
    if PLAN_SHEET not in wb.sheetnames:
        raise ValueError(f"方案文件缺少“{PLAN_SHEET}”工作表")
    ws = wb[PLAN_SHEET]
    headers = [normalize_text(cell.value) for cell in ws[1]]
    missing = [header for header in PLAN_HEADERS if header not in headers]
    if missing:
        raise ValueError(f"方案表缺少列：{', '.join(missing)}")
    positions = {header: headers.index(header) + 1 for header in headers}
    rows = []
    for row_no in range(2, ws.max_row + 1):
        record = {header: ws.cell(row_no, positions[header]).value for header in PLAN_HEADERS}
        if not normalize_text(record["源文件"]):
            continue
        record["方案行号"] = row_no
        rows.append(record)
    wb.close()
    return rows


def format_template(template: str, original: str, sequence: int) -> str:
    return template.format(序号=sequence, 原值=original, 末4位=original[-4:] if original else "")


def transform_value(method: str, template: str, original: Any, sequence: int) -> Any:
    text = normalize_text(original)
    if not text:
        return original
    if method == "删除":
        return None
    if method == "保留原值":
        return original
    if method == "固定替换":
        return format_template(template, text, sequence)
    if method in {"稳定替换", "保留末四位"}:
        return format_template(template, text, sequence)
    raise ValueError(f"不支持直接执行的处理方式：{method}")


def validate_plan(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    enabled = [record for record in records if normalize_text(record["启用"]).upper() == "Y"]
    if not enabled:
        raise ValueError("方案中没有启用的脱敏字段")
    errors = []
    for record in enabled:
        row_no = record["方案行号"]
        if normalize_text(record["复核状态"]) != "已确认":
            errors.append(f"第 {row_no} 行尚未标记为“已确认”")
        method = normalize_text(record["处理方式"])
        if method == "人工确认":
            errors.append(f"第 {row_no} 行仍为“人工确认”，请修改处理方式或停用")
        if method not in METHODS:
            errors.append(f"第 {row_no} 行处理方式无效：{method}")
        if method in {"稳定替换", "固定替换", "保留末四位"} and not normalize_text(record["脱敏后样式（可修改）"]):
            errors.append(f"第 {row_no} 行缺少脱敏后样式")
    if errors:
        raise ValueError("方案未完成确认：\n- " + "\n- ".join(errors))
    return enabled


def apply_plan(source: Path, plan_path: Path, output_dir: Path, overwrite_output: bool = False) -> dict[str, Any]:
    root, files = discover_files(source)
    records = validate_plan(load_plan(plan_path))
    by_location: dict[tuple[str, str, int], list[dict[str, Any]]] = defaultdict(list)
    for record in records:
        key = (normalize_text(record["源文件"]), normalize_text(record["工作表"]), int(record["列序号"]))
        by_location[key].append(record)

    global_maps: dict[tuple[str, str, str], dict[str, str]] = defaultdict(dict)
    audit_rows = []
    generated_files = []
    output_dir = output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    for path in files:
        relative = str(path.relative_to(root)) if path.is_relative_to(root) else path.name
        targets = [key for key in by_location if key[0] == relative]
        if not targets:
            continue
        keep_vba = path.suffix.lower() == ".xlsm"
        wb = load_workbook(path, read_only=False, data_only=False, keep_vba=keep_vba)
        changed_in_file = 0
        for _, sheet_name, column in targets:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"源文件 {relative} 中找不到工作表：{sheet_name}")
            ws = wb[sheet_name]
            for record in by_location[(relative, sheet_name, column)]:
                header_row = int(record["表头行"])
                method = normalize_text(record["处理方式"])
                template = normalize_text(record["脱敏后样式（可修改）"])
                risk_type = normalize_text(record["识别类型"])
                # 同名字段跨文件保持一致；编码列和名称列分别编号，避免两类原值混入同一序列。
                mapping = global_maps[(risk_type, normalize_text(record["字段名称"]), template)]
                changed = 0
                skipped_formulas = 0
                before_examples, after_examples = [], []
                for row_no in range(header_row + 1, ws.max_row + 1):
                    cell = ws.cell(row_no, column)
                    if cell.data_type == "f":
                        skipped_formulas += 1
                        continue
                    original = cell.value
                    text = normalize_text(original)
                    if not text:
                        continue
                    if method == "稳定替换":
                        if text not in mapping:
                            mapping[text] = transform_value(method, template, original, len(mapping) + 1)
                        transformed = mapping[text]
                    else:
                        transformed = transform_value(method, template, original, len(mapping) + 1)
                    if transformed != original:
                        cell.value = transformed
                        changed += 1
                        changed_in_file += 1
                        if len(before_examples) < 3:
                            before_examples.append(text)
                            after_examples.append(normalize_text(transformed))
                audit_rows.append({
                    "源文件": relative,
                    "工作表": sheet_name,
                    "字段名称": normalize_text(record["字段名称"]),
                    "识别类型": risk_type,
                    "处理方式": method,
                    "实际处理单元格数": changed,
                    "跳过公式单元格数": skipped_formulas,
                    "脱敏前样例": before_examples,
                    "脱敏后样例": after_examples,
                })

        destination = output_dir / Path(relative).parent / f"{path.stem}_脱敏{path.suffix}"
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists() and not overwrite_output:
            raise FileExistsError(f"输出文件已存在：{destination}。如需替换请增加 --overwrite-output")
        wb.save(destination)
        wb.close()
        generated_files.append({"源文件": relative, "输出文件": str(destination), "修改单元格数": changed_in_file})

    if not generated_files:
        raise ValueError("方案中的源文件与当前输入路径不匹配，没有生成任何文件")

    confirmed_plan = output_dir / "脱敏方案_已执行.xlsx"
    shutil.copy2(plan_path, confirmed_plan)
    log = {
        "执行时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "输入路径": str(source.expanduser().resolve()),
        "方案文件": str(plan_path.expanduser().resolve()),
        "说明": "日志包含脱敏前样例，按受控资料管理",
        "生成文件": generated_files,
        "字段处理记录": audit_rows,
    }
    log_path = output_dir / "脱敏执行记录.json"
    log_path.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"confirmed_plan": str(confirmed_plan), "log": str(log_path), **log}


def command_scan(args) -> int:
    source = Path(args.input)
    rules = load_rules(Path(args.rules))
    root, files = discover_files(source)
    all_candidates: list[Candidate] = []
    all_summaries: list[dict[str, Any]] = []
    for path in files:
        candidates, summaries = scan_workbook(path, root, args.sample_rows, rules)
        all_candidates.extend(candidates)
        all_summaries.extend(summaries)
    output = Path(args.output).expanduser().resolve()
    create_plan(output, all_candidates, all_summaries, source.expanduser().resolve(), rules["source"])
    print(json.dumps({
        "状态": "扫描完成，未修改源文件",
        "扫描文件数": len(files),
        "候选字段数": len(all_candidates),
        "识别规则": rules["source"],
        "方案文件": str(output),
        "下一步": "打开方案文件复核，将启用项标记为已确认后再运行 apply",
    }, ensure_ascii=False, indent=2))
    return 0


def command_apply(args) -> int:
    result = apply_plan(
        Path(args.input), Path(args.plan), Path(args.output_dir), overwrite_output=args.overwrite_output
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Excel 两阶段脱敏工具：先扫描生成可编辑方案，确认后再执行。"
    )
    subparsers = parser.add_subparsers(dest="command", required=True)

    scan = subparsers.add_parser("scan", help="扫描 Excel，生成待确认的脱敏方案")
    scan.add_argument("--input", required=True, help="一个 Excel 文件或包含 Excel 的目录")
    scan.add_argument("--output", required=True, help="待生成的脱敏方案 .xlsx")
    scan.add_argument("--sample-rows", type=int, default=2000, help="每列最多扫描的数据行数，默认 2000")
    scan.add_argument("--rules", default=str(DEFAULT_RULES_PATH), help="外部识别规则 JSON；默认使用工具目录中的 default_rules.json")
    scan.set_defaults(func=command_scan)

    apply_cmd = subparsers.add_parser("apply", help="读取已确认方案，生成脱敏后的新 Excel")
    apply_cmd.add_argument("--input", required=True, help="扫描时使用的原 Excel 文件或目录")
    apply_cmd.add_argument("--plan", required=True, help="人工确认后的脱敏方案 .xlsx")
    apply_cmd.add_argument("--output-dir", required=True, help="脱敏结果输出目录")
    apply_cmd.add_argument("--overwrite-output", action="store_true", help="允许替换已存在的脱敏输出文件；仍不会覆盖源文件")
    apply_cmd.set_defaults(func=command_apply)
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    try:
        return args.func(args)
    except Exception as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
