from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
import json
from pathlib import Path
import re
import shutil
from typing import Any

from openpyxl import Workbook
from sqlalchemy import select

from audit_flow_system.core.config import BASE_DIR
from audit_flow_system.core.db import SessionLocal, safe_database_label
from audit_flow_system.core.utils import list_dict, obj_dict
from audit_flow_system.models import Attachment, EnterpriseContact, Project, ProjectMember, Workpaper
from audit_flow_system.services.autofill import autofill_plan, autofill_suggestions
from audit_flow_system.services.autofill_verification import WorkpaperIdentity, verify_plan_items


SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".docx"}
C_SCOPES = {"C22", "C21", "C21-1", "C26"}


def json_default(value: Any) -> str:
    return str(value)


def safe_filename(value: str, default: str = "workpaper") -> str:
    text = re.sub(r"[^\w\u4e00-\u9fff.-]+", "_", value).strip("._")
    return (text or default)[:120]


def existing_workpaper_path(workpaper: dict[str, Any]) -> Path | None:
    text = str(workpaper.get("file_path") or "").strip()
    if not text:
        return None
    path = Path(text).expanduser()
    if path.exists() and path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES:
        return path
    return None


def project_context(project: Project, contacts: list[EnterpriseContact], members: list[ProjectMember]) -> dict[str, Any]:
    context = obj_dict(project)
    context["contacts"] = list_dict(contacts)
    context["members"] = [
        {
            **obj_dict(member),
            "username": member.user.username if member.user else "",
            "display_name": member.user.display_name if member.user else "",
            "email": member.user.email if member.user else "",
        }
        for member in members
    ]
    return context


def classify_scope(scope: str) -> str:
    if scope == "B":
        return "B"
    if scope in C_SCOPES:
        return "C"
    if scope.startswith("A"):
        return "A"
    return "Other"


def suggestion_priority(suggestion: dict[str, Any]) -> int:
    text = " ".join(
        str(suggestion.get(key) or "")
        for key in ("rule_id", "scope", "workpaper", "workpaper_code", "sheet", "name")
    )
    if "B60-2-3" in text:
        return 0
    if "B22A-4-4-1" in text:
        return 1
    if "C22" in text:
        return 0
    if "C21-1" in text:
        return 1
    if "A27" in text:
        return 0
    return 10


def plan_has_writable_item(plan: list[dict[str, Any]]) -> bool:
    return any(
        item.get("status") in {"planned", "unchanged"}
        and (item.get("cell") or item.get("locator"))
        and item.get("workbook_path")
        for item in plan
    )


def pick_sample_candidates(
    bucket: str,
    candidates: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    ordered = sorted(candidates, key=lambda item: (suggestion_priority(item), str(item.get("rule_id") or "")))
    if bucket != "C":
        return ordered
    selected: list[dict[str, Any]] = []
    used_ids: set[int] = set()
    for scope in ("C22", "C21-1", "C26", "C21"):
        for idx, item in enumerate(ordered):
            if idx in used_ids or item.get("scope") != scope:
                continue
            selected.append(item)
            used_ids.add(idx)
            break
    for idx, item in enumerate(ordered):
        if idx in used_ids:
            continue
        selected.append(item)
    return selected


def select_sample_suggestions(
    suggestions: list[dict[str, Any]],
    workpapers: list[dict[str, Any]],
    context: dict[str, Any],
    *,
    per_class: int,
    max_inspect_per_bucket: int | None = None,
) -> tuple[list[dict[str, Any]], dict[str, list[str]], list[dict[str, Any]]]:
    buckets: dict[str, list[dict[str, Any]]] = defaultdict(list)
    gaps: dict[str, list[str]] = defaultdict(list)
    inspected: list[dict[str, Any]] = []
    inspected_by_bucket: Counter[str] = Counter()
    limited_by_bucket: Counter[str] = Counter()
    for suggestion in sorted(suggestions, key=lambda item: (suggestion_priority(item), str(item.get("rule_id") or ""))):
        bucket = classify_scope(str(suggestion.get("scope") or ""))
        if bucket not in {"B", "C", "A"}:
            continue
        if max_inspect_per_bucket is not None and inspected_by_bucket[bucket] >= max_inspect_per_bucket:
            limited_by_bucket[bucket] += 1
            continue
        inspected_by_bucket[bucket] += 1
        buckets[bucket].append(suggestion)

    selected: list[dict[str, Any]] = []
    for bucket in ("B", "C", "A"):
        candidates = pick_sample_candidates(bucket, buckets.get(bucket, []))
        selected_count = 0
        if not candidates:
            gaps[bucket].append("no dry-run writable suggestion selected")
        for suggestion in candidates:
            if selected_count >= per_class:
                break
            plan = autofill_plan([suggestion], workpapers, context, apply=False)
            writable = plan_has_writable_item(plan)
            inspected.append(
                {
                    "rule_id": suggestion.get("rule_id", ""),
                    "scope": suggestion.get("scope", ""),
                    "workpaper": suggestion.get("workpaper", ""),
                    "sheet": suggestion.get("sheet", ""),
                    "bucket": bucket,
                    "plan_count": len(plan),
                    "status_counts": dict(Counter(str(item.get("status") or "") for item in plan)),
                    "writable": writable,
                }
            )
            if writable:
                selected.append(suggestion)
                selected_count += 1
            else:
                statuses = Counter(str(item.get("status") or "") for item in plan)
                gaps[bucket].append(
                    f"{suggestion.get('rule_id') or ''}: no writable target in dry-run ({dict(statuses)})"
                )
        if candidates and selected_count == 0:
            gaps[bucket].append("no dry-run writable suggestion selected")
        if limited_by_bucket.get(bucket, 0):
            gaps[bucket].append(
                f"batch dry-run inspection limited after {inspected_by_bucket[bucket]} {bucket} suggestions; "
                f"{limited_by_bucket[bucket]} suggestions not inspected"
            )
    return selected, dict(gaps), inspected


def choose_project(
    db,
    requested_project_id: int | None,
    *,
    per_class: int,
    max_inspect_per_bucket: int | None = None,
) -> tuple[Project, list[Workpaper], list[Attachment], list[EnterpriseContact], list[ProjectMember], list[dict[str, Any]], dict[str, list[str]], list[dict[str, Any]]]:
    candidates: list[Project]
    if requested_project_id is not None:
        project = db.get(Project, requested_project_id)
        if project is None:
            raise RuntimeError(f"project not found: {requested_project_id}")
        candidates = [project]
    else:
        candidates = db.execute(select(Project).order_by(Project.id.desc())).scalars().all()

    last_reason = "no projects found"
    for project in candidates:
        workpapers = db.execute(
            select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.code, Workpaper.id)
        ).scalars().all()
        attachments = db.execute(
            select(Attachment).where(Attachment.project_id == project.id).order_by(Attachment.index_no, Attachment.id)
        ).scalars().all()
        contacts = db.execute(
            select(EnterpriseContact).where(EnterpriseContact.project_id == project.id).order_by(EnterpriseContact.id)
        ).scalars().all()
        members = db.execute(
            select(ProjectMember).where(ProjectMember.project_id == project.id).order_by(ProjectMember.id)
        ).scalars().all()
        workpaper_dicts = list_dict(workpapers)
        attachment_dicts = list_dict(attachments)
        existing_count = sum(1 for item in workpaper_dicts if existing_workpaper_path(item))
        if not workpaper_dicts or not attachment_dicts or existing_count == 0:
            last_reason = (
                f"project {project.id} skipped: workpapers={len(workpaper_dicts)}, "
                f"attachments={len(attachment_dicts)}, existing_supported_workpapers={existing_count}"
            )
            continue
        context = project_context(project, contacts, members)
        suggestions = autofill_suggestions(attachment_dicts)
        selected, gaps, inspected = select_sample_suggestions(
            suggestions,
            workpaper_dicts,
            context,
            per_class=per_class,
            max_inspect_per_bucket=max_inspect_per_bucket,
        )
        if selected:
            return project, workpapers, attachments, contacts, members, selected, gaps, inspected
        last_reason = f"project {project.id} skipped: no selected autofill suggestions"
    raise RuntimeError(last_reason)


def copy_workpapers(
    workpapers: list[Workpaper],
    test_dir: Path,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, WorkpaperIdentity]]:
    copy_dir = test_dir / "workpapers"
    copy_dir.mkdir(parents=True, exist_ok=True)
    copied_workpapers: list[dict[str, Any]] = []
    mapping: list[dict[str, Any]] = []
    identity_by_path: dict[str, WorkpaperIdentity] = {}
    for workpaper in workpapers:
        data = obj_dict(workpaper)
        original_path = existing_workpaper_path(data)
        copied_path = ""
        copy_status = "not_copied"
        if original_path is not None:
            suffix = original_path.suffix
            filename = safe_filename(f"wp_{workpaper.id}_{workpaper.code}_{original_path.stem}") + suffix
            target_path = copy_dir / filename
            counter = 1
            while target_path.exists():
                target_path = copy_dir / f"{target_path.stem}_{counter}{suffix}"
                counter += 1
            shutil.copy2(original_path, target_path)
            copied_path = str(target_path)
            copy_status = "copied"
            identity_by_path[str(target_path)] = WorkpaperIdentity(
                workpaper_id=workpaper.id,
                workpaper_code=workpaper.code,
                workpaper_name=workpaper.name,
            )
        data["file_path"] = copied_path
        copied_workpapers.append(data)
        mapping.append(
            {
                "workpaper_id": workpaper.id,
                "workpaper_code": workpaper.code,
                "workpaper_name": workpaper.name,
                "original_path": str(original_path or workpaper.file_path or ""),
                "test_copy_path": copied_path,
                "copy_status": copy_status,
            }
        )
    return copied_workpapers, mapping, identity_by_path


def status_counts(records: list[dict[str, Any]], field: str) -> dict[str, int]:
    return dict(Counter(str(record.get(field) or "") for record in records))


def build_gap_records(
    inspected: list[dict[str, Any]],
    selection_gaps: dict[str, list[str]],
    verification_records: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    gaps: list[dict[str, Any]] = []
    for record in verification_records:
        verify_status = str(record.get("verification_status") or "")
        write_status = str(record.get("write_status") or "")
        if verify_status == "failed":
            priority = "P0"
        elif verify_status in {"not_supported", "missing_locator"}:
            priority = "P1"
        elif write_status in {"blocked", "skipped"}:
            priority = "P2"
        else:
            continue
        gaps.append(
            {
                "priority": priority,
                "scope": record.get("scope", ""),
                "workpaper_code": record.get("workpaper_code", ""),
                "rule_id": record.get("rule_id", ""),
                "field": record.get("field", ""),
                "issue": record.get("write_message") or record.get("verification_message") or "",
                "locator": record.get("locator", ""),
            }
        )
    for bucket, messages in selection_gaps.items():
        for message in messages:
            gaps.append(
                {
                    "priority": "P2",
                    "scope": bucket,
                    "workpaper_code": "",
                    "rule_id": "",
                    "field": "",
                    "issue": message,
                    "locator": "",
                }
            )
    for item in inspected:
        if not item.get("writable"):
            gaps.append(
                {
                    "priority": "P2",
                    "scope": item.get("scope", ""),
                    "workpaper_code": item.get("workpaper", "") or item.get("sheet", ""),
                    "rule_id": item.get("rule_id", ""),
                    "field": "",
                    "issue": f"dry-run not writable: {item.get('status_counts')}",
                    "locator": "",
                }
            )
    return gaps


def build_class_summary(records: list[dict[str, Any]], inspected: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for bucket in ("B", "C", "A"):
        bucket_records = [record for record in records if classify_scope(str(record.get("scope") or "")) == bucket]
        bucket_inspected = [item for item in inspected if item.get("bucket") == bucket]
        rows.append(
            {
                "class": bucket,
                "inspected_suggestions": len(bucket_inspected),
                "writable_suggestions": sum(1 for item in bucket_inspected if item.get("writable")),
                "write_items": len(bucket_records),
                "write_status_counts": status_counts(bucket_records, "write_status"),
                "verification_status_counts": status_counts(bucket_records, "verification_status"),
            }
        )
    return rows


def write_report_xlsx(report: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    summary_rows = [
        ("project_id", report["project"]["id"]),
        ("project_name", report["project"]["name"]),
        ("database", report["database"]),
        ("test_copy_dir", report["test_copy_dir"]),
        ("selected_suggestions", len(report["selected_suggestions"])),
        ("write_items", report["summary"]["write_items"]),
        ("passed", report["summary"]["passed"]),
        ("failed", report["summary"]["failed"]),
        ("skipped", report["summary"]["skipped"]),
        ("not_supported", report["summary"]["not_supported"]),
        ("missing_locator", report["summary"]["missing_locator"]),
    ]
    for row in summary_rows:
        ws.append(row)

    ws = wb.create_sheet("Verification")
    verification_headers = [
        "workpaper_id",
        "workpaper_code",
        "workpaper_name",
        "scope",
        "rule_id",
        "target_type",
        "sheet_or_section",
        "target_cell",
        "locator",
        "field",
        "old_value",
        "expected_value",
        "actual_value",
        "write_status",
        "verification_status",
        "verification_message",
        "workbook_path",
    ]
    ws.append(verification_headers)
    for record in report["verification_records"]:
        ws.append([record.get(header, "") for header in verification_headers])

    ws = wb.create_sheet("ClassSummary")
    ws.append(["class", "inspected_suggestions", "writable_suggestions", "write_items", "write_status_counts", "verification_status_counts"])
    for row in report["class_summary"]:
        ws.append([
            row["class"],
            row["inspected_suggestions"],
            row["writable_suggestions"],
            row["write_items"],
            json.dumps(row["write_status_counts"], ensure_ascii=False),
            json.dumps(row["verification_status_counts"], ensure_ascii=False),
        ])

    ws = wb.create_sheet("Gaps")
    gap_headers = ["priority", "scope", "workpaper_code", "rule_id", "field", "issue", "locator"]
    ws.append(gap_headers)
    for gap in report["gaps"]:
        ws.append([gap.get(header, "") for header in gap_headers])

    ws = wb.create_sheet("CopyMapping")
    mapping_headers = ["workpaper_id", "workpaper_code", "workpaper_name", "original_path", "test_copy_path", "copy_status"]
    ws.append(mapping_headers)
    for row in report["copy_mapping"]:
        ws.append([row.get(header, "") for header in mapping_headers])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            max_length = min(max(len(str(cell.value or "")) for cell in column_cells), 80)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(max_length + 2, 12)
    wb.save(path)


def run(
    project_id: int | None,
    per_class: int,
    output_root: Path,
    *,
    max_inspect_per_bucket: int | None = None,
) -> dict[str, Any]:
    safe_root = (BASE_DIR / "tmp" / "autofill_verification").resolve()
    resolved_output_root = output_root.resolve()
    if resolved_output_root != safe_root and safe_root not in resolved_output_root.parents:
        raise RuntimeError(f"output_root must be under safe autofill verification tmp dir: {safe_root}")
    db = SessionLocal()
    try:
        project, workpapers, attachments, contacts, members, selected, selection_gaps, inspected = choose_project(
            db,
            project_id,
            per_class=per_class,
            max_inspect_per_bucket=max_inspect_per_bucket,
        )
        workpaper_dicts = list_dict(workpapers)
        context = project_context(project, contacts, members)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        test_dir = output_root / f"project_{project.id}_{timestamp}"
        test_dir.mkdir(parents=True, exist_ok=True)
        copied_workpapers, copy_mapping, identity_by_path = copy_workpapers(workpapers, test_dir)
        apply_plan = autofill_plan(selected, copied_workpapers, context, apply=True)
        verification_records, summary = verify_plan_items(apply_plan, workpaper_by_path=identity_by_path)
        report = {
            "generated_at": datetime.now().isoformat(timespec="seconds"),
            "database": safe_database_label(),
            "project": {
                "id": project.id,
                "name": project.name,
                "entity_name": project.entity_name,
                "audit_year": project.audit_year,
            },
            "test_copy_dir": str(test_dir),
            "selected_suggestions": [
                {
                    "rule_id": item.get("rule_id", ""),
                    "scope": item.get("scope", ""),
                    "workpaper": item.get("workpaper", ""),
                    "sheet": item.get("sheet", ""),
                    "name": item.get("name", ""),
                }
                for item in selected
            ],
            "summary": summary,
            "class_summary": build_class_summary(verification_records, inspected),
            "verification_records": verification_records,
            "gaps": build_gap_records(inspected, selection_gaps, verification_records),
            "copy_mapping": copy_mapping,
            "inspected_suggestions": inspected,
        }
        json_path = test_dir / "autofill_verification_report.json"
        xlsx_path = test_dir / "autofill_verification_report.xlsx"
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=json_default), encoding="utf-8")
        write_report_xlsx(report, xlsx_path)
        report["report_json"] = str(json_path)
        report["report_xlsx"] = str(xlsx_path)
        return report
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Run autofill write/readback verification on safe test copies.")
    parser.add_argument("--project-id", type=int, default=None, help="Project id to verify; omitted means auto-select.")
    parser.add_argument("--per-class", type=int, default=2, help="Maximum selected suggestions per B/C/A class.")
    parser.add_argument(
        "--max-inspect-per-bucket",
        type=int,
        default=None,
        help="Optional dry-run suggestion inspection cap per B/C/A class.",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=BASE_DIR / "tmp" / "autofill_verification",
        help="Directory for test copies and reports.",
    )
    args = parser.parse_args()
    report = run(args.project_id, args.per_class, args.output_root, max_inspect_per_bucket=args.max_inspect_per_bucket)
    print(
        json.dumps(
            {
                "project_id": report["project"]["id"],
                "project_name": report["project"]["name"],
                "test_copy_dir": report["test_copy_dir"],
                "report_json": report["report_json"],
                "report_xlsx": report["report_xlsx"],
                "summary": report["summary"],
                "class_summary": report["class_summary"],
                "selected_suggestions": report["selected_suggestions"],
            },
            ensure_ascii=False,
            indent=2,
            default=json_default,
        )
    )


if __name__ == "__main__":
    main()
