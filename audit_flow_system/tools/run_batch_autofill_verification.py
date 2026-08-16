from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import datetime
import json
from pathlib import Path
import subprocess
import sys
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from audit_flow_system.core.config import BASE_DIR
from audit_flow_system.core.db import SessionLocal, safe_database_label
from audit_flow_system.models import Attachment, Project, Workpaper
from audit_flow_system.tools.run_autofill_verification import (
    classify_scope,
    existing_workpaper_path,
    json_default,
    status_counts,
)


SAFE_OUTPUT_ROOT = (BASE_DIR / "tmp" / "autofill_verification").resolve()


def safe_json_dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=json_default)


def tail_text(value: str, limit: int = 800) -> str:
    value = (value or "").strip()
    return value if len(value) <= limit else value[-limit:]


def parse_single_project_stdout(stdout: str) -> dict[str, Any]:
    text = stdout.strip()
    if not text:
        raise ValueError("single-project verification returned empty stdout")
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start < 0 or end <= start:
            raise
        return json.loads(text[start : end + 1])


def load_single_project_report(report_json: str | Path) -> dict[str, Any]:
    path = Path(report_json).expanduser().resolve()
    report = json.loads(path.read_text(encoding="utf-8"))
    report["report_json"] = str(path)
    report["report_xlsx"] = str(path.with_suffix(".xlsx"))
    return report


def run_single_project_subprocess(
    *,
    candidate: dict[str, Any],
    per_class: int,
    output_root: Path,
    max_inspect_per_bucket: int,
    project_timeout_seconds: int,
) -> dict[str, Any]:
    cmd = [
        sys.executable,
        "-m",
        "audit_flow_system.tools.run_autofill_verification",
        "--project-id",
        str(candidate["project_id"]),
        "--per-class",
        str(per_class),
        "--output-root",
        str(output_root),
        "--max-inspect-per-bucket",
        str(max_inspect_per_bucket),
    ]
    completed = subprocess.run(
        cmd,
        cwd=str(BASE_DIR.parent),
        text=True,
        capture_output=True,
        timeout=project_timeout_seconds,
        check=False,
    )
    if completed.returncode != 0:
        raise RuntimeError(
            f"single-project verification exited {completed.returncode}; "
            f"stdout={tail_text(completed.stdout)} stderr={tail_text(completed.stderr)}"
        )
    summary = parse_single_project_stdout(completed.stdout)
    return load_single_project_report(summary["report_json"])


def workpaper_class(code: str) -> str:
    code = str(code or "").upper()
    if code.startswith("B"):
        return "B"
    if code.startswith("C"):
        return "C"
    if code.startswith("A"):
        return "A"
    return "Other"


def candidate_sort_key(candidate: dict[str, Any]) -> tuple[int, int, int, int]:
    type_counts = candidate["workpaper_type_counts"]
    has_bca = all(type_counts.get(item, 0) > 0 for item in ("B", "C", "A"))
    return (
        0 if has_bca else 1,
        -candidate["supported_workpaper_count"],
        -candidate["attachment_count"],
        candidate["project_id"],
    )


def inspect_candidates() -> list[dict[str, Any]]:
    db = SessionLocal()
    try:
        projects = db.execute(select(Project).order_by(Project.id)).scalars().all()
        candidates: list[dict[str, Any]] = []
        for project in projects:
            workpapers = db.execute(
                select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.code, Workpaper.id)
            ).scalars().all()
            attachments = db.execute(
                select(Attachment).where(Attachment.project_id == project.id).order_by(Attachment.index_no, Attachment.id)
            ).scalars().all()
            type_counts: Counter[str] = Counter(workpaper_class(workpaper.code) for workpaper in workpapers)
            supported_count = 0
            missing_count = 0
            for workpaper in workpapers:
                if existing_workpaper_path(
                    {
                        "file_path": workpaper.file_path,
                    }
                ):
                    supported_count += 1
                else:
                    missing_count += 1
            skip_reasons: list[str] = []
            if not workpapers:
                skip_reasons.append("no workpapers")
            if not attachments:
                skip_reasons.append("no attachments")
            if supported_count == 0:
                skip_reasons.append("no existing supported workpaper files")
            for item in ("B", "C", "A"):
                if type_counts.get(item, 0) == 0:
                    skip_reasons.append(f"missing {item} workpapers")
            candidates.append(
                {
                    "project_id": project.id,
                    "project_name": project.name,
                    "entity_name": project.entity_name,
                    "audit_year": project.audit_year,
                    "workpaper_count": len(workpapers),
                    "attachment_count": len(attachments),
                    "supported_workpaper_count": supported_count,
                    "missing_or_unsupported_workpaper_count": missing_count,
                    "workpaper_type_counts": dict(type_counts),
                    "precheck_skip_reasons": skip_reasons,
                    "precheck_eligible": not skip_reasons,
                }
            )
        return sorted(candidates, key=candidate_sort_key)
    finally:
        db.close()


def ensure_safe_report(report: dict[str, Any]) -> list[str]:
    problems: list[str] = []
    paths = [report.get("test_copy_dir", ""), report.get("report_json", ""), report.get("report_xlsx", "")]
    for record in report.get("verification_records", []):
        paths.append(record.get("workbook_path", ""))
    for row in report.get("copy_mapping", []):
        paths.append(row.get("test_copy_path", ""))
    for raw_path in paths:
        if not raw_path:
            continue
        path = Path(str(raw_path)).expanduser().resolve()
        if path != SAFE_OUTPUT_ROOT and SAFE_OUTPUT_ROOT not in path.parents:
            problems.append(str(path))
    return problems


def classify_gap_flags(gap: dict[str, Any]) -> dict[str, bool]:
    issue = str(gap.get("issue") or "")
    issue_lower = issue.lower()
    rule_missing = any(text in issue_lower for text in ("no dry-run writable suggestion", "no selected autofill"))
    template_missing_row = any(
        text in issue_lower
        for text in (
            "not found",
            "missing locator",
            "headers not found",
            "label not found",
            "table not found",
            "control code row not found",
            "control code not present",
        )
    )
    evidence_insufficient = any(
        text in issue_lower
        for text in ("missing_evidence", "no attachments", "matched evidence", "evidence", "no writable target")
    )
    writer_unsupported = any(text in issue_lower for text in ("not supported", "unsupported file type", "only single-cell"))
    return {
        "rule_missing": rule_missing,
        "template_missing_row": template_missing_row,
        "evidence_insufficient": evidence_insufficient,
        "writer_unsupported": writer_unsupported,
    }


def project_dimension_rows(reports: list[dict[str, Any]], candidate_by_id: dict[int, dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for report in reports:
        project_id = int(report["project"]["id"])
        candidate = candidate_by_id.get(project_id, {})
        summary = report["summary"]
        rows.append(
            {
                "project_id": project_id,
                "project_name": report["project"]["name"],
                "entity_name": report["project"].get("entity_name", ""),
                "audit_year": report["project"].get("audit_year", ""),
                "workpaper_count": candidate.get("workpaper_count", len(report.get("copy_mapping", []))),
                "attachment_count": candidate.get("attachment_count", ""),
                "supported_workpaper_count": candidate.get("supported_workpaper_count", ""),
                "workpaper_type_counts": candidate.get("workpaper_type_counts", {}),
                "verification_items": summary["write_items"],
                "passed": summary["passed"],
                "failed": summary["failed"],
                "skipped": summary["skipped"],
                "not_supported": summary["not_supported"],
                "missing_locator": summary["missing_locator"],
                "report_json": report.get("report_json", ""),
                "report_xlsx": report.get("report_xlsx", ""),
                "test_copy_dir": report.get("test_copy_dir", ""),
            }
        )
    return rows


def workpaper_dimension_rows(reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for report in reports:
        project_id = report["project"]["id"]
        project_name = report["project"]["name"]
        for record in report.get("verification_records", []):
            key = (
                project_id,
                project_name,
                record.get("workpaper_id"),
                record.get("workpaper_code", ""),
                record.get("workpaper_name", ""),
                classify_scope(str(record.get("scope") or "")),
            )
            grouped[key].append(record)
    rows: list[dict[str, Any]] = []
    for key, records in sorted(grouped.items(), key=lambda item: (item[0][0], str(item[0][3]), str(item[0][2]))):
        counts = status_counts(records, "verification_status")
        rows.append(
            {
                "project_id": key[0],
                "project_name": key[1],
                "workpaper_id": key[2],
                "workpaper_code": key[3],
                "workpaper_name": key[4],
                "class": key[5],
                "verification_items": len(records),
                "passed": counts.get("passed", 0),
                "failed": counts.get("failed", 0),
                "skipped": counts.get("skipped", 0),
                "not_supported": counts.get("not_supported", 0),
                "missing_locator": counts.get("missing_locator", 0),
            }
        )
    return rows


def field_dimension_rows(reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for report in reports:
        project_id = report["project"]["id"]
        project_name = report["project"]["name"]
        for record in report.get("verification_records", []):
            rows.append(
                {
                    "project_id": project_id,
                    "project_name": project_name,
                    "workpaper_id": record.get("workpaper_id", ""),
                    "workpaper_code": record.get("workpaper_code", ""),
                    "workpaper_name": record.get("workpaper_name", ""),
                    "class": classify_scope(str(record.get("scope") or "")),
                    "scope": record.get("scope", ""),
                    "rule_id": record.get("rule_id", ""),
                    "field": record.get("field", ""),
                    "sheet_or_section": record.get("sheet_or_section", ""),
                    "locator": record.get("locator", ""),
                    "cell": record.get("target_cell", ""),
                    "expected": record.get("expected_value", ""),
                    "actual": record.get("actual_value", ""),
                    "write_status": record.get("write_status", ""),
                    "verification_status": record.get("verification_status", ""),
                    "message": record.get("write_message", "") or record.get("verification_message", ""),
                    "workbook_path": record.get("workbook_path", ""),
                }
            )
    return rows


def gap_dimension_rows(reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for report in reports:
        project_id = report["project"]["id"]
        project_name = report["project"]["name"]
        verification_issue_by_key: dict[tuple[str, str, str, str], str] = {}
        for record in report.get("verification_records", []):
            key = (
                str(record.get("workpaper_code") or ""),
                str(record.get("rule_id") or ""),
                str(record.get("field") or ""),
                str(record.get("locator") or ""),
            )
            issue = record.get("write_message") or record.get("verification_message") or ""
            if issue:
                verification_issue_by_key[key] = issue
        for gap in report.get("gaps", []):
            key = (
                str(gap.get("workpaper_code") or ""),
                str(gap.get("rule_id") or ""),
                str(gap.get("field") or ""),
                str(gap.get("locator") or ""),
            )
            issue = verification_issue_by_key.get(key) or gap.get("issue", "")
            flags = classify_gap_flags({**gap, "issue": issue})
            rows.append(
                {
                    "project_id": project_id,
                    "project_name": project_name,
                    "priority": gap.get("priority", ""),
                    "class": classify_scope(str(gap.get("scope") or "")),
                    "scope": gap.get("scope", ""),
                    "workpaper_code": gap.get("workpaper_code", ""),
                    "rule_id": gap.get("rule_id", ""),
                    "field": gap.get("field", ""),
                    "issue": issue,
                    "locator": gap.get("locator", ""),
                    **flags,
                }
            )
    return rows


def class_dimension_rows(reports: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    inspected_counter: dict[str, int] = defaultdict(int)
    writable_counter: dict[str, int] = defaultdict(int)
    for report in reports:
        for record in report.get("verification_records", []):
            grouped[classify_scope(str(record.get("scope") or ""))].append(record)
        for row in report.get("class_summary", []):
            item_class = row["class"]
            inspected_counter[item_class] += row["inspected_suggestions"]
            writable_counter[item_class] += row["writable_suggestions"]
    rows: list[dict[str, Any]] = []
    for item_class in ("B", "C", "A", "Other"):
        records = grouped.get(item_class, [])
        if not records and item_class == "Other":
            continue
        counts = status_counts(records, "verification_status")
        rows.append(
            {
                "class": item_class,
                "inspected_suggestions": inspected_counter.get(item_class, 0),
                "writable_suggestions": writable_counter.get(item_class, 0),
                "verification_items": len(records),
                "passed": counts.get("passed", 0),
                "failed": counts.get("failed", 0),
                "skipped": counts.get("skipped", 0),
                "not_supported": counts.get("not_supported", 0),
                "missing_locator": counts.get("missing_locator", 0),
            }
        )
    return rows


def append_rows(ws, rows: list[dict[str, Any]], headers: list[str]) -> None:
    ws.append(headers)
    for row in rows:
        values = []
        for header in headers:
            value = row.get(header, "")
            if isinstance(value, (dict, list, tuple)):
                value = safe_json_dump(value)
            values.append(value)
        ws.append(values)


def write_batch_xlsx(report: dict[str, Any], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "BatchSummary"
    summary = report["summary"]
    for key in (
        "generated_at",
        "database",
        "batch_dir",
        "requested_project_count",
        "covered_project_count",
        "skipped_project_count",
        "safe_output_root",
        "all_write_targets_safe",
    ):
        ws.append([key, report.get(key, summary.get(key, ""))])
    for key in ("write_items", "passed", "failed", "skipped", "not_supported", "missing_locator"):
        ws.append([key, summary.get(key, 0)])

    sheet_specs = [
        (
            "ProjectDimension",
            report["project_dimension"],
            [
                "project_id",
                "project_name",
                "entity_name",
                "audit_year",
                "workpaper_count",
                "attachment_count",
                "supported_workpaper_count",
                "workpaper_type_counts",
                "verification_items",
                "passed",
                "failed",
                "skipped",
                "not_supported",
                "missing_locator",
                "report_json",
                "report_xlsx",
                "test_copy_dir",
            ],
        ),
        (
            "WorkpaperDimension",
            report["workpaper_dimension"],
            [
                "project_id",
                "project_name",
                "workpaper_id",
                "workpaper_code",
                "workpaper_name",
                "class",
                "verification_items",
                "passed",
                "failed",
                "skipped",
                "not_supported",
                "missing_locator",
            ],
        ),
        (
            "FieldDimension",
            report["field_dimension"],
            [
                "project_id",
                "project_name",
                "workpaper_id",
                "workpaper_code",
                "workpaper_name",
                "class",
                "scope",
                "rule_id",
                "field",
                "sheet_or_section",
                "locator",
                "cell",
                "expected",
                "actual",
                "write_status",
                "verification_status",
                "message",
                "workbook_path",
            ],
        ),
        (
            "GapDimension",
            report["gap_dimension"],
            [
                "project_id",
                "project_name",
                "priority",
                "class",
                "scope",
                "workpaper_code",
                "rule_id",
                "field",
                "issue",
                "locator",
                "rule_missing",
                "template_missing_row",
                "evidence_insufficient",
                "writer_unsupported",
            ],
        ),
        (
            "ClassDimension",
            report["class_dimension"],
            [
                "class",
                "inspected_suggestions",
                "writable_suggestions",
                "verification_items",
                "passed",
                "failed",
                "skipped",
                "not_supported",
                "missing_locator",
            ],
        ),
        (
            "SkippedProjects",
            report["skipped_projects"],
            [
                "project_id",
                "project_name",
                "workpaper_count",
                "attachment_count",
                "supported_workpaper_count",
                "workpaper_type_counts",
                "skip_reason",
            ],
        ),
    ]
    for title, rows, headers in sheet_specs:
        ws = wb.create_sheet(title)
        append_rows(ws, rows, headers)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for column_cells in sheet.columns:
            max_length = min(max(len(str(cell.value or "")) for cell in column_cells), 80)
            sheet.column_dimensions[column_cells[0].column_letter].width = max(max_length + 2, 12)
    wb.save(path)


def build_batch_report(
    *,
    batch_dir: Path,
    candidates: list[dict[str, Any]],
    reports: list[dict[str, Any]],
    skipped_projects: list[dict[str, Any]],
    requested_project_count: int,
    coverage_status: str = "complete",
) -> dict[str, Any]:
    candidate_by_id = {int(item["project_id"]): item for item in candidates}
    summary_counter: Counter[str] = Counter()
    for report in reports:
        summary_counter.update(report["summary"])
    unsafe_paths: list[str] = []
    for report in reports:
        unsafe_paths.extend(ensure_safe_report(report))
    summary = {
        "write_items": summary_counter.get("write_items", 0),
        "passed": summary_counter.get("passed", 0),
        "failed": summary_counter.get("failed", 0),
        "skipped": summary_counter.get("skipped", 0),
        "not_supported": summary_counter.get("not_supported", 0),
        "missing_locator": summary_counter.get("missing_locator", 0),
    }
    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "database": safe_database_label(),
        "batch_dir": str(batch_dir),
        "safe_output_root": str(SAFE_OUTPUT_ROOT),
        "requested_project_count": requested_project_count,
        "covered_project_count": len(reports),
        "skipped_project_count": len(skipped_projects),
        "coverage_status": coverage_status,
        "all_write_targets_safe": not unsafe_paths,
        "unsafe_paths": unsafe_paths,
        "summary": summary,
        "covered_projects": [
            {
                "project_id": report["project"]["id"],
                "project_name": report["project"]["name"],
                "test_copy_dir": report["test_copy_dir"],
                "report_json": report["report_json"],
                "report_xlsx": report["report_xlsx"],
                "summary": report["summary"],
            }
            for report in reports
        ],
        "skipped_projects": skipped_projects,
        "project_dimension": project_dimension_rows(reports, candidate_by_id),
        "workpaper_dimension": workpaper_dimension_rows(reports),
        "field_dimension": field_dimension_rows(reports),
        "gap_dimension": gap_dimension_rows(reports),
        "class_dimension": class_dimension_rows(reports),
    }


def save_batch_report(batch_report: dict[str, Any], batch_dir: Path) -> dict[str, Any]:
    json_path = batch_dir / "batch_autofill_verification_report.json"
    xlsx_path = batch_dir / "batch_autofill_verification_report.xlsx"
    json_path.write_text(json.dumps(batch_report, ensure_ascii=False, indent=2, default=json_default), encoding="utf-8")
    write_batch_xlsx(batch_report, xlsx_path)
    batch_report["report_json"] = str(json_path)
    batch_report["report_xlsx"] = str(xlsx_path)
    return batch_report


def run_batch(
    project_count: int,
    per_class: int,
    output_root: Path,
    max_inspect_per_bucket: int,
    project_timeout_seconds: int = 300,
) -> dict[str, Any]:
    resolved_output_root = output_root.resolve()
    if resolved_output_root != SAFE_OUTPUT_ROOT:
        raise RuntimeError(f"batch output_root must be {SAFE_OUTPUT_ROOT}")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    batch_dir = resolved_output_root / f"batch_{timestamp}"
    batch_dir.mkdir(parents=True, exist_ok=True)
    candidates = inspect_candidates()
    reports: list[dict[str, Any]] = []
    skipped_projects: list[dict[str, Any]] = []

    for candidate in candidates:
        if len(reports) >= project_count:
            break
        if not candidate["precheck_eligible"]:
            skipped_projects.append(
                {
                    **candidate,
                    "skip_reason": "; ".join(candidate["precheck_skip_reasons"]),
                }
            )
            continue
        print(
            f"starting project {candidate['project_id']} {candidate['project_name']} "
            f"(workpapers={candidate['workpaper_count']}, attachments={candidate['attachment_count']})",
            flush=True,
        )
        try:
            report = run_single_project_subprocess(
                candidate=candidate,
                per_class=per_class,
                output_root=resolved_output_root,
                max_inspect_per_bucket=max_inspect_per_bucket,
                project_timeout_seconds=project_timeout_seconds,
            )
        except subprocess.TimeoutExpired:
            skipped_projects.append(
                {
                    **candidate,
                    "skip_reason": f"verification timed out after {project_timeout_seconds}s",
                }
            )
            print(
                f"skipped project {candidate['project_id']} timeout after {project_timeout_seconds}s",
                flush=True,
            )
            continue
        except Exception as exc:
            skipped_projects.append(
                {
                    **candidate,
                    "skip_reason": f"verification failed before report: {exc}",
                }
            )
            continue
        unsafe_paths = ensure_safe_report(report)
        if unsafe_paths:
            raise RuntimeError("unsafe autofill verification paths detected: " + "; ".join(unsafe_paths))
        reports.append(report)
        print(
            f"finished project {candidate['project_id']} summary={safe_json_dump(report['summary'])}",
            flush=True,
        )

    batch_report = build_batch_report(
        batch_dir=batch_dir,
        candidates=candidates,
        reports=reports,
        skipped_projects=skipped_projects,
        requested_project_count=project_count,
    )
    return save_batch_report(batch_report, batch_dir)


def parse_skip_arg(value: str) -> dict[str, Any]:
    parts = value.split("|", 2)
    if len(parts) != 3:
        raise ValueError("--skip-project must use project_id|project_name|reason")
    project_id, project_name, reason = parts
    return {
        "project_id": int(project_id),
        "project_name": project_name,
        "workpaper_count": "",
        "attachment_count": "",
        "supported_workpaper_count": "",
        "workpaper_type_counts": {},
        "skip_reason": reason,
    }


def recover_partial_batch(report_paths: list[Path], skip_project_args: list[str], output_root: Path) -> dict[str, Any]:
    resolved_output_root = output_root.resolve()
    if resolved_output_root != SAFE_OUTPUT_ROOT:
        raise RuntimeError(f"recovery output_root must be {SAFE_OUTPUT_ROOT}")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    batch_dir = resolved_output_root / f"batch_{timestamp}_partial"
    batch_dir.mkdir(parents=True, exist_ok=True)
    reports: list[dict[str, Any]] = []
    for path in report_paths:
        report = json.loads(path.read_text(encoding="utf-8"))
        report["report_json"] = str(path)
        report["report_xlsx"] = str(path.with_suffix(".xlsx"))
        reports.append(report)
    candidates = inspect_candidates()
    candidate_by_id = {int(item["project_id"]): item for item in candidates}
    skipped_projects: list[dict[str, Any]] = []
    for raw in skip_project_args:
        row = parse_skip_arg(raw)
        candidate = candidate_by_id.get(int(row["project_id"]))
        if candidate:
            row = {**candidate, "skip_reason": row["skip_reason"]}
        skipped_projects.append(row)
    batch_report = build_batch_report(
        batch_dir=batch_dir,
        candidates=candidates,
        reports=reports,
        skipped_projects=skipped_projects,
        requested_project_count=len(reports) + len(skipped_projects),
        coverage_status="partial",
    )
    return save_batch_report(batch_report, batch_dir)


def print_report_summary(report: dict[str, Any]) -> None:
    print(
        json.dumps(
            {
                "batch_dir": report["batch_dir"],
                "coverage_status": report.get("coverage_status", ""),
                "report_json": report["report_json"],
                "report_xlsx": report["report_xlsx"],
                "covered_projects": report["covered_projects"],
                "skipped_project_count": report["skipped_project_count"],
                "skipped_projects": report["skipped_projects"],
                "summary": report["summary"],
                "class_dimension": report["class_dimension"],
                "all_write_targets_safe": report["all_write_targets_safe"],
            },
            ensure_ascii=False,
            indent=2,
            default=json_default,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Run batch autofill verification on safe test copies.")
    parser.add_argument("--project-count", type=int, default=5, help="Number of eligible projects to verify.")
    parser.add_argument("--per-class", type=int, default=2, help="Maximum selected suggestions per B/C/A class.")
    parser.add_argument(
        "--max-inspect-per-bucket",
        type=int,
        default=12,
        help="Dry-run suggestion inspection cap per B/C/A class for batch performance.",
    )
    parser.add_argument(
        "--output-root",
        type=Path,
        default=SAFE_OUTPUT_ROOT,
        help="Safe autofill verification root directory.",
    )
    parser.add_argument(
        "--recover-report-json",
        type=Path,
        action="append",
        default=[],
        help="Existing single-project report JSON to aggregate without rerunning verification.",
    )
    parser.add_argument(
        "--skip-project",
        action="append",
        default=[],
        help="Skipped project in project_id|project_name|reason format for recovery mode.",
    )
    parser.add_argument(
        "--project-timeout-seconds",
        type=int,
        default=300,
        help="Maximum seconds for one project verification subprocess before it is skipped.",
    )
    args = parser.parse_args()
    if args.recover_report_json:
        report = recover_partial_batch(args.recover_report_json, args.skip_project, args.output_root)
    else:
        report = run_batch(
            args.project_count,
            args.per_class,
            args.output_root,
            args.max_inspect_per_bucket,
            project_timeout_seconds=args.project_timeout_seconds,
        )
    print_report_summary(report)


if __name__ == "__main__":
    main()
