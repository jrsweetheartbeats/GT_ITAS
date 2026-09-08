from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


SHEETS = (("C22 (问题标准集)", "C22"), ("非C22 (问题标准集)", "非C22"))


def build_catalog(source: Path) -> list[dict[str, object]]:
    workbook = load_workbook(source, data_only=True, read_only=True)
    rows: list[dict[str, object]] = []
    for sheet_name, scope in SHEETS:
        sheet = workbook[sheet_name]
        current_workpaper = ""
        current_workpaper_no = ""
        for row_no, values in enumerate(sheet.iter_rows(min_row=3, max_col=12, values_only=True), start=3):
            workpaper = str(values[1] or "").strip()
            if workpaper:
                current_workpaper = workpaper
            if values[2] not in (None, ""):
                current_workpaper_no = str(values[2]).strip()
            description = str(values[6] or "").strip()
            if not description:
                continue
            index_code = str(values[3] or "").strip()
            rows.append(
                {
                    "id": f"{scope}-{index_code or row_no}",
                    "scope": scope,
                    "workpaper": current_workpaper,
                    "workpaper_no": current_workpaper_no,
                    "index_code": index_code,
                    "issue_no": str(values[4] or "").strip(),
                    "source_step": str(values[5] or "").strip(),
                    "description": description,
                    "category": str(values[7] or "").strip(),
                    "self_conflict": str(values[8] or "").strip(),
                    "updated": str(values[9] or "").strip(),
                    "new": str(values[10] or "").strip(),
                    "maintainer_note": str(values[11] or "").strip(),
                    "source_sheet": sheet_name,
                    "source_row": row_no,
                }
            )
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the ITAS review issue catalog from the review workbook.")
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "data" / "review_issue_catalog.json")
    args = parser.parse_args()
    payload = build_catalog(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {len(payload)} review issue standards to {args.output}")


if __name__ == "__main__":
    main()
