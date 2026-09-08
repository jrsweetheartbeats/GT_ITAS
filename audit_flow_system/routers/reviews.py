from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import sys
from typing import Any, Optional

from fastapi import APIRouter, Body, Depends, File, Header, HTTPException, Query, Response, UploadFile
from fastapi.responses import FileResponse
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import and_, func, or_, select
from sqlalchemy.orm import Session

from ..core.config import BASE_DIR, WORKSPACE_ROOT
from ..core.db import get_db, safe_database_label
from ..core.security import (
    DEFAULT_MODULE_ORDER,
    DEFAULT_PASSWORD_POLICY,
    can_edit_project,
    can_review_project,
    can_upload_documents,
    current_user,
    default_audit_scope,
    ensure_document_uploader,
    ensure_project_editor,
    ensure_project_reviewer,
    ensure_project_viewer,
    ensure_workpaper_viewer,
    get_setting,
    hash_password,
    is_admin,
    is_project_member,
    normalize_module_order,
    require_admin,
    set_setting,
    supervised_project_ids,
    validate_password_policy,
    visible_project_ids,
    verify_password,
)
from ..core.utils import apply_patch_to_model, get_or_404, list_dict, obj_dict
from ..models import (
    Attachment,
    AutomationRule,
    AutofillPlanItem,
    AutofillRun,
    Client,
    ClientITContact,
    DocumentRequest,
    EnterpriseContact,
    LoginSession,
    Project,
    ProjectMember,
    ReviewFinding,
    ReviewFindingHistory,
    ReviewFindingResponse,
    ReviewRun,
    ReviewStep,
    ReviewStepHistory,
    Role,
    Task,
    User,
    Workpaper,
    WorkpaperVersion,
)
from ..schemas import (
    AttachmentIn,
    AttachmentReferenceScanIn,
    AttachmentScanIn,
    AutomationRuleIn,
    AutofillPlanIn,
    ClientITContactIn,
    ClientIn,
    ContactIn,
    DocumentRequestUploadIn,
    InitFromPriorIn,
    LoginIn,
    MemberIn,
    PasswordPolicyIn,
    ProjectIn,
    ReviewFindingDecisionIn,
    ReviewDecisionIn,
    ReviewFindingIn,
    ReviewFindingPatchIn,
    ReviewFindingResponseIn,
    ReviewRunIn,
    RoleIn,
    TaskIn,
    UserIn,
    WorkpaperIn,
    WorkpaperSubmitIn,
    BatchWorkpaperSubmitIn,
)


from ..services.attachments import (
    append_reference,
    candidate_reference_tokens,
    extract_workpaper_text,
    file_type_for_path,
    guess_workpaper_for_file,
    likely_attachment_reference,
    next_attachment_index,
    safe_project_code,
    sanitize_workpaper_code,
    scan_project_attachment_files,
)
from ..services.materials import c22_document_requests_from_rules, save_upload_files
from ..services.projects import copy_workpaper_file, replace_audit_year, seed_project_template_workpapers
from ..services.review import add_finding, run_external_rules, run_internal_review
from ..services.review_catalog import filter_review_issue_catalog, review_issue_catalog
from ..services.deepseek_review import DeepSeekReviewError, run_deepseek_review
from ..services.timeliness import days_after, overdue_payload, waiting_days
from ..services.privacy import project_redaction_terms, redact_text
from ..services.audit_log import record_audit_log


router = APIRouter()


OPEN_FINDING_STATUSES = {
    "open",
    "assigned",
    "changes_requested",
    "responded",
    "retained",
    "revised",
    "待处理",
    "已分派",
    "保留",
    "已修订",
}

STANDARD_FINDING_HEADERS = [
    "问题编号", "审计阶段", "问题类型", "严重程度", "是否C22相关", "底稿文件", "页签/位置",
    "问题描述", "具体描述补充", "问题步骤归属", "问题类别", "项目组回复", "确认复核问题已满意解决",
    "问题出现的复核阶段", "现场负责人", "项目组内复核人", "建议处理", "复核状态",
]

_FINDING_HEADER_ALIASES = {
    "issue_no": {"问题编号", "issue_no", "编号"},
    "audit_stage": {"审计阶段", "阶段", "audit_stage"},
    "finding_type": {"问题类型", "finding_type", "例如ITGCITAC等"},
    "rule_code": {"规则代码", "rule_code", "规则"},
    "severity": {"严重程度", "severity", "优先级", "级别"},
    "c22_related": {"是否C22相关", "c22_related", "C22相关"},
    "workpaper_file": {"底稿文件", "workpaper_file", "来源文件"},
    "location": {"页签/位置", "location", "页签/单元格", "底稿位置"},
    "issue": {"问题描述", "issue", "问题"},
    "evidence": {"当前证据", "evidence", "证据", "具体描述补充"},
    "issue_step": {"问题步骤归属", "issue_step"},
    "issue_category": {"问题类别", "issue_category", "类别"},
    "recommendation": {"建议处理", "recommendation", "建议", "整改建议"},
    "status": {"复核状态", "status", "状态"},
    "project_reply": {"项目组回复", "project_reply", "回复"},
    "resolution_confirmed": {"确认复核问题已满意解决", "resolution_confirmed", "确认解决"},
    "review_stage": {"问题出现的复核阶段", "review_stage", "复核阶段"},
    "field_lead": {"现场负责人", "现场负人", "field_lead"},
    "project_reviewer": {"项目组内复核人", "合伙人/总监/高经-项目审核人", "project_reviewer"},
}


def _normalise_import_header(value: Any) -> str:
    return re.sub(r"[\s（）()、，,：:]+", "", str(value or "").strip())


def _normalise_severity(value: Any) -> str:
    text = str(value or "medium").strip().lower()
    return {"高": "high", "中": "medium", "低": "low", "重大": "high", "一般": "medium"}.get(text, text if text in {"high", "medium", "low"} else "medium")


def _normalise_bool(value: Any) -> bool:
    return str(value or "").strip().lower() in {"是", "y", "yes", "true", "1", "相关"}


def _normalise_status(value: Any) -> str:
    text = str(value or "open").strip()
    return {"待补充/待整改": "open", "待处理": "open", "整改中": "assigned", "已分派": "assigned", "已整改": "resolved", "已关闭": "resolved"}.get(text, text or "open")


def _parse_full_review_record_sheet(sheet: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for row in sheet.iter_rows(min_row=2, max_col=24, values_only=True):
        issue_no, audit_stage, finding_type, issue, evidence, issue_step, issue_category, workpaper_file = row[:8]
        if not str(issue or "").strip() or not str(workpaper_file or "").strip():
            continue
        confirmed = _normalise_bool(row[10])
        project_reply = str(row[9] or "").strip()
        scope = str(row[12] or "").strip()
        rows.append(
            {
                "issue_no": str(issue_no or "").strip(),
                "standard_index_code": "",
                "source": "review_workbook",
                "rule_code": "REVIEW_WORKBOOK",
                "audit_stage": str(audit_stage or "").strip(),
                "finding_type": str(finding_type or "").strip(),
                "issue_step": str(issue_step or "").strip(),
                "issue_category": str(issue_category or "").strip(),
                "severity": "high" if _normalise_bool(row[20]) else "medium",
                "c22_related": scope.upper() == "C22",
                "workpaper_file": str(workpaper_file or "").strip(),
                "location": str(issue_step or "").strip(),
                "target": str(workpaper_file or "").strip(),
                "issue": str(issue or "").strip(),
                "evidence": str(evidence or "").strip(),
                "recommendation": "",
                "project_reply": project_reply,
                "resolution_confirmed": confirmed,
                "review_stage": str(row[8] or "").strip(),
                "field_lead": str(row[13] or "").strip(),
                "project_reviewer": str(row[14] or "").strip(),
                "note": "",
                "status": "resolved" if confirmed else ("responded" if project_reply else "open"),
                "review_comment": "",
            }
        )
    return rows


def _parse_import_rows(contents: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    workbook = load_workbook(BytesIO(contents), data_only=True, read_only=True)
    if "问题汇总" in workbook.sheetnames:
        rows = _parse_full_review_record_sheet(workbook["问题汇总"])
        return (rows, []) if rows else ([], ["问题汇总页未识别到可导入的复核问题"])
    sheet = workbook.active
    values = list(sheet.iter_rows(values_only=True))
    if not values:
        return [], ["工作簿为空"]
    header_row = [_normalise_import_header(value) for value in values[0]]
    index_map: dict[str, int] = {}
    for field, aliases in _FINDING_HEADER_ALIASES.items():
        for index, header in enumerate(header_row):
            if header in {_normalise_import_header(alias) for alias in aliases}:
                index_map[field] = index
                break
    required = {"issue_no", "issue"}
    missing = sorted(field for field in required if field not in index_map)
    if missing:
        return [], [f"缺少必填列：{'、'.join(missing)}；标准列为：{'、'.join(STANDARD_FINDING_HEADERS)}"]
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    for row_no, row in enumerate(values[1:], start=2):
        if not any(value not in (None, "") for value in row):
            continue
        item: dict[str, Any] = {}
        for field, index in index_map.items():
            item[field] = row[index] if index < len(row) else None
        item["issue_no"] = str(item.get("issue_no") or "").strip()
        item["issue"] = str(item.get("issue") or "").strip()
        if not item["issue_no"] or not item["issue"]:
            errors.append(f"第{row_no}行缺少问题编号或问题描述")
            continue
        item["severity"] = _normalise_severity(item.get("severity"))
        item["c22_related"] = _normalise_bool(item.get("c22_related"))
        item["resolution_confirmed"] = _normalise_bool(item.get("resolution_confirmed"))
        item["status"] = _normalise_status(item.get("status"))
        item["source"] = "codex_import"
        item["workpaper_file"] = str(item.get("workpaper_file") or "").strip()
        item["location"] = str(item.get("location") or "").strip()
        item["target"] = " / ".join(value for value in [item["workpaper_file"], item["location"]] if value)
        item["evidence"] = str(item.get("evidence") or "").strip()
        item["recommendation"] = str(item.get("recommendation") or "").strip()
        item["project_reply"] = str(item.get("project_reply") or "").strip()
        for field in ("audit_stage", "finding_type", "issue_step", "issue_category", "review_stage", "field_lead", "project_reviewer", "standard_index_code"):
            item[field] = str(item.get(field) or "").strip()
        item["note"] = ""
        item["review_comment"] = ""
        if item["resolution_confirmed"]:
            item["status"] = "resolved"
        elif item["project_reply"] and item["status"] == "open":
            item["status"] = "responded"
        rows.append(item)
    return rows, errors


def _history_payload(history: ReviewFindingHistory) -> dict[str, Any]:
    data = obj_dict(history)
    data["operator_name"] = history.operator.display_name if history.operator else ""
    return data


_FINDING_FIELD_LABELS = {
    "issue_no": "序号", "standard_index_code": "标准问题编号", "source": "来源", "rule_code": "问题规则类型",
    "audit_stage": "阶段", "finding_type": "问题类型", "issue_step": "问题步骤归属", "issue_category": "问题类别",
    "severity": "严重程度", "c22_related": "C22与否", "workpaper_file": "对应底稿", "location": "位置",
    "target": "底稿/位置", "issue": "复核问题", "evidence": "具体描述补充", "recommendation": "建议处理",
    "project_reply": "项目组回复", "resolution_confirmed": "确认满意解决", "review_stage": "复核阶段",
    "field_lead": "现场负责人", "project_reviewer": "项目组内复核人", "status": "状态",
    "assignee_user_id": "责任人", "due_date": "截止日期", "review_comment": "复核意见",
}


def _finding_change_text(key: str, old_value: Any, new_value: Any) -> str:
    def display(value: Any) -> str:
        if value is True:
            return "是"
        if value is False:
            return "否"
        return str(value) if value not in (None, "") else "空"

    return f"{_FINDING_FIELD_LABELS.get(key, key)}：{display(old_value)} -> {display(new_value)}"


def _matched_workpaper_for_finding(db: Session, finding: ReviewFinding, run: ReviewRun) -> Workpaper | None:
    if finding.workpaper_id:
        item = db.get(Workpaper, finding.workpaper_id)
        if item:
            return item
    if run.workpaper_id:
        item = db.get(Workpaper, run.workpaper_id)
        if item:
            return item
    text = f"{finding.target or ''} {finding.issue or ''}".lower()
    if not text.strip():
        return None
    attachments = db.execute(
        select(Attachment).where(Attachment.project_id == run.project_id).order_by(Attachment.index_no)
    ).scalars().all()
    for attachment in attachments:
        if attachment.workpaper_id and attachment.index_no and attachment.index_no.lower() in text:
            item = db.get(Workpaper, attachment.workpaper_id)
            if item:
                return item
    workpapers = db.execute(
        select(Workpaper).where(Workpaper.project_id == run.project_id).order_by(Workpaper.code)
    ).scalars().all()
    for item in sorted(workpapers, key=lambda row: len(row.code or ""), reverse=True):
        code = (item.code or "").lower()
        name = (item.name or "").lower()
        if (code and code in text) or (name and name in text):
            return item
    return None


def _finding_payload(
    finding: ReviewFinding,
    run: ReviewRun,
    project: Project,
    db: Session | None = None,
    current_user: User | None = None,
) -> dict[str, Any]:
    data = obj_dict(finding)
    data["run_id"] = run.id
    data["project_id"] = project.id
    data["project_name"] = project.name
    data["entity_name"] = project.entity_name
    data["run_mode"] = run.mode
    data["run_status"] = run.status
    data["run_summary"] = run.summary
    discovered_at = finding.created_at or run.started_at or datetime.utcnow()
    due_date = finding.due_date or days_after(discovered_at.date(), 5)
    data["discovered_at"] = discovered_at
    data.update(overdue_payload(finding.status or "open", due_date))
    data["review_comment"] = finding.review_comment
    data["created_by_name"] = finding.created_by.display_name if finding.created_by else "未记录"
    data["created_by_username"] = finding.created_by.username if finding.created_by else ""
    data["assignee_user_id"] = finding.assignee_user_id
    data["assignee_name"] = finding.assignee.display_name if finding.assignee else ""
    if db is not None:
        workpaper = _matched_workpaper_for_finding(db, finding, run)
        owner = finding.assignee or (workpaper.preparer if workpaper else project.project_leader)
        data["workpaper_id"] = workpaper.id if workpaper else run.workpaper_id
        data["workpaper_code"] = workpaper.code if workpaper else ""
        data["workpaper_name"] = workpaper.name if workpaper else ""
        data["owner_user_id"] = owner.id if owner else None
        data["owner_name"] = owner.display_name if owner else ""
        data["workpaper_version_id"] = finding.workpaper_version_id
        data["resolved_workpaper_version_id"] = finding.resolved_workpaper_version_id
        data["review_step_id"] = finding.review_step_id
        version = db.get(WorkpaperVersion, finding.workpaper_version_id) if finding.workpaper_version_id else None
        resolved_version = db.get(WorkpaperVersion, finding.resolved_workpaper_version_id) if finding.resolved_workpaper_version_id else None
        data["workpaper_version_no"] = version.version_no if version else None
        data["resolved_workpaper_version_no"] = resolved_version.version_no if resolved_version else None
        response_rows = db.execute(
            select(ReviewFindingResponse)
            .where(ReviewFindingResponse.finding_id == finding.id)
            .order_by(ReviewFindingResponse.round_no)
        ).scalars().all()
        data["responses"] = [
            {
                **obj_dict(response),
                "responder_name": response.responder.display_name if response.responder else "",
            }
            for response in response_rows
        ]
        histories = db.execute(
            select(ReviewFindingHistory)
            .where(ReviewFindingHistory.finding_id == finding.id)
            .order_by(ReviewFindingHistory.id.desc())
        ).scalars().all()
        data["histories"] = [_history_payload(row) for row in histories]
        if current_user is not None:
            data["can_review"] = can_review_project(current_user, project)
            data["can_reply"] = bool(
                can_edit_project(current_user, project)
                or current_user.id == project.field_leader_user_id
                or is_project_member(db, project.id, current_user.id)
            )
    return data


def _setting_int(db: Session, key: str, default: int) -> int:
    try:
        return int(get_setting(db, key, default))
    except (TypeError, ValueError):
        return default


def _role_sla_days(db: Session, role_code: str) -> int:
    default_sla = _setting_int(db, "review_sla_days", 3)
    return _setting_int(db, f"review_sla_days_{role_code}", default_sla)


def _enter_review_step(db: Session, step: ReviewStep, now: datetime | None = None) -> None:
    current = now or datetime.utcnow()
    step.entered_at = current
    step.sla_days = _role_sla_days(db, step.reviewer_role_code)
    step.due_date = days_after(current.date(), step.sla_days)


def _project_review_assignments(project: Project) -> list[tuple[str, int]]:
    candidates = [
        ("project_manager", project.project_leader_user_id),
        ("responsible_manager", project.manager_user_id),
        ("partner", project.partner_user_id),
        ("quality", project.quality_reviewer_user_id),
    ]
    return [(role_code, user_id) for role_code, user_id in candidates if user_id is not None]


def _assignments_from_selected_reviewer(project: Project, reviewer_user_id: int | None) -> list[tuple[str, int]]:
    """Return the configured review chain, optionally starting at the selected reviewer.

    A preparer may only choose someone already configured on this project's
    review chain.  Selecting a later reviewer deliberately skips earlier
    configured levels, while all later levels are still retained.
    """
    assignments = _project_review_assignments(project)
    if not assignments:
        raise HTTPException(status_code=400, detail="请先在项目中配置项目经理、项目负责经理、合伙人或质控复核人")
    if reviewer_user_id is None:
        return assignments
    for index, (_, configured_user_id) in enumerate(assignments):
        if configured_user_id == reviewer_user_id:
            return assignments[index:]
    raise HTTPException(status_code=400, detail="所选下一层复核人不在本项目已配置的复核链中")


def _quality_reviewer_ids(project: Project) -> set[int]:
    try:
        values = json.loads(project.quality_reviewer_user_ids_json or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        values = []
    values = values if isinstance(values, list) else []
    if project.quality_reviewer_user_id:
        values = [project.quality_reviewer_user_id, *values]
    return {int(value) for value in values if str(value).isdigit() and int(value) > 0}


def _ensure_step_actor(step: ReviewStep, user: User, project: Project) -> None:
    is_quality_reviewer = step.reviewer_role_code == "quality" and user.id in _quality_reviewer_ids(project)
    if not is_admin(user) and step.reviewer_user_id != user.id and not is_quality_reviewer:
        raise HTTPException(status_code=403, detail="仅当前复核步骤指定人员可操作")
    if step.status != "pending":
        raise HTTPException(status_code=409, detail=f"复核步骤当前为 {step.status}，不可重复或越级处理")


def _add_step_history(
    db: Session,
    step: ReviewStep,
    user: User,
    action: str,
    old_status: str,
    new_status: str,
    comment: str = "",
) -> None:
    db.add(
        ReviewStepHistory(
            step_id=step.id,
            workpaper_version_id=step.workpaper_version_id,
            operator_user_id=user.id,
            round_no=step.round_no,
            action=action,
            old_status=old_status,
            new_status=new_status,
            comment=comment,
        )
    )


def _review_step_payload(
    step: ReviewStep,
    workpaper: Workpaper | None = None,
    project: Project | None = None,
    current_user: User | None = None,
) -> dict[str, Any]:
    data = obj_dict(step)
    data.update(overdue_payload(step.status, step.due_date))
    data["waiting_days"] = waiting_days(step.entered_at)
    if workpaper is not None:
        data["workpaper_id"] = workpaper.id
        data["workpaper_code"] = workpaper.code
        data["workpaper_name"] = workpaper.name
        data["project_id"] = workpaper.project_id
    if project is not None:
        data["project_name"] = project.name
        data["entity_name"] = project.entity_name
    data["reviewer_name"] = step.reviewer.display_name if step.reviewer else ""
    data["histories"] = [
        {
            **obj_dict(history),
            "operator_name": history.operator.display_name if history.operator else "",
        }
        for history in sorted(step.histories, key=lambda item: item.id)
    ]
    if current_user is not None:
        data["can_decide"] = bool(
            step.status == "pending" and (is_admin(current_user) or step.reviewer_user_id == current_user.id)
        )
    return data


def _validate_resubmission_version(rejected: ReviewStep, latest_version: WorkpaperVersion) -> None:
    if rejected.workpaper_version_id == latest_version.id:
        raise HTTPException(status_code=409, detail="退回后必须先上传新版本底稿，才能重新提交")
    if rejected.workpaper_version_id is None and rejected.reviewed_at and latest_version.created_at < rejected.reviewed_at:
        raise HTTPException(status_code=409, detail="未检测到退回后上传的新版本底稿")


def _finding_query():
    return (
        select(ReviewFinding, ReviewRun, Project)
        .join(ReviewRun, ReviewFinding.run_id == ReviewRun.id)
        .join(Project, ReviewRun.project_id == Project.id)
    )


def _is_open_finding(status: str) -> bool:
    return (status or "open") in OPEN_FINDING_STATUSES


_CLOSED_REVIEW_STAGE_STATUSES = {"closed", "resolved"}
_CHINESE_STAGE_NUMBERS = {
    "一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
    "六": 6, "七": 7, "八": 8, "九": 9, "十": 10,
}


def _review_stage_number(value: str) -> int | None:
    text = str(value or "").strip()
    digit = re.search(r"第\s*(\d+)\s*(?:阶段|轮)", text)
    if digit:
        return int(digit.group(1))
    chinese = re.search(r"第\s*([一二三四五六七八九十])\s*(?:阶段|轮)", text)
    return _CHINESE_STAGE_NUMBERS.get(chinese.group(1)) if chinese else None


def _review_stage_label(number: int) -> str:
    chinese = next((label for label, value in _CHINESE_STAGE_NUMBERS.items() if value == number), None)
    return f"第{chinese or number}阶段"


def _next_review_stage_from_rows(rows: list[tuple[str, str]]) -> str:
    if not rows:
        return _review_stage_label(1)
    normalized = [(str(stage or "").strip() or _review_stage_label(1), str(status or "open").lower()) for stage, status in rows]
    current_stage = normalized[-1][0]
    current_statuses = [status for stage, status in normalized if stage == current_stage]
    if current_statuses and not all(status in _CLOSED_REVIEW_STAGE_STATUSES for status in current_statuses):
        return current_stage
    known_numbers = [number for stage, _ in normalized if (number := _review_stage_number(stage)) is not None]
    next_number = max(known_numbers, default=len({stage for stage, _ in normalized})) + 1
    return _review_stage_label(next_number)


def _automatic_review_stage(db: Session, project_id: int) -> str:
    rows = db.execute(
        select(ReviewFinding.review_stage, ReviewFinding.status)
        .join(ReviewRun, ReviewFinding.run_id == ReviewRun.id)
        .where(ReviewRun.project_id == project_id)
        .order_by(ReviewFinding.id)
    ).all()
    return _next_review_stage_from_rows(rows)


@router.get("/api/automation-rules")
def list_automation_rules(
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    return list_dict(db.execute(select(AutomationRule).order_by(AutomationRule.code)).scalars().all())


@router.post("/api/automation-rules", status_code=201)
def create_automation_rule(
    body: AutomationRuleIn,
    db: Session = Depends(get_db),
    user: User = Depends(require_admin),
) -> dict[str, Any]:
    item = AutomationRule(**body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.post("/api/review-runs", status_code=201)
def create_review_run(body: ReviewRunIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    project = get_or_404(db, Project, body.project_id, "项目")
    ensure_project_reviewer(project, user)
    engine = "external_rules" if body.use_external_rules else body.review_engine
    if body.workpaper_id is not None:
        workpaper = get_or_404(db, Workpaper, body.workpaper_id, "底稿")
        if workpaper.project_id != project.id:
            raise HTTPException(status_code=400, detail="指定底稿不属于当前项目")
        if engine == "external_rules":
            raise HTTPException(status_code=400, detail="外部规则当前仅支持项目级复核，请取消指定底稿")
    run = ReviewRun(
        project_id=body.project_id,
        workpaper_id=body.workpaper_id,
        mode="ai" if engine == "deepseek" else "auto",
        status="running",
    )
    db.add(run)
    db.commit()
    db.refresh(run)

    run.started_at = datetime.utcnow()
    try:
        if engine == "deepseek":
            result = run_deepseek_review(db, run)
            count = int(result["count"])
            rejected = int(result["rejected_count"])
            details = str(result["summary"] or "").strip()
            run.summary = f"DeepSeek 智能复核完成，形成 {count} 项有证据的问题"
            if rejected:
                run.summary += f"，忽略 {rejected} 项字段不完整或底稿编码无效的结果"
            if details:
                run.summary += f"。{details}"
        elif engine == "external_rules":
            count = run_external_rules(db, run)
            run.summary = f"外部规则复核完成，发现 {count} 项问题"
        else:
            count = run_internal_review(db, run)
            run.summary = f"自动规则复核完成，发现 {count} 项问题"
        run.status = "completed"
    except DeepSeekReviewError as exc:
        run.status = "failed"
        run.summary = f"DeepSeek 智能复核失败：{exc}"
        add_finding(db, run, "AI-SYS-001", "high", "review-run", run.summary, source="ai")
    except Exception as exc:
        run.status = "failed"
        run.summary = f"自动复核失败：{exc}"
        add_finding(db, run, "SYS-001", "high", "review-run", run.summary)
    finally:
        run.finished_at = datetime.utcnow()
        db.commit()
        db.refresh(run)
    return obj_dict(run)


@router.get("/api/review-runs")
def list_review_runs(
    projectId: Optional[int] = None,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    visible_ids = visible_project_ids(db, user)
    if not visible_ids:
        return []
    stmt = select(ReviewRun).where(ReviewRun.project_id.in_(visible_ids)).order_by(ReviewRun.id.desc())
    if projectId is not None:
        project = get_or_404(db, Project, projectId, "项目")
        ensure_project_viewer(db, project, user)
        stmt = stmt.where(ReviewRun.project_id == projectId)
    rows = db.execute(stmt).scalars().all()
    supervised_ids = supervised_project_ids(db, user)
    rows = [
        row
        for row in rows
        if row.project_id in supervised_ids or any(finding.assignee_user_id == user.id for finding in row.findings)
    ]
    payload = []
    for row in rows:
        data = obj_dict(row)
        data["finding_count"] = len(row.findings)
        payload.append(data)
    return payload


@router.get("/api/projects/{project_id}/llm-review/redacted-preview")
def llm_review_redacted_preview(
    project_id: int,
    workpaperId: Optional[int] = Query(default=None),
    maxChars: int = Query(default=2000, ge=200, le=20000),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    terms = project_redaction_terms(db, project)
    stmt = select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.code)
    if project.id not in supervised_project_ids(db, user):
        stmt = stmt.where(Workpaper.preparer_user_id == user.id)
    if workpaperId is not None:
        stmt = stmt.where(Workpaper.id == workpaperId)
    workpapers = db.execute(stmt.limit(5)).scalars().all()
    items: list[dict[str, Any]] = []
    total_stats: dict[str, int] = {}
    for workpaper in workpapers:
        path = Path(workpaper.file_path or "").expanduser()
        if not workpaper.file_path or not path.exists() or not path.is_file():
            continue
        text = extract_workpaper_text(path)
        redacted_text, stats = redact_text(text, terms)
        for key, count in stats.items():
            total_stats[key] = total_stats.get(key, 0) + count
        preview_text = redacted_text[:maxChars]
        items.append(
            {
                "workpaper_id": workpaper.id,
                "code": workpaper.code,
                "name": workpaper.name,
                "file_type": path.suffix.lower().lstrip("."),
                "redacted_text": preview_text,
                "char_count": len(preview_text),
            }
        )
    return {
        "project_id": project.id,
        "project_name": "[项目名称]",
        "redaction_enabled": True,
        "redaction_stats": total_stats,
        "items": items,
    }


@router.get("/api/review-runs/{run_id}/findings")
def list_review_findings(
    run_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    run = get_or_404(db, ReviewRun, run_id, "复核任务")
    project = get_or_404(db, Project, run.project_id, "项目")
    ensure_project_viewer(db, project, user)
    rows = db.execute(
        select(ReviewFinding).where(ReviewFinding.run_id == run_id).order_by(ReviewFinding.severity.desc(), ReviewFinding.id)
    ).scalars().all()
    return [_finding_payload(finding, run, project, db, user) for finding in rows]


@router.get("/api/review-findings")
def list_all_review_findings(
    projectId: Optional[int] = Query(default=None),
    status: Optional[str] = Query(default=None),
    severity: Optional[str] = Query(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    visible_ids = visible_project_ids(db, user)
    if not visible_ids:
        return []
    stmt = (
        _finding_query()
        .where(ReviewRun.project_id.in_(visible_ids))
        .order_by(ReviewFinding.id.desc())
    )
    if projectId is not None:
        project = get_or_404(db, Project, projectId, "项目")
        ensure_project_viewer(db, project, user)
        stmt = stmt.where(ReviewRun.project_id == projectId)
    if status:
        stmt = stmt.where(ReviewFinding.status == status)
    if severity:
        stmt = stmt.where(ReviewFinding.severity == severity)
    rows = db.execute(stmt).all()
    return [_finding_payload(finding, run, project, db, user) for finding, run, project in rows]


@router.get("/api/review-issue-catalog")
def list_review_issue_catalog(
    scope: str = Query(default=""),
    workpaper: str = Query(default=""),
    keyword: str = Query(default=""),
    limit: int = Query(default=2000, ge=1, le=3000),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    rows = filter_review_issue_catalog(scope=scope, workpaper=workpaper, keyword=keyword, limit=limit)
    return {"total": len(review_issue_catalog()), "count": len(rows), "items": rows}


@router.post("/api/review-findings", status_code=201)
def create_manual_review_finding(
    body: ReviewFindingIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    if body.run_id:
        run = get_or_404(db, ReviewRun, body.run_id, "复核任务")
        project = get_or_404(db, Project, run.project_id, "项目")
    else:
        if body.project_id is None:
            raise HTTPException(status_code=400, detail="人工问题记录必须选择项目")
        project = get_or_404(db, Project, body.project_id, "项目")
        run = db.execute(
            select(ReviewRun)
            .where(ReviewRun.project_id == project.id, ReviewRun.mode == "manual")
            .order_by(ReviewRun.id.desc())
        ).scalars().first()
        if run is None:
            run = ReviewRun(
                project_id=project.id,
                mode="manual",
                status="completed",
                started_at=datetime.utcnow(),
                finished_at=datetime.utcnow(),
                summary="人工复核问题记录",
            )
            db.add(run)
            db.flush()
    ensure_project_reviewer(project, user)
    workpaper: Workpaper | None = None
    review_step: ReviewStep | None = None
    if body.review_step_id is not None:
        review_step = get_or_404(db, ReviewStep, body.review_step_id, "复核步骤")
        workpaper = get_or_404(db, Workpaper, review_step.workpaper_id, "底稿")
        if body.workpaper_id is not None and body.workpaper_id != workpaper.id:
            raise HTTPException(status_code=400, detail="复核步骤与所选底稿不匹配")
    elif body.workpaper_id is not None:
        workpaper = get_or_404(db, Workpaper, body.workpaper_id, "底稿")
    if workpaper is not None and workpaper.project_id != project.id:
        raise HTTPException(status_code=400, detail="底稿不属于当前项目")
    if body.assignee_user_id is not None:
        assignee = get_or_404(db, User, body.assignee_user_id, "责任人")
        if not can_view_project(db, assignee, project):
            raise HTTPException(status_code=400, detail="责任人必须是本项目成员或项目负责人")
    latest_version = None
    if workpaper is not None:
        latest_version = db.execute(
            select(WorkpaperVersion)
            .where(WorkpaperVersion.workpaper_id == workpaper.id)
            .order_by(WorkpaperVersion.version_no.desc())
        ).scalars().first()
    finding_count = db.execute(
        select(func.count(ReviewFinding.id)).join(ReviewRun).where(ReviewRun.project_id == project.id)
    ).scalar_one()
    issue_no = body.issue_no or f"{'C22' if body.c22_related else 'NC22'}-{finding_count + 1:03d}"
    review_stage = _automatic_review_stage(db, project.id)
    finding = ReviewFinding(
        run_id=run.id,
        workpaper_id=workpaper.id if workpaper else None,
        workpaper_version_id=(review_step.workpaper_version_id if review_step else None) or (latest_version.id if latest_version else None),
        review_step_id=review_step.id if review_step else None,
        issue_no=issue_no,
        created_by_user_id=user.id,
        standard_index_code=body.standard_index_code or "",
        source=body.source or "manual",
        rule_code=body.rule_code or "MANUAL",
        audit_stage=body.audit_stage or "",
        finding_type=body.finding_type or "",
        issue_step=body.issue_step or "",
        issue_category=body.issue_category or "",
        severity=body.severity or "medium",
        c22_related=body.c22_related,
        workpaper_file=body.workpaper_file or "",
        location=body.location or "",
        target=body.target or (f"{workpaper.code} / {workpaper.name}" if workpaper else ""),
        issue=body.issue,
        evidence=body.evidence or "",
        recommendation=body.recommendation or "",
        project_reply=body.project_reply or "",
        resolution_confirmed=body.resolution_confirmed,
        review_stage=review_stage,
        field_lead=body.field_lead or "",
        project_reviewer=user.display_name if (body.source or "manual") == "manual" else (body.project_reviewer or ""),
        note="",
        status="resolved" if body.resolution_confirmed else ("responded" if body.project_reply else "open"),
        assignee_user_id=None,
        due_date=None,
        review_comment=body.review_comment or "",
    )
    db.add(finding)
    db.flush()
    db.add(
        ReviewFindingHistory(
            finding_id=finding.id,
            operator_user_id=user.id,
            old_status="",
            new_status=finding.status,
            comment="创建人工复核问题",
            change_summary=(
                f"创建问题；复核阶段：{finding.review_stage or '空'}；"
                f"对应底稿：{finding.target or finding.workpaper_file or '空'}；"
                f"复核问题：{finding.issue or '空'}"
            ),
        )
    )
    count = db.execute(select(func.count(ReviewFinding.id)).where(ReviewFinding.run_id == run.id)).scalar_one()
    run.status = "completed"
    run.finished_at = datetime.utcnow()
    run.summary = f"人工复核问题记录，累计 {count} 项"
    record_audit_log(db, None, "review_finding_created", user=user, target_type="review_finding", target_id=finding.id, project_id=project.id)
    db.commit()
    db.refresh(finding)
    db.refresh(run)
    return _finding_payload(finding, run, project, db, user)


@router.get("/api/projects/{project_id}/review-findings/import-template")
def download_review_finding_import_template(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> Response:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "问题清单"
    sheet.append(STANDARD_FINDING_HEADERS)
    sheet.append([
        "B-001", "计划阶段", "ITGC", "高", "否", "B60-2-1 IT复杂性判断表", "表1",
        "问题描述示例", "具体描述或证据", "E1", "内容填写不正确或不完整或错误", "", "否",
        "第一轮", "", "", "建议处理示例", "待补充/待整改",
    ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(STANDARD_FINDING_HEADERS))}2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    widths = [14, 14, 14, 10, 14, 34, 18, 42, 50, 16, 34, 30, 18, 18, 16, 18, 42, 18]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="ITAS_review_findings_template.xlsx"'},
    )


@router.post("/api/projects/{project_id}/review-findings/import")
async def import_review_findings(
    project_id: int,
    file: UploadFile = File(...),
    dry_run: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_reviewer(project, user)
    filename = file.filename or ""
    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="问题清单必须为 .xlsx 或 .xlsm 文件")
    rows, parse_errors = _parse_import_rows(await file.read())
    if parse_errors:
        raise HTTPException(status_code=400, detail="；".join(parse_errors))
    existing_runs = db.execute(
        select(ReviewRun).where(ReviewRun.project_id == project_id, ReviewRun.mode == "manual").order_by(ReviewRun.id.desc())
    ).scalars().all()
    run = existing_runs[0] if existing_runs else None
    existing: dict[str, ReviewFinding] = {}
    if run:
        for finding in db.execute(select(ReviewFinding).where(ReviewFinding.run_id.in_([item.id for item in existing_runs]))).scalars().all():
            if finding.issue_no:
                existing[finding.issue_no] = finding
    preview: list[dict[str, Any]] = []
    created = updated = skipped = 0
    if not dry_run and run is None:
        run = ReviewRun(
            project_id=project_id,
            mode="manual",
            status="completed",
            started_at=datetime.utcnow(),
            finished_at=datetime.utcnow(),
            summary="标准问题清单导入",
        )
        db.add(run)
        db.flush()
    for item in rows:
        old = existing.get(item["issue_no"])
        action = "update" if old else "create"
        preview.append({"issue_no": item["issue_no"], "action": action, "target": item["target"], "issue": item["issue"]})
        if dry_run:
            continue
        if old:
            previous_status = old.status or ""
            changes: list[str] = []
            for key in (
                "standard_index_code", "source", "rule_code", "audit_stage", "finding_type", "issue_step", "issue_category",
                "severity", "c22_related", "workpaper_file", "location", "target", "issue", "evidence", "recommendation",
                "project_reply", "resolution_confirmed", "review_stage", "field_lead", "project_reviewer", "status", "review_comment",
            ):
                new_value = item.get(key, getattr(old, key))
                if getattr(old, key) != new_value:
                    changes.append(_finding_change_text(key, getattr(old, key), new_value))
                setattr(old, key, new_value)
            if old.created_by_user_id is None:
                old.created_by_user_id = user.id
            db.add(
                ReviewFindingHistory(
                    finding_id=old.id,
                    operator_user_id=user.id,
                    old_status=previous_status,
                    new_status=old.status or "",
                    comment=f"导入文件：{filename}",
                    change_summary=f"标准问题清单更新：{'；'.join(changes) or '无字段变化'}",
                )
            )
            updated += 1
        else:
            finding = ReviewFinding(run_id=run.id, **{key: item.get(key, "") for key in (
                "issue_no", "standard_index_code", "source", "rule_code", "audit_stage", "finding_type", "issue_step", "issue_category",
                "severity", "c22_related", "workpaper_file", "location", "target", "issue", "evidence", "recommendation",
                "project_reply", "resolution_confirmed", "review_stage", "field_lead", "project_reviewer", "status", "review_comment",
            )})
            finding.created_by_user_id = user.id
            db.add(finding)
            db.flush()
            db.add(ReviewFindingHistory(
                finding_id=finding.id,
                operator_user_id=user.id,
                old_status="",
                new_status=finding.status,
                comment=f"导入文件：{filename}",
                change_summary=f"导入问题；对应底稿：{finding.target or finding.workpaper_file or '空'}；复核问题：{finding.issue or '空'}",
            ))
            created += 1
    if not dry_run:
        count = db.execute(select(func.count(ReviewFinding.id)).where(ReviewFinding.run_id == run.id)).scalar_one()
        run.status = "completed"
        run.finished_at = datetime.utcnow()
        run.summary = f"标准问题清单导入，新增 {created} 项、更新 {updated} 项，累计 {count} 项"
        record_audit_log(db, None, "review_findings_imported", user=user, target_type="review_run", target_id=run.id, project_id=project.id, details={"filename": filename, "created": created, "updated": updated})
        db.commit()
    return {"project_id": project_id, "project_name": project.name, "filename": filename, "dry_run": dry_run, "created": created, "updated": updated, "skipped": skipped, "rows": len(rows), "preview": preview}


@router.patch("/api/review-findings/{finding_id}")
def update_review_finding(
    finding_id: int,
    body: ReviewFindingPatchIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    row = db.execute(_finding_query().where(ReviewFinding.id == finding_id)).first()
    if row is None:
        raise HTTPException(status_code=404, detail="复核问题不存在")
    finding, run, project = row
    ensure_project_reviewer(project, user)
    payload = body.model_dump(exclude_unset=True)
    payload.pop("review_stage", None)
    if "status" in payload:
        raise HTTPException(status_code=400, detail="问题状态请通过分派、成员回复或复核决定操作变更")
    old_status = finding.status or ""
    changes: list[str] = []
    for key, value in payload.items():
        if value is not None:
            old_value = getattr(finding, key, None)
            if old_value != value:
                changes.append(_finding_change_text(key, old_value, value))
            setattr(finding, key, value)
    if run.mode == "manual":
        run.finished_at = datetime.utcnow()
    if changes:
        db.add(
            ReviewFindingHistory(
                finding_id=finding.id,
                operator_user_id=user.id,
                old_status=old_status,
                new_status=finding.status or "",
                comment="",
                change_summary="；".join(changes),
            )
        )
    db.commit()
    db.refresh(finding)
    return _finding_payload(finding, run, project, db, user)


@router.post("/api/review-findings/{finding_id}/responses", status_code=201)
def respond_to_review_finding(
    finding_id: int,
    body: ReviewFindingResponseIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    response_text = body.response_text.strip()
    if not response_text:
        raise HTTPException(status_code=400, detail="回复内容不能为空")
    row = db.execute(_finding_query().where(ReviewFinding.id == finding_id)).first()
    if row is None:
        raise HTTPException(status_code=404, detail="复核问题不存在")
    finding, run, project = row
    ensure_project_viewer(db, project, user)
    can_reply = bool(
        can_edit_project(user, project)
        or user.id == project.field_leader_user_id
        or is_project_member(db, project.id, user.id)
    )
    if not can_reply:
        raise HTTPException(status_code=403, detail="仅当前项目组成员可提交问题回复")
    if finding.status not in {"open", "assigned", "changes_requested"}:
        raise HTTPException(status_code=409, detail="当前问题状态不允许提交整改回复")
    next_round = int(
        db.execute(
            select(func.coalesce(func.max(ReviewFindingResponse.round_no), 0)).where(
                ReviewFindingResponse.finding_id == finding.id
            )
        ).scalar_one()
    ) + 1
    old_status = finding.status or ""
    response = ReviewFindingResponse(
        finding_id=finding.id,
        responder_user_id=user.id,
        round_no=next_round,
        response_text=response_text,
        attachment=body.attachment.strip(),
    )
    finding.status = "responded"
    db.add(response)
    db.add(
        ReviewFindingHistory(
            finding_id=finding.id,
            operator_user_id=user.id,
            old_status=old_status,
            new_status="responded",
            comment=response_text,
            change_summary=f"提交第 {next_round} 轮整改回复",
        )
    )
    record_audit_log(db, None, "review_finding_responded", user=user, target_type="review_finding", target_id=finding.id, project_id=project.id, details={"round_no": next_round})
    db.commit()
    db.refresh(finding)
    return _finding_payload(finding, run, project, db, user)


@router.post("/api/review-findings/{finding_id}/decision")
def decide_review_finding(
    finding_id: int,
    body: ReviewFindingDecisionIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    row = db.execute(_finding_query().where(ReviewFinding.id == finding_id)).first()
    if row is None:
        raise HTTPException(status_code=404, detail="复核问题不存在")
    finding, run, project = row
    ensure_project_reviewer(project, user)
    comment = body.comment.strip()
    if not comment:
        raise HTTPException(status_code=400, detail="复核决定说明不能为空")
    old_status = finding.status or ""
    transitions = {
        "changes_requested": ({"responded"}, "changes_requested", "整改回复被退回"),
        "closed": ({"responded"}, "closed", "复核确认关闭"),
        "reopened": ({"closed", "resolved"}, "changes_requested", "重新打开问题"),
    }
    allowed_from, new_status, summary = transitions[body.result]
    if old_status not in allowed_from:
        raise HTTPException(status_code=409, detail=f"问题状态 {old_status or '空'} 不允许执行该复核决定")
    resolved_version = None
    if body.result == "closed" and run.mode in {"auto", "ai"} and finding.workpaper_id and finding.workpaper_version_id:
        baseline_version = db.get(WorkpaperVersion, finding.workpaper_version_id)
        resolved_version = db.execute(
            select(WorkpaperVersion)
            .where(WorkpaperVersion.workpaper_id == finding.workpaper_id)
            .order_by(WorkpaperVersion.version_no.desc())
        ).scalars().first()
        if (
            baseline_version is None
            or resolved_version is None
            or resolved_version.version_no <= baseline_version.version_no
        ):
            raise HTTPException(status_code=409, detail="请先上传整改后的新版本底稿，再确认关闭自动复核问题")
    finding.status = new_status
    finding.review_comment = comment
    if resolved_version is not None:
        finding.resolved_workpaper_version_id = resolved_version.id
        summary = f"{summary}；整改版本 V{resolved_version.version_no}"
    db.add(
        ReviewFindingHistory(
            finding_id=finding.id,
            operator_user_id=user.id,
            old_status=old_status,
            new_status=new_status,
            comment=comment,
            change_summary=summary,
        )
    )
    record_audit_log(db, None, "review_finding_decided", user=user, target_type="review_finding", target_id=finding.id, project_id=project.id, details={"result": body.result, "new_status": new_status})
    db.commit()
    db.refresh(finding)
    return _finding_payload(finding, run, project, db, user)


@router.get("/api/review-findings/{finding_id}/history")
def list_review_finding_history(
    finding_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    row = db.execute(_finding_query().where(ReviewFinding.id == finding_id)).first()
    if row is None:
        raise HTTPException(status_code=404, detail="复核问题不存在")
    _finding, _run, project = row
    ensure_project_viewer(db, project, user)
    if project.id not in supervised_project_ids(db, user) and _finding.assignee_user_id != user.id:
        raise HTTPException(status_code=403, detail="仅可查看分派给本人的问题记录")
    histories = db.execute(
        select(ReviewFindingHistory)
        .where(ReviewFindingHistory.finding_id == finding_id)
        .order_by(ReviewFindingHistory.id.desc())
    ).scalars().all()
    return [_history_payload(item) for item in histories]


@router.get("/api/review-dashboard")
def review_dashboard(
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    visible_ids = visible_project_ids(db, user)
    supervised_ids = supervised_project_ids(db, user)
    own_only_ids = visible_ids - supervised_ids
    rows = db.execute(
        _finding_query()
        .where(
            or_(
                ReviewRun.project_id.in_(supervised_ids),
                and_(ReviewRun.project_id.in_(own_only_ids), ReviewFinding.assignee_user_id == user.id),
            )
        )
        .order_by(ReviewFinding.id.desc())
    ).all() if visible_ids else []
    by_project: dict[int, dict[str, Any]] = {}
    by_rule: dict[str, int] = {}
    by_severity: dict[str, int] = {}
    by_status: dict[str, int] = {}
    recent: list[dict[str, Any]] = []
    overdue: list[dict[str, Any]] = []
    manual_count = 0
    open_count = 0
    for finding, run, project in rows:
        project_bucket = by_project.setdefault(
            project.id,
            {
                "project_id": project.id,
                "project_name": project.name,
                "entity_name": project.entity_name,
                "total": 0,
                "open": 0,
                "high": 0,
                "medium": 0,
                "low": 0,
                "manual": 0,
            },
        )
        status = finding.status or "open"
        severity = finding.severity or "medium"
        project_bucket["total"] += 1
        project_bucket[severity] = project_bucket.get(severity, 0) + 1
        if _is_open_finding(status):
            project_bucket["open"] += 1
            open_count += 1
        if run.mode == "manual":
            project_bucket["manual"] += 1
            manual_count += 1
        by_rule[finding.rule_code or "未分类"] = by_rule.get(finding.rule_code or "未分类", 0) + 1
        by_severity[severity] = by_severity.get(severity, 0) + 1
        by_status[status] = by_status.get(status, 0) + 1
        payload = _finding_payload(finding, run, project, db, user)
        if payload.get("is_overdue"):
            overdue.append(payload)
        if len(recent) < 20:
            recent.append(payload)
    material_overdue: list[dict[str, Any]] = []
    material_rows = db.execute(
        select(DocumentRequest, Project)
        .join(Project, DocumentRequest.project_id == Project.id)
        .where(DocumentRequest.due_date.is_not(None), Project.id.in_(visible_ids))
        .order_by(DocumentRequest.due_date, DocumentRequest.id)
    ).all()
    for request_item, project in material_rows:
        payload = obj_dict(request_item)
        payload.update(overdue_payload(request_item.status, request_item.due_date))
        if not payload.get("is_overdue"):
            continue
        payload["project_id"] = project.id
        payload["project_name"] = project.name
        payload["entity_name"] = project.entity_name
        material_overdue.append(payload)
        bucket = by_project.setdefault(
            project.id,
            {
                "project_id": project.id,
                "project_name": project.name,
                "entity_name": project.entity_name,
                "total": 0,
                "open": 0,
                "high": 0,
                "medium": 0,
                "low": 0,
                "manual": 0,
            },
        )
        bucket["material_overdue_count"] = bucket.get("material_overdue_count", 0) + 1
    step_overdue: list[dict[str, Any]] = []
    step_rows = db.execute(
        select(ReviewStep, Workpaper, Project)
        .join(Workpaper, ReviewStep.workpaper_id == Workpaper.id)
        .join(Project, Workpaper.project_id == Project.id)
        .where(
            ReviewStep.due_date.is_not(None),
            or_(
                Project.id.in_(supervised_ids),
                and_(Project.id.in_(own_only_ids), Workpaper.preparer_user_id == user.id),
            ),
        )
        .order_by(ReviewStep.due_date, ReviewStep.id)
    ).all()
    for step, workpaper, project in step_rows:
        payload = _review_step_payload(step, workpaper, project, user)
        if not payload.get("is_overdue"):
            continue
        step_overdue.append(payload)
        bucket = by_project.setdefault(
            project.id,
            {
                "project_id": project.id,
                "project_name": project.name,
                "entity_name": project.entity_name,
                "total": 0,
                "open": 0,
                "high": 0,
                "medium": 0,
                "low": 0,
                "manual": 0,
            },
        )
        bucket["step_overdue_count"] = bucket.get("step_overdue_count", 0) + 1
    by_project_rows = sorted(by_project.values(), key=lambda item: (item["open"], item["total"]), reverse=True)
    return {
        "total": len(rows),
        "open": open_count,
        "manual": manual_count,
        "auto": len(rows) - manual_count,
        "overdue": len(overdue),
        "overdue_findings": sorted(overdue, key=lambda item: item.get("overdue_days", 0), reverse=True)[:20],
        "material_overdue": sorted(material_overdue, key=lambda item: item.get("overdue_days", 0), reverse=True)[:20],
        "step_overdue": sorted(step_overdue, key=lambda item: item.get("overdue_days", 0), reverse=True)[:20],
        "projects_with_findings": len(by_project_rows),
        "by_project": by_project_rows,
        "by_rule": [
            {"rule_code": code, "count": count}
            for code, count in sorted(by_rule.items(), key=lambda item: item[1], reverse=True)[:30]
        ],
        "by_severity": [{"severity": key, "count": value} for key, value in sorted(by_severity.items())],
        "by_status": [{"status": key, "count": value} for key, value in sorted(by_status.items())],
        "recent": recent,
    }


def _submit_workpaper_for_review(
    db: Session,
    wp: Workpaper,
    project: Project,
    user: User,
    next_reviewer_user_id: int | None = None,
) -> dict[str, Any]:
    ensure_document_uploader(db, project, user)
    if wp.preparer_user_id and wp.preparer_user_id != user.id and not can_edit_project(user, project) and not is_admin(user):
        raise HTTPException(status_code=403, detail="仅底稿编制人或项目负责人可提交该底稿")
    existing = db.execute(
        select(ReviewStep).where(ReviewStep.workpaper_id == wp.id).order_by(ReviewStep.sequence_no)
    ).scalars().all()
    latest_version = db.execute(
        select(WorkpaperVersion)
        .where(WorkpaperVersion.workpaper_id == wp.id)
        .order_by(WorkpaperVersion.version_no.desc())
    ).scalars().first()
    if latest_version is None:
        raise HTTPException(status_code=409, detail="请先上传底稿文件，系统需要保留可追溯的版本后才能提交复核")
    if not existing:
        if wp.status != "draft":
            raise HTTPException(status_code=409, detail=f"底稿当前为 {wp.status}，不能发起首次复核")
        now = datetime.utcnow()
        assignments = _assignments_from_selected_reviewer(project, next_reviewer_user_id)
        for index, (role_code, reviewer_user_id) in enumerate(assignments, start=1):
            step = ReviewStep(
                workpaper_id=wp.id,
                sequence_no=index,
                reviewer_role_code=role_code,
                reviewer_user_id=reviewer_user_id,
                workpaper_version_id=latest_version.id,
                round_no=1,
                status="pending" if index == 1 else "waiting",
            )
            if index == 1:
                _enter_review_step(db, step, now)
            db.add(step)
            db.flush()
            _add_step_history(
                db,
                step,
                user,
                "submitted" if index == 1 else "queued",
                "",
                step.status,
                "底稿提交复核",
            )
    elif wp.status == "returned":
        if next_reviewer_user_id is not None:
            raise HTTPException(status_code=409, detail="整改后重提将沿用原复核链，无需重新选择复核人")
        rejected = next((step for step in existing if step.status == "rejected"), None)
        if rejected is None:
            raise HTTPException(status_code=409, detail="退回底稿缺少被退回复核步骤，无法重新提交")
        _validate_resubmission_version(rejected, latest_version)
        pending_findings = db.execute(
            select(ReviewFinding).where(
                ReviewFinding.review_step_id == rejected.id,
                ReviewFinding.status.not_in({"responded", "closed", "resolved"}),
            )
        ).scalars().all()
        if pending_findings:
            raise HTTPException(status_code=409, detail="请先对本轮复核问题提交整改回复，再重新提交底稿")
        now = datetime.utcnow()
        for step in existing:
            if step.sequence_no < rejected.sequence_no:
                continue
            old_status = step.status
            step.round_no += 1
            step.workpaper_version_id = latest_version.id
            step.status = "pending" if step.id == rejected.id else "waiting"
            step.comment = ""
            step.reviewed_at = None
            step.entered_at = None
            step.due_date = None
            if step.id == rejected.id:
                _enter_review_step(db, step, now)
            _add_step_history(db, step, user, "resubmitted", old_status, step.status, "底稿整改后重新提交")
    else:
        raise HTTPException(status_code=409, detail=f"底稿当前为 {wp.status}，不能重复提交复核")
    wp.status = "submitted"
    record_audit_log(
        db,
        None,
        "workpaper_submitted",
        user=user,
        target_type="workpaper",
        target_id=wp.id,
        project_id=project.id,
        details={"resubmission": bool(existing), "next_reviewer_user_id": next_reviewer_user_id},
    )
    return {"workpaper_id": wp.id, "status": "submitted"}


@router.get("/api/projects/{project_id}/reviewer-options")
def list_project_reviewer_options(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    role_names = {
        "project_manager": "项目经理",
        "responsible_manager": "项目负责经理",
        "director": "总监",
        "partner": "合伙人",
        "quality": "质控",
    }
    options: list[dict[str, Any]] = []
    seen_user_ids: set[int] = set()
    for role_code, reviewer_user_id in _project_review_assignments(project):
        if reviewer_user_id in seen_user_ids:
            continue
        reviewer = db.get(User, reviewer_user_id)
        if reviewer is None or reviewer.status != "active":
            continue
        seen_user_ids.add(reviewer_user_id)
        options.append({
            "user_id": reviewer.id,
            "display_name": reviewer.display_name,
            "role_code": role_code,
            "role_name": role_names.get(role_code, role_code),
        })
    return options


@router.post("/api/workpapers/{workpaper_id}/submit")
def submit_workpaper(
    workpaper_id: int,
    body: Optional[WorkpaperSubmitIn] = Body(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    wp = get_or_404(db, Workpaper, workpaper_id, "底稿")
    project = get_or_404(db, Project, wp.project_id, "项目")
    result = _submit_workpaper_for_review(db, wp, project, user, body.next_reviewer_user_id if body else None)
    db.commit()
    return result


@router.post("/api/workpapers/batch-submit")
def batch_submit_workpapers(
    body: BatchWorkpaperSubmitIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    workpaper_ids = list(dict.fromkeys(body.workpaper_ids))
    workpapers = [get_or_404(db, Workpaper, workpaper_id, "底稿") for workpaper_id in workpaper_ids]
    project_ids = {workpaper.project_id for workpaper in workpapers}
    if len(project_ids) != 1:
        raise HTTPException(status_code=400, detail="批量提交的底稿必须属于同一项目")
    project = get_or_404(db, Project, workpapers[0].project_id, "项目")
    try:
        results = [
            _submit_workpaper_for_review(db, workpaper, project, user, body.next_reviewer_user_id)
            for workpaper in workpapers
        ]
        db.commit()
    except Exception:
        db.rollback()
        raise
    return {"count": len(results), "workpapers": results, "status": "submitted"}


@router.get("/api/workpapers/{workpaper_id}/review-steps")
def list_review_steps(
    workpaper_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    workpaper = get_or_404(db, Workpaper, workpaper_id, "底稿")
    project = get_or_404(db, Project, workpaper.project_id, "项目")
    ensure_workpaper_viewer(db, project, workpaper, user)
    rows = db.execute(
        select(ReviewStep).where(ReviewStep.workpaper_id == workpaper_id).order_by(ReviewStep.sequence_no)
    ).scalars().all()
    return [_review_step_payload(row, workpaper, project, user) for row in rows]


@router.get("/api/projects/{project_id}/review-steps")
def list_project_review_steps(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    statement = (
        select(ReviewStep, Workpaper)
        .join(Workpaper, ReviewStep.workpaper_id == Workpaper.id)
        .where(Workpaper.project_id == project.id)
        .order_by(Workpaper.code, ReviewStep.sequence_no)
    )
    if project.id not in supervised_project_ids(db, user):
        statement = statement.where(Workpaper.preparer_user_id == user.id)
    rows = db.execute(statement).all()
    status_order = {"pending": 0, "rejected": 1, "waiting": 2, "approved": 3}
    payload = [_review_step_payload(step, workpaper, project, user) for step, workpaper in rows]
    return sorted(
        payload,
        key=lambda item: (
            0 if item.get("can_decide") else 1,
            status_order.get(str(item.get("status") or ""), 9),
            bool(not item.get("is_overdue")),
            str(item.get("workpaper_code") or ""),
            int(item.get("sequence_no") or 0),
        ),
    )


@router.post("/api/review-steps/{step_id}/approve")
def approve_review_step(
    step_id: int,
    body: ReviewDecisionIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    step = get_or_404(db, ReviewStep, step_id, "复核步骤")
    wp = get_or_404(db, Workpaper, step.workpaper_id, "底稿")
    project = get_or_404(db, Project, wp.project_id, "项目")
    ensure_project_viewer(db, project, user)
    _ensure_step_actor(step, user, project)
    linked_findings = db.execute(
        select(ReviewFinding).where(ReviewFinding.review_step_id == step.id)
    ).scalars().all()
    awaiting_member = [
        finding for finding in linked_findings
        if finding.status not in {"responded", "closed", "resolved"}
    ]
    if awaiting_member:
        raise HTTPException(status_code=409, detail="本复核步骤仍有问题尚未完成成员回复，不能通过")
    old_status = step.status
    step.status = "approved"
    step.comment = body.comment.strip()
    step.reviewed_at = datetime.utcnow()
    _add_step_history(db, step, user, "approved", old_status, "approved", step.comment)
    closed_finding_ids: list[int] = []
    for finding in linked_findings:
        if finding.status != "responded":
            continue
        finding.status = "closed"
        finding.resolved_workpaper_version_id = step.workpaper_version_id
        finding.review_comment = step.comment or "底稿复核通过，整改确认关闭"
        closed_finding_ids.append(finding.id)
        db.add(
            ReviewFindingHistory(
                finding_id=finding.id,
                operator_user_id=user.id,
                old_status="responded",
                new_status="closed",
                comment=finding.review_comment,
                change_summary=f"第 {step.round_no} 轮底稿复核通过，自动确认整改关闭",
            )
        )
    next_step = db.execute(
        select(ReviewStep)
        .where(ReviewStep.workpaper_id == step.workpaper_id, ReviewStep.sequence_no > step.sequence_no)
        .order_by(ReviewStep.sequence_no)
    ).scalars().first()
    if next_step:
        next_old_status = next_step.status
        next_step.status = "pending"
        _enter_review_step(db, next_step)
        _add_step_history(db, next_step, user, "activated", next_old_status, "pending", "上一级复核已通过")
        wp.status = "in_review"
    else:
        wp.status = "approved"
    record_audit_log(db, None, "review_step_approved", user=user, target_type="review_step", target_id=step.id, project_id=project.id, details={"workpaper_id": wp.id, "round_no": step.round_no, "closed_finding_ids": closed_finding_ids})
    db.commit()
    return {"status": wp.status}


@router.post("/api/review-steps/{step_id}/reject")
def reject_review_step(
    step_id: int,
    body: ReviewDecisionIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    step = get_or_404(db, ReviewStep, step_id, "复核步骤")
    wp = get_or_404(db, Workpaper, step.workpaper_id, "底稿")
    project = get_or_404(db, Project, wp.project_id, "项目")
    ensure_project_viewer(db, project, user)
    _ensure_step_actor(step, user, project)
    comment = body.comment.strip()
    if not comment:
        raise HTTPException(status_code=400, detail="退回底稿必须填写复核意见")
    old_status = step.status
    step.status = "rejected"
    step.comment = comment
    step.reviewed_at = datetime.utcnow()
    _add_step_history(db, step, user, "rejected", old_status, "rejected", comment)
    wp.status = "returned"
    run = db.execute(
        select(ReviewRun)
        .where(ReviewRun.project_id == project.id, ReviewRun.mode == "manual")
        .order_by(ReviewRun.id.desc())
    ).scalars().first()
    if run is None:
        run = ReviewRun(
            project_id=project.id,
            mode="manual",
            status="completed",
            started_at=datetime.utcnow(),
            finished_at=datetime.utcnow(),
            summary="人工复核问题记录",
        )
        db.add(run)
        db.flush()
    version = db.get(WorkpaperVersion, step.workpaper_version_id) if step.workpaper_version_id else None
    finding = ReviewFinding(
        run_id=run.id,
        workpaper_id=wp.id,
        workpaper_version_id=step.workpaper_version_id,
        review_step_id=step.id,
        issue_no=f"WP-{wp.id}-S{step.id}-R{step.round_no}",
        source="review_step",
        rule_code="WORKPAPER_REVIEW",
        severity="medium",
        workpaper_file=Path(wp.file_path or "").name,
        location=step.reviewer_role_code,
        target=f"{wp.code} / {wp.name}",
        issue=comment,
        evidence=f"被复核版本：V{version.version_no}" if version else "被复核版本：历史版本",
        recommendation="请根据复核意见修订底稿，上传新版本并提交整改回复。",
        status="assigned" if wp.preparer_user_id else "open",
        assignee_user_id=wp.preparer_user_id,
        due_date=days_after(datetime.utcnow().date(), 5),
        review_comment=comment,
    )
    db.add(finding)
    db.flush()
    db.add(
        ReviewFindingHistory(
            finding_id=finding.id,
            operator_user_id=user.id,
            old_status="",
            new_status=finding.status,
            comment=comment,
            change_summary=f"第 {step.round_no} 轮底稿复核退回并生成整改问题",
        )
    )
    count = db.execute(select(func.count(ReviewFinding.id)).where(ReviewFinding.run_id == run.id)).scalar_one()
    run.status = "completed"
    run.finished_at = datetime.utcnow()
    run.summary = f"人工复核问题记录，累计 {count} 项"
    record_audit_log(db, None, "review_step_rejected", user=user, target_type="review_step", target_id=step.id, project_id=project.id, details={"workpaper_id": wp.id, "round_no": step.round_no, "finding_id": finding.id})
    db.commit()
    return {"status": "returned", "finding_id": finding.id}
