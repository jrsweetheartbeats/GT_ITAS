from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
import json
from pathlib import Path
import re
import secrets
import shutil
import subprocess
import sys
from typing import Any, Optional, Sequence

from fastapi import APIRouter, Body, Depends, File, Header, HTTPException, Query, Response, UploadFile
from fastapi.responses import FileResponse
from docx import Document
from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..core.config import BASE_DIR, WORKSPACE_ROOT
from ..core.db import get_db, safe_database_label
from ..core.security import (
    DEFAULT_MODULE_ORDER,
    DEFAULT_PASSWORD_POLICY,
    can_edit_project,
    can_view_project,
    can_upload_documents,
    current_user,
    default_audit_scope,
    ensure_document_uploader,
    ensure_feature_permission,
    ensure_project_editor,
    ensure_project_viewer,
    get_setting,
    hash_password,
    is_admin,
    normalize_module_order,
    require_admin,
    set_setting,
    validate_password_policy,
    visible_project_ids,
    supervised_project_ids,
    verify_password,
)
from ..core.utils import apply_patch_to_model, get_or_404, list_dict, obj_dict
from ..models import (
    Attachment,
    AutomationRule,
    AutofillPlanItem,
    AutofillRun,
    Client,
    ClientITContact,
    DocumentRequest,
    EnterpriseContact,
    LoginSession,
    ImsContact,
    Project,
    ProjectMember,
    ProjectMemberClaim,
    ReviewFinding,
    ReviewRun,
    ReviewStep,
    Role,
    Task,
    User,
    Workpaper,
)
from ..schemas import (
    AttachmentIn,
    AttachmentReferenceScanIn,
    AttachmentScanIn,
    AutomationRuleIn,
    AutofillPlanIn,
    ClientITContactIn,
    ClientIn,
    ContactIn,
    DocumentRequestUploadIn,
    InitFromPriorIn,
    LoginIn,
    HomeProjectMemberBatchIn,
    HomeProjectClaimApprovalIn,
    HomeProjectUpdateIn,
    MemberIn,
    PasswordPolicyIn,
    ProjectIn,
    ReviewDecisionIn,
    ReviewRunIn,
    RoleIn,
    TaskIn,
    UserIn,
    WorkpaperIn,
    WorkpaperTreeNodeOut,
)


from ..services.attachments import (
    append_reference,
    candidate_reference_tokens,
    extract_workpaper_text,
    file_type_for_path,
    guess_workpaper_for_file,
    likely_attachment_reference,
    next_attachment_index,
    safe_project_code,
    sanitize_workpaper_code,
    scan_project_attachment_files,
)
from ..services.materials import c22_document_requests_from_rules, save_upload_files
from ..services.project_scope import infer_audit_scope_from_workpapers
from ..services.projects import (
    copy_workpaper_file,
    default_project_root,
    initialize_project_workspace,
    public_workspace_report,
    replace_audit_year,
    rollback_project_workspace,
)
from ..services.review import add_finding, run_external_rules, run_internal_review
from ..services.delivery_deadline import delivery_info
from ..services.workpaper_reader import build_workpaper_tree
from ..services.audit_log import record_audit_log


router = APIRouter()


PROJECT_REVIEWER_ROLE_REQUIREMENTS = {
    "manager_user_id": ({"manager", "senior_manager", "admin"}, "项目负责经理"),
    "quality_reviewer_user_id": ({"quality", "admin"}, "质控复核人"),
    "partner_user_id": ({"partner", "admin"}, "合伙人复核人"),
}


def _default_user_id(db: Session, display_name: str) -> int | None:
    user = db.execute(select(User).where(User.display_name == display_name, User.status == "active")).scalar_one_or_none()
    return user.id if user else None


def _quality_reviewer_ids(project: Project) -> list[int]:
    try:
        values = json.loads(project.quality_reviewer_user_ids_json or "[]")
    except (TypeError, ValueError, json.JSONDecodeError):
        values = []
    values = values if isinstance(values, list) else []
    if project.quality_reviewer_user_id:
        values = [project.quality_reviewer_user_id, *values]
    return list(dict.fromkeys(int(value) for value in values if str(value).isdigit() and int(value) > 0))


def _set_quality_reviewer_ids(project: Project, user_ids: list[int]) -> None:
    ids = list(dict.fromkeys(int(user_id) for user_id in user_ids if int(user_id) > 0))
    project.quality_reviewer_user_ids_json = json.dumps(ids)
    project.quality_reviewer_user_id = ids[0] if ids else None


def _validate_project_reviewers(db: Session, payload: dict[str, Any]) -> None:
    leader_user_id = payload.get("project_leader_user_id")
    if leader_user_id:
        leader = get_or_404(db, User, int(leader_user_id), "项目负责人")
        if not leader.role or leader.role.rank < 780:
            raise HTTPException(status_code=400, detail="项目负责人仅可选择高级经理及以上人员")
    for field, (allowed_roles, label) in PROJECT_REVIEWER_ROLE_REQUIREMENTS.items():
        user_id = payload.get(field)
        if not user_id:
            continue
        reviewer = get_or_404(db, User, int(user_id), label)
        role_code = reviewer.role.code if reviewer.role else ""
        if role_code not in allowed_roles:
            raise HTTPException(status_code=400, detail=f"{label}所选人员角色不匹配")
    for user_id in payload.get("quality_reviewer_user_ids", []):
        reviewer = get_or_404(db, User, int(user_id), "质控复核人")
        if not reviewer.role or reviewer.role.code not in {"quality", "manager", "senior_manager", "director", "partner", "admin"}:
            raise HTTPException(status_code=400, detail="质控复核人所选人员角色不匹配")


def _team_maps(db: Session, project_ids: set[int], user: User) -> tuple[dict[int, list[dict[str, Any]]], dict[int, list[dict[str, Any]]]]:
    members_by_project: dict[int, list[dict[str, Any]]] = defaultdict(list)
    contacts_by_project: dict[int, list[dict[str, Any]]] = defaultdict(list)
    if not project_ids:
        return members_by_project, contacts_by_project
    supervised = supervised_project_ids(db, user)
    member_rows = db.execute(
        select(ProjectMember).where(ProjectMember.project_id.in_(project_ids)).order_by(ProjectMember.id)
    ).scalars().all()
    for row in member_rows:
        if row.project_id not in supervised and row.user_id != user.id:
            continue
        members_by_project[int(row.project_id)].append({
            "id": row.id,
            "user_id": row.user_id,
            "display_name": row.user.display_name if row.user else "",
            "username": row.user.username if row.user else "",
            "role_name": row.user.role.name if row.user and row.user.role else "",
            "role_on_project": row.role_on_project or "",
            "module": row.module or "",
            "workload": row.workload or "",
        })
    contact_rows = db.execute(
        select(EnterpriseContact).where(EnterpriseContact.project_id.in_(project_ids)).order_by(EnterpriseContact.id)
    ).scalars().all()
    for row in contact_rows:
        contacts_by_project[int(row.project_id)].append(obj_dict(row))
    return members_by_project, contacts_by_project


@router.get("/api/projects")
def list_projects(
    status: Optional[str] = Query(default=None),
    activeOnly: bool = Query(default=False),
    includeDelivery: bool = Query(default=False),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    visible_ids = visible_project_ids(db, user)
    if not visible_ids:
        return []
    stmt = select(Project).where(Project.id.in_(visible_ids)).order_by(Project.id.desc())
    # 项目选择和搜索不再按“活跃状态”隐藏项目。
    if status:
        stmt = stmt.where(Project.status == status)
    rows = db.execute(stmt).scalars().all()
    project_ids = {row.id for row in rows}
    member_counts = dict(
        db.execute(select(ProjectMember.project_id, func.count(ProjectMember.id)).group_by(ProjectMember.project_id)).all()
    )
    workpaper_counts = dict(
        db.execute(select(Workpaper.project_id, func.count(Workpaper.id)).group_by(Workpaper.project_id)).all()
    )
    attachment_counts = dict(
        db.execute(select(Attachment.project_id, func.count(Attachment.id)).group_by(Attachment.project_id)).all()
    )
    task_counts = dict(db.execute(select(Task.project_id, func.count(Task.id)).group_by(Task.project_id)).all())
    members_by_project, contacts_by_project = _team_maps(db, project_ids, user)
    delivery_workpapers: dict[int, list[Workpaper]] = {}
    if includeDelivery and project_ids:
        delivery_rows = db.execute(
            select(Workpaper).where(
                Workpaper.project_id.in_(project_ids),
                or_(Workpaper.code.like("B60-2-3%"), Workpaper.name.like("%计划备忘录%")),
            )
        ).scalars().all()
        for workpaper in delivery_rows:
            delivery_workpapers.setdefault(workpaper.project_id, []).append(workpaper)
    payload = []
    for row in rows:
        data = obj_dict(row)
        data["quality_reviewer_user_ids"] = _quality_reviewer_ids(row)
        quality_users = db.execute(select(User).where(User.id.in_(data["quality_reviewer_user_ids"]))).scalars().all() if data["quality_reviewer_user_ids"] else []
        by_id = {item.id: item.display_name for item in quality_users}
        data["quality_reviewer_names"] = [by_id[user_id] for user_id in data["quality_reviewer_user_ids"] if user_id in by_id]
        delivery = delivery_info(row, delivery_workpapers.get(row.id, [])) if includeDelivery else {}
        data["manager_name"] = row.manager.display_name if row.manager else ""
        data["project_leader_name"] = row.project_leader.display_name if row.project_leader else ""
        data["quality_reviewer_name"] = row.quality_reviewer.display_name if row.quality_reviewer else ""
        data["field_leader_name"] = row.field_leader.display_name if row.field_leader else ""
        data["director_name"] = row.director.display_name if row.director else ""
        data["partner_name"] = row.partner.display_name if row.partner else ""
        data["client_name"] = row.client.entity_name if row.client else ""
        data["member_count"] = int(member_counts.get(row.id, 0))
        data["workpaper_count"] = int(workpaper_counts.get(row.id, 0))
        data["attachment_count"] = int(attachment_counts.get(row.id, 0))
        data["task_count"] = int(task_counts.get(row.id, 0))
        data["members"] = members_by_project.get(row.id, [])
        data["contacts"] = contacts_by_project.get(row.id, [])
        data["can_edit"] = can_edit_project(user, row)
        if includeDelivery:
            data["due_days"] = delivery.get("due_days")
            data["delivery_date"] = delivery.get("delivery_date").isoformat() if delivery.get("delivery_date") else ""
            data["project_exit_date"] = delivery.get("project_exit_date").isoformat() if delivery.get("project_exit_date") else ""
            data["delivery_source"] = delivery.get("delivery_source", "")
            data["delivery_source_type"] = delivery.get("delivery_source_type", "")
            data["delivery_source_workpaper_id"] = delivery.get("delivery_source_workpaper_id")
            data["delivery_source_workpaper_code"] = delivery.get("delivery_source_workpaper_code", "")
            data["delivery_source_workpaper_name"] = delivery.get("delivery_source_workpaper_name", "")
        payload.append(data)
    return payload


@router.get("/api/home/projects")
def list_home_projects(db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    """Return firm-wide project-directory fields for the homepage only."""
    rows = db.execute(select(Project)).scalars().all()
    rows.sort(key=lambda row: (row.start_date or row.audit_scope_start or date.min, row.id), reverse=True)
    return [_home_project_payload(row) for row in rows]


def _home_systems_from_it_overview(db: Session, project_id: int) -> tuple[list[str], str]:
    """Read system names from column B of the B22A-4-1 IT overview workbook."""
    workpapers = db.execute(
        select(Workpaper)
        .where(Workpaper.project_id == project_id, Workpaper.code == "B22A-4-1")
        .order_by(Workpaper.id.desc())
    ).scalars().all()
    if not workpapers:
        return [], "未登记 B22A-4-1 IT概要底稿。"

    errors: list[str] = []
    for workpaper in workpapers:
        file_path = Path(workpaper.file_path or "")
        if file_path.suffix.lower() not in {".xlsx", ".xlsm"}:
            errors.append("B22A-4-1 不是可读取的 Excel 文件")
            continue
        if not file_path.is_file():
            errors.append("B22A-4-1 底稿文件不存在")
            continue
        try:
            workbook = load_workbook(file_path, read_only=True, data_only=True, keep_vba=file_path.suffix.lower() == ".xlsm")
        except Exception:
            errors.append("B22A-4-1 底稿无法读取")
            continue
        try:
            systems: list[str] = []
            seen: set[str] = set()
            for worksheet in workbook.worksheets:
                for serial, system_name in worksheet.iter_rows(min_row=6, max_col=2, values_only=True):
                    serial_text = str(serial or "").strip()
                    if not serial_text or not serial_text.replace(".", "", 1).isdigit():
                        continue
                    name = str(system_name or "").strip()
                    key = name.casefold()
                    if name and key not in seen:
                        seen.add(key)
                        systems.append(name)
            if systems:
                return systems, ""
            errors.append("B22A-4-1 的 B 列未识别到系统名称")
        finally:
            workbook.close()
    return [], "；".join(dict.fromkeys(errors)) or "未识别到公司系统清单。"


def _ims_contact_match_filters(users: Sequence[User]):
    emails = sorted({str(item.email or "").strip().casefold() for item in users if str(item.email or "").strip()})
    names = sorted({str(item.display_name or "").strip() for item in users if str(item.display_name or "").strip()})
    filters = []
    if emails:
        filters.append(func.lower(ImsContact.email).in_(emails))
    if names:
        filters.append(ImsContact.lastname.in_(names))
    return filters


def _home_staff_directory(
    db: Session,
    users: Sequence[User] | None = None,
) -> tuple[list[dict[str, Any]], dict[int, dict[str, Any]]]:
    """Active ITAS users enriched with office and job fields from the IMS sync."""
    if users is None:
        users = db.execute(select(User).where(User.status == "active").order_by(User.display_name, User.id)).scalars().all()
    filters = _ims_contact_match_filters(users)
    contacts = db.execute(select(ImsContact).where(or_(*filters))).scalars().all() if filters else []
    contact_by_email = {str(item.email or "").strip().casefold(): item for item in contacts if str(item.email or "").strip()}
    contacts_by_name: dict[str, list[ImsContact]] = defaultdict(list)
    for item in contacts:
        name = str(item.lastname or "").strip()
        if name:
            contacts_by_name[name].append(item)
    rows: list[dict[str, Any]] = []
    by_user_id: dict[int, dict[str, Any]] = {}
    for item in users:
        email_key = str(item.email or "").strip().casefold()
        candidates = contacts_by_name.get(str(item.display_name or "").strip(), [])
        contact = contact_by_email.get(email_key) or (candidates[0] if len(candidates) == 1 else None)
        row = {
            "id": item.id,
            "name": item.display_name or item.username,
            "username": item.username,
            "role_code": item.role.code if item.role else "",
            "role_name": item.role.name if item.role else "",
            "office": contact.subcompany if contact else "",
            "department": contact.department if contact else "",
            "job_title": contact.job_title if contact else "",
            "workcode": contact.workcode if contact else "",
        }
        rows.append(row)
        by_user_id[item.id] = row
    return rows, by_user_id


def _home_project_payload(row: Project) -> dict[str, Any]:
    return {
        "id": row.id,
        "name": row.name,
        "code": row.code,
        "oa_project_no": row.oa_project_no,
        "ims_project_no": row.ims_project_no,
        "entity_name": row.entity_name,
        "client_name": row.client.entity_name if row.client else "",
        "audit_year": row.audit_year,
        "audit_scope_start": row.audit_scope_start.isoformat() if row.audit_scope_start else "",
        "audit_scope_end": row.audit_scope_end.isoformat() if row.audit_scope_end else "",
        "start_date": row.start_date.isoformat() if row.start_date else "",
        "end_date": row.end_date.isoformat() if row.end_date else "",
        "status": row.status,
        "project_leader_name": row.project_leader.display_name if row.project_leader else "",
        "manager_name": row.manager.display_name if row.manager else "",
        "field_leader_name": row.field_leader.display_name if row.field_leader else "",
        "department": row.department or "",
        "charge_with_tax": row.charge_with_tax or "",
        "charge_without_tax": row.charge_without_tax or "",
        "can_edit_home": False,
    }


@router.get("/api/home/projects/{project_id}")
def get_home_project_detail(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    """Return directory-level project details for every signed-in user."""
    project = get_or_404(db, Project, project_id, "项目")
    payload = _home_project_payload(project)
    members = db.execute(
        select(ProjectMember).where(ProjectMember.project_id == project.id).order_by(ProjectMember.id)
    ).scalars().all()
    can_edit = bool(is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id})
    if can_edit:
        staff_rows, staff_by_user_id = _home_staff_directory(db)
    else:
        needed_ids = {item.user_id for item in members} | {user.id}
        needed_users = db.execute(select(User).where(User.id.in_(needed_ids))).scalars().all() if needed_ids else []
        staff_rows, staff_by_user_id = _home_staff_directory(db, users=needed_users)
    payload["members"] = [
        {
            "id": item.id,
            "user_id": item.user_id,
            "name": item.user.display_name if item.user else "",
            "username": item.user.username if item.user else "",
            "role": item.role_on_project,
            "module": item.module,
            "workload": item.workload,
            "office": staff_by_user_id.get(item.user_id, {}).get("office", ""),
            "department": staff_by_user_id.get(item.user_id, {}).get("department", ""),
            "job_title": staff_by_user_id.get(item.user_id, {}).get("job_title", ""),
            "workcode": staff_by_user_id.get(item.user_id, {}).get("workcode", ""),
        }
        for item in members
    ]
    payload["scope"] = {
        "start": project.audit_scope_start.isoformat() if project.audit_scope_start else "",
        "end": project.audit_scope_end.isoformat() if project.audit_scope_end else "",
        "description": project.scope_description or "",
    }
    systems, systems_message = _home_systems_from_it_overview(db, project.id)
    payload["systems"] = systems
    payload["business_revenue"] = project.business_revenue or ""
    payload["can_edit_home"] = can_edit
    member_user_ids = {item.user_id for item in members}
    claim = db.execute(
        select(ProjectMemberClaim).where(
            ProjectMemberClaim.project_id == project.id,
            ProjectMemberClaim.user_id == user.id,
        )
    ).scalar_one_or_none()
    payload["claim_status"] = "member" if user.id in member_user_ids else (claim.status if claim else "none")
    if payload["can_edit_home"]:
        pending_claims = db.execute(
            select(ProjectMemberClaim)
            .where(ProjectMemberClaim.project_id == project.id, ProjectMemberClaim.status == "pending")
            .order_by(ProjectMemberClaim.created_at, ProjectMemberClaim.id)
        ).scalars().all()
        payload["pending_claims"] = [
            {
                "id": item.id,
                "user_id": item.user_id,
                "name": item.user.display_name if item.user else "",
                "username": item.user.username if item.user else "",
                "office": staff_by_user_id.get(item.user_id, {}).get("office", ""),
                "department": staff_by_user_id.get(item.user_id, {}).get("department", ""),
                "job_title": staff_by_user_id.get(item.user_id, {}).get("job_title", ""),
                "workcode": staff_by_user_id.get(item.user_id, {}).get("workcode", ""),
                "created_at": item.created_at.isoformat() if item.created_at else "",
            }
            for item in pending_claims
        ]
        payload["role_assignments"] = {
            "project_leader_user_id": project.project_leader_user_id,
            "manager_user_id": project.manager_user_id,
            "field_leader_user_id": project.field_leader_user_id,
            "quality_reviewer_user_id": project.quality_reviewer_user_id,
            "director_user_id": project.director_user_id,
            "partner_user_id": project.partner_user_id,
        }
        payload["available_users"] = staff_rows
    payload["data_availability"] = {
        "systems": systems_message,
        "business_revenue": "未同步主营业务收入；项目收款金额不能替代被审计单位的主营业务收入。",
    }
    return payload


@router.post("/api/home/projects/{project_id}/claims", status_code=201)
def create_home_project_claim(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    """Submit (or resubmit) the current user's request to join a project."""
    project = get_or_404(db, Project, project_id, "项目")
    if user.status != "active":
        raise HTTPException(status_code=400, detail="仅在职账号可领用项目")
    already_member = db.execute(
        select(ProjectMember.id).where(ProjectMember.project_id == project.id, ProjectMember.user_id == user.id)
    ).scalar_one_or_none()
    if already_member:
        return {"status": "member", "message": "您已在该项目成员中"}
    claim = db.execute(
        select(ProjectMemberClaim).where(
            ProjectMemberClaim.project_id == project.id,
            ProjectMemberClaim.user_id == user.id,
        )
    ).scalar_one_or_none()
    if claim and claim.status == "pending":
        return {"status": "pending", "message": "领用申请已提交，等待审批"}
    if claim is None:
        claim = ProjectMemberClaim(project_id=project.id, user_id=user.id, status="pending")
        db.add(claim)
    else:
        claim.status = "pending"
        claim.reviewer_user_id = None
        claim.reviewed_at = None
    db.flush()
    record_audit_log(
        db, None, "home_project_claim_submitted", user=user, target_type="project_member_claim",
        target_id=claim.id, project_id=project.id, details={"claim_user_id": user.id},
    )
    db.commit()
    return {"status": "pending", "id": claim.id, "message": "领用申请已提交，等待项目负责人、项目负责经理或管理员审批"}


@router.post("/api/home/projects/{project_id}/claims/approve")
def approve_home_project_claims(
    project_id: int,
    body: HomeProjectClaimApprovalIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    """Batch approve pending claims and create the corresponding project members."""
    project = get_or_404(db, Project, project_id, "项目")
    if not (is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id}):
        raise HTTPException(status_code=403, detail="仅项目负责人、项目负责经理或管理员可审批领用申请")
    claim_ids = list(dict.fromkeys(int(item) for item in body.claim_ids if int(item) > 0))
    claims = db.execute(
        select(ProjectMemberClaim).where(
            ProjectMemberClaim.project_id == project.id,
            ProjectMemberClaim.status == "pending",
            ProjectMemberClaim.id.in_(claim_ids),
        )
    ).scalars().all()
    if not claims:
        raise HTTPException(status_code=400, detail="未找到可审批的领用申请")
    existing_user_ids = set(db.execute(
        select(ProjectMember.user_id).where(ProjectMember.project_id == project.id)
    ).scalars().all())
    created = 0
    now = datetime.utcnow()
    for claim in claims:
        if claim.user_id not in existing_user_ids:
            db.add(ProjectMember(
                project_id=project.id,
                user_id=claim.user_id,
                role_on_project="项目成员",
                module="待分工",
                workload="",
            ))
            existing_user_ids.add(claim.user_id)
            created += 1
        claim.status = "approved"
        claim.reviewer_user_id = user.id
        claim.reviewed_at = now
        record_audit_log(
            db, None, "home_project_claim_approved", user=user, target_type="project_member_claim",
            target_id=claim.id, project_id=project.id, details={"claim_user_id": claim.user_id},
        )
    db.commit()
    return {"approved": len(claims), "members_created": created}


@router.patch("/api/home/projects/{project_id}")
def update_home_project(
    project_id: int,
    body: HomeProjectUpdateIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    """Update directory fields from the homepage for the assigned leaders or an administrator."""
    project = get_or_404(db, Project, project_id, "项目")
    if not (is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id}):
        raise HTTPException(status_code=403, detail="仅项目负责人、项目负责经理或管理员可修改首页项目资料")

    payload = body.model_dump(exclude_unset=True)
    member_workloads = payload.pop("member_workloads", None)
    role_fields = {
        "project_leader_user_id", "manager_user_id", "field_leader_user_id",
        "quality_reviewer_user_id", "director_user_id", "partner_user_id",
    }
    role_updates = {key: payload[key] for key in role_fields if key in payload}
    for field, user_id in role_updates.items():
        if user_id is not None:
            assigned = get_or_404(db, User, int(user_id), "项目角色人员")
            if assigned.status != "active":
                raise HTTPException(status_code=400, detail="项目角色人员必须为在职账号")
    changed_role_updates = {
        field: user_id
        for field, user_id in role_updates.items()
        if user_id != getattr(project, field)
    }
    _validate_project_reviewers(db, changed_role_updates)
    if "quality_reviewer_user_id" in role_updates:
        selected_quality_id = role_updates["quality_reviewer_user_id"]
        existing_quality_ids = _quality_reviewer_ids(project)
        _set_quality_reviewer_ids(
            project,
            [selected_quality_id, *(item for item in existing_quality_ids if item != selected_quality_id)] if selected_quality_id else [],
        )
    if member_workloads is not None:
        members = {
            item.id: item
            for item in db.execute(select(ProjectMember).where(ProjectMember.project_id == project.id)).scalars().all()
        }
        for item in member_workloads:
            member = members.get(item["member_id"])
            if member is None:
                raise HTTPException(status_code=400, detail="人员安排记录不属于当前项目")
            member.workload = str(item.get("workload") or "").strip()
    for key, value in payload.items():
        setattr(project, key, value.strip() if isinstance(value, str) else value)
    if project.start_date and project.end_date and project.start_date > project.end_date:
        raise HTTPException(status_code=400, detail="预计开始日期不能晚于预计结束日期")
    if project.audit_scope_start and project.audit_scope_end and project.audit_scope_start > project.audit_scope_end:
        raise HTTPException(status_code=400, detail="审计期间开始日期不能晚于结束日期")
    record_audit_log(
        db, None, "home_project_updated", user=user, target_type="project", target_id=project.id, project_id=project.id,
        details={"fields": sorted([*payload.keys(), *(["member_workloads"] if member_workloads is not None else [])])},
    )
    db.commit()
    db.refresh(project)
    return _home_project_payload(project)


@router.post("/api/home/projects/{project_id}/members/batch", status_code=201)
def create_home_project_members_batch(
    project_id: int,
    body: HomeProjectMemberBatchIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    if not (is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id}):
        raise HTTPException(status_code=403, detail="仅项目负责人、项目负责经理或管理员可维护人员安排")
    if not body.items:
        raise HTTPException(status_code=400, detail="请至少选择一名成员")
    existing = {
        (item.user_id, str(item.module or "").strip())
        for item in db.execute(select(ProjectMember).where(ProjectMember.project_id == project.id)).scalars().all()
    }
    pending: set[tuple[int, str]] = set()
    prepared: list[MemberIn] = []
    for item in body.items:
        member_user = get_or_404(db, User, item.user_id, "成员")
        if member_user.status != "active":
            raise HTTPException(status_code=400, detail="仅可添加在职账号")
        key = (item.user_id, str(item.module or "").strip())
        if key in existing or key in pending:
            raise HTTPException(status_code=400, detail="存在已安排或重复选择的成员及模块")
        pending.add(key)
        prepared.append(item)
    records = [ProjectMember(project_id=project.id, **item.model_dump()) for item in prepared]
    db.add_all(records)
    db.flush()
    for item in records:
        record_audit_log(db, None, "home_project_member_added", user=user, target_type="project_member", target_id=item.id, project_id=project.id, details={"member_user_id": item.user_id})
    db.commit()
    return {"created": len(records)}


@router.post("/api/home/projects/{project_id}/members", status_code=201)
def create_home_project_member(
    project_id: int,
    body: MemberIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    if not (is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id}):
        raise HTTPException(status_code=403, detail="仅项目负责人、项目负责经理或管理员可维护人员安排")
    member_user = get_or_404(db, User, body.user_id, "成员")
    if member_user.status != "active":
        raise HTTPException(status_code=400, detail="仅可添加在职账号")
    module = str(body.module or "").strip()
    exists = db.execute(
        select(ProjectMember.id).where(
            ProjectMember.project_id == project.id,
            ProjectMember.user_id == body.user_id,
            ProjectMember.module == module,
        )
    ).scalar_one_or_none()
    if exists:
        raise HTTPException(status_code=400, detail="该成员已在当前项目的相同模块中")
    item = ProjectMember(
        project_id=project.id,
        user_id=body.user_id,
        role_on_project=str(body.role_on_project or "").strip(),
        module=module,
        workload=str(body.workload or "").strip(),
    )
    db.add(item)
    db.flush()
    record_audit_log(db, None, "home_project_member_added", user=user, target_type="project_member", target_id=item.id, project_id=project.id, details={"member_user_id": item.user_id})
    db.commit()
    db.refresh(item)
    return {"id": item.id, "user_id": item.user_id, "role": item.role_on_project, "module": item.module, "workload": item.workload}


@router.delete("/api/home/projects/{project_id}/members/{member_id}", status_code=204)
def delete_home_project_member(
    project_id: int,
    member_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> Response:
    project = get_or_404(db, Project, project_id, "项目")
    if not (is_admin(user) or user.id in {project.project_leader_user_id, project.manager_user_id}):
        raise HTTPException(status_code=403, detail="仅项目负责人、项目负责经理或管理员可维护人员安排")
    item = get_or_404(db, ProjectMember, member_id, "项目成员")
    if item.project_id != project.id:
        raise HTTPException(status_code=400, detail="人员安排记录不属于当前项目")
    record_audit_log(db, None, "home_project_member_removed", user=user, target_type="project_member", target_id=item.id, project_id=project.id, details={"member_user_id": item.user_id})
    db.delete(item)
    db.commit()
    return Response(status_code=204)


@router.post("/api/projects", status_code=201)
def create_project(body: ProjectIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "projects", "edit")
    payload = body.model_dump()
    quality_reviewer_user_ids = payload.pop("quality_reviewer_user_ids", [])
    payload["director_user_id"] = None
    payload["partner_user_id"] = payload.get("partner_user_id") or _default_user_id(db, "任德昀")
    payload["project_leader_user_id"] = payload.get("project_leader_user_id") or _default_user_id(db, "任德昀")
    if not quality_reviewer_user_ids:
        quality_reviewer_user_ids = [user_id for user_id in (_default_user_id(db, "占志权"), _default_user_id(db, "李瑞"), _default_user_id(db, "闫晓濛")) if user_id]
    if not quality_reviewer_user_ids and payload.get("quality_reviewer_user_id"):
        quality_reviewer_user_ids = [payload["quality_reviewer_user_id"]]
    payload["quality_reviewer_user_ids"] = quality_reviewer_user_ids
    _validate_project_reviewers(db, payload)
    payload.pop("quality_reviewer_user_ids", None)
    if payload.get("client_id"):
        client = get_or_404(db, Client, payload["client_id"], "客户")
        payload["entity_name"] = payload.get("entity_name") or client.entity_name
    if not payload.get("audit_scope_start") or not payload.get("audit_scope_end"):
        start, end = default_audit_scope()
        payload["audit_scope_start"] = payload.get("audit_scope_start") or start
        payload["audit_scope_end"] = payload.get("audit_scope_end") or end
    item = Project(**payload, creator_user_id=user.id)
    _set_quality_reviewer_ids(item, quality_reviewer_user_ids)
    workspace: dict[str, Any] | None = None
    committed = False
    try:
        db.add(item)
        db.flush()
        if not str(item.project_root or "").strip():
            item.project_root = str(default_project_root(item))
        workspace = initialize_project_workspace(db, item)
        record_audit_log(
            db,
            None,
            "project_created",
            user=user,
            target_type="project",
            target_id=item.id,
            project_id=item.id,
            details={
                "name": item.name,
                "project_root": item.project_root,
                "workpapers_created": workspace["workpapers_created"],
                "header_fields_changed": workspace["header_fields_changed"],
            },
        )
        db.commit()
        committed = True
        db.refresh(item)
    except HTTPException:
        db.rollback()
        if not committed:
            rollback_project_workspace(workspace)
        raise
    except Exception as exc:
        db.rollback()
        if not committed:
            rollback_project_workspace(workspace)
        raise HTTPException(
            status_code=500,
            detail=f"项目创建失败，数据库与本次生成文件已回滚：{exc}",
        ) from exc
    assert workspace is not None
    workspace = public_workspace_report(workspace)
    data = obj_dict(item)
    data["template_workpapers_created"] = workspace["workpapers_created"]
    data["workspace"] = workspace
    return data


@router.patch("/api/projects/{project_id}")
def update_project(project_id: int, body: ProjectIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    item = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(item, user)
    payload = body.model_dump()
    quality_reviewer_user_ids = payload.pop("quality_reviewer_user_ids", [])
    payload["director_user_id"] = None
    payload["partner_user_id"] = payload.get("partner_user_id") or _default_user_id(db, "任德昀")
    if not quality_reviewer_user_ids and payload.get("quality_reviewer_user_id"):
        quality_reviewer_user_ids = [payload["quality_reviewer_user_id"]]
    payload["quality_reviewer_user_ids"] = quality_reviewer_user_ids
    _validate_project_reviewers(db, payload)
    payload.pop("quality_reviewer_user_ids", None)
    if payload.get("client_id"):
        client = get_or_404(db, Client, payload["client_id"], "客户")
        payload["entity_name"] = payload.get("entity_name") or client.entity_name
    apply_patch_to_model(item, payload)
    _set_quality_reviewer_ids(item, quality_reviewer_user_ids)
    record_audit_log(db, None, "project_updated", user=user, target_type="project", target_id=item.id, project_id=item.id)
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.post("/api/projects/{project_id}/audit-scope/infer")
def infer_project_audit_scope(project_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    item = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(item, user)
    workpapers = db.execute(select(Workpaper).where(Workpaper.project_id == project_id).order_by(Workpaper.code)).scalars().all()
    inferred = infer_audit_scope_from_workpapers(workpapers)
    if not inferred:
        raise HTTPException(status_code=404, detail="未从底稿首页识别到截止日或审计期间")
    item.audit_year = inferred.audit_year
    item.audit_scope_start = inferred.start
    item.audit_scope_end = inferred.end
    db.commit()
    db.refresh(item)
    data = obj_dict(item)
    data["audit_scope_source"] = inferred.source
    data["audit_scope_raw_value"] = inferred.raw_value
    return data


@router.delete("/api/projects/{project_id}", status_code=204)
def delete_project(project_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> Response:
    item = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "manage")
    ensure_project_editor(item, user)
    db.delete(item)
    db.commit()
    return Response(status_code=204)


@router.get("/api/projects/{project_id}/contacts")
def list_contacts(project_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    rows = db.execute(
        select(EnterpriseContact).where(EnterpriseContact.project_id == project_id).order_by(EnterpriseContact.id)
    ).scalars().all()
    return list_dict(rows)


@router.post("/api/projects/{project_id}/contacts", status_code=201)
def create_contact(project_id: int, body: ContactIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(project, user)
    item = EnterpriseContact(project_id=project_id, **body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.delete("/api/contacts/{contact_id}", status_code=204)
def delete_contact(contact_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> Response:
    item = get_or_404(db, EnterpriseContact, contact_id, "对接人")
    project = get_or_404(db, Project, item.project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(project, user)
    db.delete(item)
    db.commit()
    return Response(status_code=204)


@router.get("/api/projects/{project_id}/members")
def list_members(project_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    stmt = select(ProjectMember).where(ProjectMember.project_id == project_id).order_by(ProjectMember.id)
    if project.id not in supervised_project_ids(db, user):
        stmt = stmt.where(ProjectMember.user_id == user.id)
    rows = db.execute(stmt).scalars().all()
    payload = []
    for row in rows:
        data = obj_dict(row)
        data["display_name"] = row.user.display_name
        data["username"] = row.user.username
        data["role_name"] = row.user.role.name if row.user.role else ""
        payload.append(data)
    return payload


@router.get("/api/projects/{project_id}/workpaper-tree", response_model=list[WorkpaperTreeNodeOut])
def get_project_workpaper_tree(
    project_id: int,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    workpaper_stmt = select(Workpaper).where(Workpaper.project_id == project_id).order_by(Workpaper.stage, Workpaper.code)
    workpapers = db.execute(workpaper_stmt).scalars().all()
    attachment_stmt = select(Attachment).where(Attachment.project_id == project_id).order_by(Attachment.index_no, Attachment.id)
    attachments = db.execute(attachment_stmt).scalars().all()
    return build_workpaper_tree(workpapers, attachments, project.project_root or "")


@router.get("/api/projects/{project_id}/planned-workpapers")
def list_planned_workpapers(project_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_viewer(db, project, user)
    rows = db.execute(select(Workpaper).where(Workpaper.project_id == project_id).order_by(Workpaper.stage, Workpaper.code)).scalars().all()
    return [obj_dict(row) for row in rows]


@router.put("/api/projects/{project_id}/planned-workpapers")
def update_planned_workpapers(
    project_id: int,
    body: dict[str, Any] = Body(default_factory=dict),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(project, user)
    items = body.get("items") if isinstance(body, dict) else []
    if not isinstance(items, list):
        raise HTTPException(status_code=400, detail="底稿目录格式必须是列表")
    existing = {row.code: row for row in db.execute(select(Workpaper).where(Workpaper.project_id == project_id)).scalars().all()}
    seen: set[str] = set()
    for raw in items:
        if not isinstance(raw, dict):
            continue
        code = str(raw.get("code") or "").strip().upper()
        name = str(raw.get("name") or "").strip()
        stage = str(raw.get("stage") or "execution").strip() or "execution"
        if not code or not name:
            continue
        if code in seen:
            continue
        seen.add(code)
        row = existing.get(code)
        if row is None:
            row = Workpaper(project_id=project_id, code=code, name=name, stage=stage, status="draft", preparer_user_id=user.id)
            db.add(row)
        else:
            row.name = name
            row.stage = stage
    db.commit()
    rows = db.execute(select(Workpaper).where(Workpaper.project_id == project_id).order_by(Workpaper.stage, Workpaper.code)).scalars().all()
    return [obj_dict(row) for row in rows]


@router.post("/api/projects/{project_id}/members", status_code=201)
def create_member(project_id: int, body: MemberIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(project, user)
    get_or_404(db, User, body.user_id, "用户")
    item = ProjectMember(project_id=project_id, **body.model_dump())
    db.add(item)
    db.flush()
    record_audit_log(db, None, "project_member_added", user=user, target_type="project_member", target_id=item.id, project_id=project_id, details={"member_user_id": body.user_id})
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.delete("/api/members/{member_id}", status_code=204)
def delete_member(member_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> Response:
    item = get_or_404(db, ProjectMember, member_id, "项目成员")
    project = get_or_404(db, Project, item.project_id, "项目")
    ensure_feature_permission(db, user, "projects", "edit")
    ensure_project_editor(project, user)
    record_audit_log(db, None, "project_member_removed", user=user, target_type="project_member", target_id=item.id, project_id=project.id, details={"member_user_id": item.user_id})
    db.delete(item)
    db.commit()
    return Response(status_code=204)


@router.get("/api/tasks")
def list_tasks(
    projectId: Optional[int] = Query(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> list[dict[str, Any]]:
    visible_ids = visible_project_ids(db, user)
    if not visible_ids:
        return []
    supervised_ids = supervised_project_ids(db, user)
    own_only_ids = visible_ids - supervised_ids
    stmt = select(Task).where(
        or_(
            Task.project_id.in_(supervised_ids),
            (Task.project_id.in_(own_only_ids)) & (Task.owner_user_id == user.id),
        )
    ).order_by(Task.id)
    if projectId is not None:
        project = get_or_404(db, Project, projectId, "项目")
        ensure_project_viewer(db, project, user)
        stmt = stmt.where(Task.project_id == projectId)
    rows = db.execute(stmt).scalars().all()
    project_ids = {row.project_id for row in rows}
    projects_by_id = {
        row.id: row
        for row in db.execute(select(Project).where(Project.id.in_(project_ids))).scalars().all()
    } if project_ids else {}
    payload = []
    for row in rows:
        data = obj_dict(row)
        data["owner_name"] = row.owner.display_name if row.owner else ""
        project = projects_by_id.get(row.project_id)
        data["project_name"] = project.name if project else ""
        data["entity_name"] = project.entity_name if project else ""
        payload.append(data)
    return payload


@router.post("/api/tasks", status_code=201)
def create_task(body: TaskIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    project = get_or_404(db, Project, body.project_id, "项目")
    ensure_feature_permission(db, user, "resourcePlan", "edit")
    ensure_project_editor(project, user)
    item = Task(**body.model_dump())
    db.add(item)
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.patch("/api/tasks/{task_id}")
def update_task(task_id: int, body: TaskIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    item = get_or_404(db, Task, task_id, "任务")
    if body.project_id != item.project_id:
        raise HTTPException(status_code=400, detail="任务不支持跨项目修改；请在目标项目重新创建任务")
    project = get_or_404(db, Project, item.project_id, "项目")
    ensure_feature_permission(db, user, "resourcePlan", "edit")
    ensure_project_editor(project, user)
    apply_patch_to_model(item, body.model_dump())
    db.commit()
    db.refresh(item)
    return obj_dict(item)


@router.delete("/api/tasks/{task_id}", status_code=204)
def delete_task(task_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> Response:
    item = get_or_404(db, Task, task_id, "任务")
    project = get_or_404(db, Project, item.project_id, "项目")
    ensure_feature_permission(db, user, "resourcePlan", "edit")
    ensure_project_editor(project, user)
    db.delete(item)
    db.commit()
    return Response(status_code=204)


@router.post("/api/projects/{project_id}/init-from-prior")
def init_from_prior(
    project_id: int,
    body: InitFromPriorIn,
    db: Session = Depends(get_db),
    user: User = Depends(current_user),
) -> dict[str, Any]:
    project = get_or_404(db, Project, project_id, "项目")
    ensure_project_editor(project, user)
    prior_id = body.prior_project_id or project.prior_project_id
    if not prior_id:
        raise HTTPException(status_code=400, detail="未指定以前年度项目")
    prior = get_or_404(db, Project, prior_id, "以前年度项目")

    prior_workpapers = db.execute(
        select(Workpaper).where(Workpaper.project_id == prior.id).order_by(Workpaper.id)
    ).scalars().all()
    created = 0
    copied_files = 0
    for old in prior_workpapers:
        new_name, name_changed = replace_audit_year(old.name, prior.audit_year, project.audit_year)
        new_file_path = ""
        if body.copy_files:
            new_file_path = copy_workpaper_file(old.file_path, project, new_name, prior.audit_year)
            copied_files += 1 if new_file_path else 0
        item = Workpaper(
            project_id=project.id,
            code=old.code,
            name=new_name,
            stage=old.stage,
            file_path=new_file_path or old.file_path,
            source_workpaper_id=old.id,
            status="draft",
            preparer_user_id=old.preparer_user_id,
            year_updated=name_changed,
            extracted_fields_json=old.extracted_fields_json,
        )
        db.add(item)
        created += 1
    db.commit()
    return {"created_workpapers": created, "copied_files": copied_files}
