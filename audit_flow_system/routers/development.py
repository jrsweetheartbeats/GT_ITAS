from __future__ import annotations

from datetime import date, datetime, timedelta
from hashlib import sha256
from io import BytesIO
import json
import re
from typing import Any, Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..core.db import get_db
from ..core.security import current_user, ensure_feature_permission
from ..models import (
    CompetencyDimension,
    DevelopmentEmployee,
    DevelopmentExercise,
    DevelopmentExerciseSubmission,
    DevelopmentIssue,
    DevelopmentPlan,
    DevelopmentReview,
    EmployeeCompetency,
    MonthlyAssessment,
    MonthlyPlan,
    User,
    WeeklyTask,
)
from ..schemas import (
    CompetencyAssessmentIn,
    DevelopmentEmployeeIn,
    DevelopmentEmployeePatchIn,
    DevelopmentExerciseIn,
    DevelopmentIssueIn,
    DevelopmentIssuePatchIn,
    DevelopmentPlanIn,
    DevelopmentReviewIn,
    DevelopmentSubmissionIn,
    DraftPlanIn,
    DevelopmentMonthlyAssessmentIn,
    MonthlyPlanIn,
    WeeklyTaskIn,
    WeeklyTaskPatchIn,
)


router = APIRouter(prefix="/api/development", tags=["development"])

TRAINING_MANAGEMENT_ROLES = {"admin", "partner", "quality", "director", "senior_manager", "manager"}


def _can_manage_training(user: User) -> bool:
    return bool(user.role and user.role.code in TRAINING_MANAGEMENT_ROLES)


def _require_training_management(user: User) -> None:
    if not _can_manage_training(user):
        raise HTTPException(status_code=403, detail="只有经理及以上层级可以管理培养档案和计划")

TEMPLATES: dict[str, dict[str, Any]] = {
    "it_audit_foundation": {
        "name": "IT审计基础型",
        "direction": "ITGC + ITAC + 证据判断",
        "topics": ["财审到IT审计", "ITGC", "ITAC", "Interface/IPE"],
    },
    "finance_to_it": {
        "name": "财审转IT审计型",
        "direction": "业务财务映射 + ITGC",
        "topics": ["业务财务映射", "ITGC", "ITAC", "综合案例"],
    },
    "itac_focus": {
        "name": "ITAC强化型",
        "direction": "ITAC + Interface/IPE",
        "topics": ["业务流程与控制", "ITAC设计", "Interface/IPE", "陌生案例"],
    },
    "data_focus": {
        "name": "数据分析强化型",
        "direction": "数据需求 + 规则 + 审计结论",
        "topics": ["数据需求", "数据完整性", "异常规则", "结论与底稿"],
    },
    "project_manager": {
        "name": "项目经理培养型",
        "direction": "Scope + PBC + 复核 + 沟通",
        "topics": ["Scope与计划", "PBC与客户沟通", "复核与问题定性", "项目交付"],
    },
}

TOPIC_LIBRARY: dict[str, dict[str, str]] = {
    "财审到IT审计": {"learning": "从审计目标、业务流程、系统处理和数据证据建立完整链路。", "exercise": "选择一个熟悉财审程序，改造成系统或数据驱动的补充程序。", "deliverable": "业务—系统—科目/认定映射表", "acceptance": "能说明为什么测、测什么、证据为何充分。"},
    "业务财务映射": {"learning": "连接业务发生、系统处理、会计处理、科目及报表认定。", "exercise": "完成收入或采购流程的端到端映射。", "deliverable": "端到端流程映射", "acceptance": "随机抽取三个环节均能独立说明。"},
    "ITGC": {"learning": "访问管理、程序变更、系统运维、备份恢复的控制目标、风险和证据。", "exercise": "离职账号未及时停用案例。", "deliverable": "完整ITGC测试底稿", "acceptance": "从零写出控制目标、程序、证据、异常和结论。"},
    "ITAC": {"learning": "控制逻辑、配置、输入输出、依赖数据和异常处理。", "exercise": "自动审批、三单匹配或自动计算案例。", "deliverable": "ITAC测试设计", "acceptance": "测试覆盖配置、运行、依赖和异常。"},
    "ITAC设计": {"learning": "从风险和业务规则拆解ITAC，明确配置、输入、处理和输出。", "exercise": "不给模板独立设计一项ITAC。", "deliverable": "可Review的ITAC底稿", "acceptance": "逻辑、样本、IPE依赖和结论相互一致。"},
    "Interface/IPE": {"learning": "接口字段映射、批次数量、失败重跑，以及IPE总体完整性与准确性。", "exercise": "ERP报表作为审计证据的完整性测试。", "deliverable": "Interface/IPE案例", "acceptance": "覆盖来源、参数、范围、完整性、准确性和导出后修改。"},
    "系统技术基础": {"learning": "应用、数据库、操作系统、接口、日志和备份之间的关系。", "exercise": "绘制一个常见ERP技术链路并标注审计证据位置。", "deliverable": "系统技术关系图", "acceptance": "能解释各层职责及可获取证据。"},
    "数据需求": {"learning": "从审计问题推导字段、期间、总体、关联键和取数口径。", "exercise": "为权限或业务异常设计数据需求清单。", "deliverable": "数据需求说明", "acceptance": "需求可由客户按字段和口径直接执行。"},
    "数据完整性": {"learning": "记录数、期间范围、主键、汇总勾稽和源数据独立核对。", "exercise": "识别一份系统导出数据的完整性缺口。", "deliverable": "数据完整性验证底稿", "acceptance": "总体边界清楚且有独立核对。"},
    "异常规则": {"learning": "将业务规则转成可复核的筛选条件，并区分线索和结论。", "exercise": "设计三个异常规则并说明误报处理。", "deliverable": "异常规则与核验清单", "acceptance": "规则、数据字段、阈值和后续核验可追溯。"},
    "结论与底稿": {"learning": "从分析结果回到审计目标、证据充分性和影响结论。", "exercise": "把异常清单写成可Review底稿。", "deliverable": "数据分析底稿", "acceptance": "不把异常线索直接等同控制缺陷。"},
    "Scope与计划": {"learning": "识别财务相关系统、服务组织、接口与关键报表，形成范围理由。", "exercise": "根据陌生客户背景设计Scope和项目计划。", "deliverable": "Scope与计划草案", "acceptance": "范围与重大流程、科目和系统依赖一致。"},
    "PBC与客户沟通": {"learning": "把测试目的转成客户可理解、可执行的PBC需求。", "exercise": "模拟首次资料沟通和缺口追问。", "deliverable": "PBC清单与沟通记录", "acceptance": "需求含对象、期间、字段、格式和用途。"},
    "复核与问题定性": {"learning": "区分资料不足、个别例外、执行缺陷和设计缺陷。", "exercise": "复核一份含证据缺口的底稿。", "deliverable": "Review意见与问题定性", "acceptance": "事实、判断、影响和改进要求分开表达。"},
    "项目交付": {"learning": "进度、风险、问题升级、复核关闭和交付检查。", "exercise": "模拟项目收尾和管理层汇报。", "deliverable": "项目状态与交付清单", "acceptance": "阻塞、责任人、截止日和下一步明确。"},
    "业务流程与控制": {"learning": "从业务风险识别人工控制、自动控制、接口和关键报表。", "exercise": "拆解一个销售或采购流程。", "deliverable": "业务控制矩阵", "acceptance": "风险、控制和审计应对一一对应。"},
    "陌生案例": {"learning": "综合运用Scope、ITGC、ITAC、IPE和异常判断。", "exercise": "20分钟陌生ERP案例现场验收。", "deliverable": "综合案例输出", "acceptance": "不依赖历史底稿完成结构化判断。"},
    "综合案例": {"learning": "综合运用本月能力并形成可Review成果。", "exercise": "陌生业务、多个系统和接口的综合案例。", "deliverable": "综合案例底稿", "acceptance": "范围、程序、证据、异常和结论闭环。"},
}

COMPETENCY_NAMES = {
    "financial_audit": "财务/会计",
    "itgc": "ITGC",
    "itac": "ITAC",
    "interface_ipe": "Interface/IPE",
    "data_analysis": "数据分析",
    "system_foundation": "IT技术基础",
    "communication": "客户沟通",
    "project_management": "项目管理",
    "ai_application": "AI应用",
}


def _json(value: str, fallback: Any) -> Any:
    try:
        return json.loads(value or "")
    except (TypeError, json.JSONDecodeError):
        return fallback


def _require(db: Session, model: Any, row_id: int, label: str) -> Any:
    row = db.get(model, row_id)
    if row is None:
        raise HTTPException(status_code=404, detail=f"{label}不存在")
    return row


def _latest_competencies(db: Session, employee_id: int) -> tuple[dict[str, float], list[dict[str, Any]]]:
    rows = db.execute(
        select(EmployeeCompetency)
        .where(EmployeeCompetency.employee_id == employee_id)
        .order_by(EmployeeCompetency.assessed_on, EmployeeCompetency.id)
    ).scalars().all()
    latest: dict[str, float] = {}
    history: list[dict[str, Any]] = []
    for row in rows:
        code = row.competency.code
        latest[code] = round(float(row.score), 1)
        history.append({
            "id": row.id, "code": code, "name": row.competency.name,
            "score": round(float(row.score), 1), "assessedOn": row.assessed_on,
            "source": row.source, "note": row.note,
        })
    return latest, history


def _employee_payload(db: Session, row: DevelopmentEmployee, include_detail: bool = False) -> dict[str, Any]:
    latest, history = _latest_competencies(db, row.id)
    payload = {
        "id": row.id, "userId": row.user_id, "code": row.code, "name": row.name,
        "currentRole": row.current_role, "hiredOn": row.hired_on,
        "mentorUserId": row.mentor_user_id,
        "mentorName": row.mentor.display_name if row.mentor else "",
        "direction": row.direction, "halfYearGoal": row.half_year_goal, "yearGoal": row.year_goal,
        "advantages": row.advantages, "weaknesses": row.weaknesses,
        "currentFocus": row.current_focus, "recentIssues": row.recent_issues,
        "mentorObservation": row.mentor_observation, "developmentIntent": row.development_intent,
        "status": row.status, "competencies": latest, "sourceReference": row.source_reference,
    }
    if include_detail:
        payload["competencyHistory"] = history
        payload["questionnaire"] = _json(row.questionnaire_json, {})
    return payload


def _task_payload(row: WeeklyTask) -> dict[str, Any]:
    return {
        "id": row.id, "monthlyPlanId": row.monthly_plan_id, "weekNo": row.week_no,
        "topic": row.topic, "learningContent": row.learning_content,
        "exerciseCase": row.exercise_case, "deliverable": row.deliverable,
        "acceptanceCriteria": row.acceptance_criteria, "ownerUserId": row.owner_user_id,
        "ownerName": row.owner.display_name if row.owner else "", "dueDate": row.due_date,
        "status": row.status, "progress": row.progress,
    }


def _monthly_payload(db: Session, row: MonthlyPlan) -> dict[str, Any]:
    tasks = db.execute(
        select(WeeklyTask).where(WeeklyTask.monthly_plan_id == row.id).order_by(WeeklyTask.week_no, WeeklyTask.id)
    ).scalars().all()
    assessment = db.execute(select(MonthlyAssessment).where(MonthlyAssessment.monthly_plan_id == row.id)).scalar_one_or_none()
    return {
        "id": row.id, "developmentPlanId": row.development_plan_id, "monthNo": row.month_no,
        "title": row.title, "monthGoal": row.month_goal, "generationBasis": row.generation_basis,
        "rationale": row.rationale, "expectedProject": row.expected_project,
        "startsOn": row.starts_on, "endsOn": row.ends_on, "status": row.status,
        "generatedBy": row.generated_by, "publishedAt": row.published_at,
        "tasks": [_task_payload(task) for task in tasks],
        "assessment": _assessment_payload(assessment) if assessment else None,
    }


def _plan_payload(db: Session, row: DevelopmentPlan) -> dict[str, Any]:
    months = db.execute(
        select(MonthlyPlan).where(MonthlyPlan.development_plan_id == row.id).order_by(MonthlyPlan.month_no, MonthlyPlan.id)
    ).scalars().all()
    return {
        "id": row.id, "employeeId": row.employee_id, "title": row.title, "planType": row.plan_type,
        "startsOn": row.starts_on, "endsOn": row.ends_on, "halfYearGoal": row.half_year_goal,
        "yearGoal": row.year_goal, "status": row.status, "sourceType": row.source_type,
        "sourceReference": row.source_reference,
        "months": [_monthly_payload(db, month) for month in months],
    }


def _issue_payload(db: Session, row: DevelopmentIssue) -> dict[str, Any]:
    employee = db.get(DevelopmentEmployee, row.employee_id)
    return {
        "id": row.id, "employeeId": row.employee_id, "employeeName": employee.name if employee else "",
        "reviewId": row.review_id, "source": row.source, "issueType": row.issue_type,
        "description": row.description, "severity": row.severity, "competencyCode": row.competency_code,
        "competencyName": COMPETENCY_NAMES.get(row.competency_code, row.competency_code),
        "improvementRequirement": row.improvement_requirement, "repeatCount": row.repeat_count,
        "firstSeenOn": row.first_seen_on, "lastSeenOn": row.last_seen_on,
        "status": row.status, "verificationMethod": row.verification_method,
    }


def _assessment_payload(row: MonthlyAssessment) -> dict[str, Any]:
    return {
        "id": row.id, "monthlyPlanId": row.monthly_plan_id,
        "employeeSelfReview": row.employee_self_review, "mentorReview": row.mentor_review,
        "completionScore": row.completion_score, "exerciseScore": row.exercise_score,
        "unfamiliarCaseScore": row.unfamiliar_case_score,
        "projectPerformanceScore": row.project_performance_score,
        "initiativeScore": row.initiative_score, "totalScore": row.total_score,
        "competencyChanges": _json(row.competency_changes_json, {}),
        "unresolvedIssues": _json(row.unresolved_issues_json, []),
        "nextMonthSuggestion": row.next_month_suggestion, "status": row.status,
    }


@router.get("/templates")
def list_templates(db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_feature_permission(db, user, "development", "view")
    return [{"code": code, **spec} for code, spec in TEMPLATES.items()]


@router.get("/dashboard")
def dashboard(db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "view")
    employee_query = select(DevelopmentEmployee).where(DevelopmentEmployee.status == "active").order_by(DevelopmentEmployee.code)
    if not _can_manage_training(user):
        employee_query = employee_query.where(DevelopmentEmployee.user_id == user.id)
    employees = db.execute(employee_query).scalars().all()
    tasks = db.execute(select(WeeklyTask)).scalars().all()
    pending_reviews = db.execute(select(func.count(DevelopmentReview.id)).where(DevelopmentReview.status == "pending")).scalar_one()
    open_issues = db.execute(select(DevelopmentIssue).where(DevelopmentIssue.status.in_(["open", "pending_verification"]))).scalars().all()
    completed = sum(1 for task in tasks if task.status == "completed")
    due_tasks = [task for task in tasks if task.due_date]
    on_time = sum(1 for task in due_tasks if task.status == "completed" and (not task.completed_at or task.completed_at.date() <= task.due_date))
    team = []
    today = date.today()
    for employee in employees:
        plans = db.execute(select(DevelopmentPlan).where(DevelopmentPlan.employee_id == employee.id)).scalars().all()
        plan_ids = [row.id for row in plans]
        months = db.execute(select(MonthlyPlan).where(MonthlyPlan.development_plan_id.in_(plan_ids))).scalars().all() if plan_ids else []
        month_ids = [row.id for row in months]
        person_tasks = db.execute(select(WeeklyTask).where(WeeklyTask.monthly_plan_id.in_(month_ids))).scalars().all() if month_ids else []
        person_done = sum(1 for task in person_tasks if task.status == "completed")
        progress = round(sum(task.progress for task in person_tasks) / len(person_tasks)) if person_tasks else 0
        person_issues = [item for item in open_issues if item.employee_id == employee.id]
        repeated = [item for item in person_issues if item.repeat_count >= 2]
        latest_month = max(months, key=lambda item: (item.month_no, item.id), default=None)
        assessment = db.execute(select(MonthlyAssessment).where(MonthlyAssessment.monthly_plan_id == latest_month.id)).scalar_one_or_none() if latest_month else None
        stalled = any(task.status != "completed" and task.due_date and task.due_date <= today - timedelta(days=14) for task in person_tasks)
        team.append({
            "employee": _employee_payload(db, employee), "progress": progress,
            "taskCompleted": person_done, "taskTotal": len(person_tasks),
            "pendingReview": db.execute(select(func.count(DevelopmentReview.id)).where(DevelopmentReview.employee_id == employee.id, DevelopmentReview.status == "pending")).scalar_one(),
            "openIssueCount": len(person_issues), "repeatIssueCount": len(repeated),
            "risk": repeated[0].description if repeated else (person_issues[0].description if person_issues else ""),
            "monthlyScore": round(assessment.total_score, 1) if assessment else None,
            "stalledTwoWeeks": stalled,
        })
    return {
        "summary": {
            "employeeCount": len(employees),
            "completionRate": round(completed * 100 / len(tasks)) if tasks else 0,
            "onTimeRate": round(on_time * 100 / len(due_tasks)) if due_tasks else 0,
            "unfinishedTaskCount": len(tasks) - completed, "pendingReviewCount": pending_reviews,
            "openIssueCount": len(open_issues), "stalledEmployeeCount": sum(1 for row in team if row["stalledTwoWeeks"]),
        },
        "team": team,
    }


@router.get("/employees")
def list_employees(db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_feature_permission(db, user, "development", "view")
    query = select(DevelopmentEmployee).order_by(DevelopmentEmployee.status, DevelopmentEmployee.code)
    if not _can_manage_training(user):
        query = query.where(DevelopmentEmployee.user_id == user.id)
    rows = db.execute(query).scalars().all()
    return [_employee_payload(db, row) for row in rows]


@router.get("/employees/{employee_id}")
def employee_detail(employee_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "view")
    employee = _require(db, DevelopmentEmployee, employee_id, "人员档案")
    if employee.user_id != user.id:
        _require_training_management(user)
    payload = _employee_payload(db, employee, include_detail=True)
    plans = db.execute(select(DevelopmentPlan).where(DevelopmentPlan.employee_id == employee.id).order_by(DevelopmentPlan.id.desc())).scalars().all()
    issues = db.execute(select(DevelopmentIssue).where(DevelopmentIssue.employee_id == employee.id).order_by(DevelopmentIssue.last_seen_on.desc(), DevelopmentIssue.id.desc())).scalars().all()
    payload["plans"] = [_plan_payload(db, row) for row in plans]
    payload["issues"] = [_issue_payload(db, row) for row in issues]
    return payload


@router.post("/employees")
def create_employee(payload: DevelopmentEmployeeIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    if db.execute(select(DevelopmentEmployee).where(DevelopmentEmployee.code == payload.code.strip().upper())).scalar_one_or_none():
        raise HTTPException(status_code=409, detail="人员简称已存在")
    row = DevelopmentEmployee(**payload.model_dump())
    row.code = row.code.strip().upper()
    db.add(row)
    db.commit()
    db.refresh(row)
    return _employee_payload(db, row, include_detail=True)


@router.patch("/employees/{employee_id}")
def patch_employee(employee_id: int, payload: DevelopmentEmployeePatchIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    row = _require(db, DevelopmentEmployee, employee_id, "人员档案")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(row, key, value)
    db.commit()
    return _employee_payload(db, row, include_detail=True)


@router.post("/employees/{employee_id}/competencies")
def add_competency(employee_id: int, payload: CompetencyAssessmentIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, DevelopmentEmployee, employee_id, "人员档案")
    dimension = db.execute(select(CompetencyDimension).where(CompetencyDimension.code == payload.competency_code)).scalar_one_or_none()
    if dimension is None:
        raise HTTPException(status_code=422, detail="能力维度不存在")
    row = EmployeeCompetency(employee_id=employee_id, competency_id=dimension.id, assessor_user_id=user.id, **payload.model_dump(exclude={"competency_code"}))
    db.add(row)
    db.commit()
    return {"id": row.id, "employeeId": employee_id, "code": dimension.code, "score": row.score, "assessedOn": row.assessed_on}


@router.get("/plans")
def list_plans(employee_id: Optional[int] = None, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_feature_permission(db, user, "development", "view")
    query = select(DevelopmentPlan)
    if not _can_manage_training(user):
        employee = db.execute(select(DevelopmentEmployee).where(DevelopmentEmployee.user_id == user.id)).scalar_one_or_none()
        employee_id = employee.id if employee else -1
    if employee_id:
        query = query.where(DevelopmentPlan.employee_id == employee_id)
    rows = db.execute(query.order_by(DevelopmentPlan.employee_id, DevelopmentPlan.id.desc())).scalars().all()
    return [_plan_payload(db, row) for row in rows]


@router.post("/plans")
def create_plan(payload: DevelopmentPlanIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, DevelopmentEmployee, payload.employee_id, "人员档案")
    row = DevelopmentPlan(**payload.model_dump())
    db.add(row)
    db.commit()
    db.refresh(row)
    return _plan_payload(db, row)


@router.post("/monthly-plans")
def create_monthly_plan(payload: MonthlyPlanIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, DevelopmentPlan, payload.development_plan_id, "培养计划")
    row = MonthlyPlan(**payload.model_dump())
    db.add(row)
    db.commit()
    db.refresh(row)
    return _monthly_payload(db, row)


@router.post("/weekly-tasks")
def create_weekly_task(payload: WeeklyTaskIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, MonthlyPlan, payload.monthly_plan_id, "月度计划")
    row = WeeklyTask(**payload.model_dump())
    db.add(row)
    db.commit()
    db.refresh(row)
    return _task_payload(row)


@router.patch("/weekly-tasks/{task_id}")
def patch_weekly_task(task_id: int, payload: WeeklyTaskPatchIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    row = _require(db, WeeklyTask, task_id, "周任务")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(row, key, value)
    if row.status == "completed":
        row.progress = 100
        row.completed_at = row.completed_at or datetime.utcnow()
    db.commit()
    return _task_payload(row)


@router.post("/exercises")
def create_exercise(payload: DevelopmentExerciseIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, WeeklyTask, payload.weekly_task_id, "周任务")
    values = payload.model_dump(exclude={"competency_codes"})
    row = DevelopmentExercise(**values, competency_codes_json=json.dumps(payload.competency_codes, ensure_ascii=False))
    db.add(row)
    db.commit()
    return {"id": row.id, "code": row.code, "topic": row.topic}


@router.post("/submissions")
def create_submission(payload: DevelopmentSubmissionIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "view")
    _require(db, DevelopmentExercise, payload.exercise_id, "练习")
    employee = _require(db, DevelopmentEmployee, payload.employee_id, "人员档案")
    if employee.user_id != user.id and not (user.role and user.role.code in {"admin", "partner", "quality", "director", "senior_manager", "manager"}):
        raise HTTPException(status_code=403, detail="只能提交自己的练习")
    latest = db.execute(select(func.max(DevelopmentExerciseSubmission.attempt_no)).where(DevelopmentExerciseSubmission.exercise_id == payload.exercise_id, DevelopmentExerciseSubmission.employee_id == payload.employee_id)).scalar_one()
    row = DevelopmentExerciseSubmission(
        exercise_id=payload.exercise_id, employee_id=payload.employee_id,
        attempt_no=int(latest or 0) + 1, content=payload.content,
        attachment_reference=payload.attachment_reference, used_ai=payload.used_ai,
        status="submitted" if payload.submit else "draft",
        submitted_at=datetime.utcnow() if payload.submit else None,
    )
    db.add(row)
    db.commit()
    return {"id": row.id, "attemptNo": row.attempt_no, "status": row.status}


@router.post("/reviews")
def create_review(payload: DevelopmentReviewIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    row = DevelopmentReview(**payload.model_dump(), reviewer_user_id=user.id, reviewed_at=datetime.utcnow())
    db.add(row)
    db.commit()
    return {"id": row.id, "status": row.status, "passed": row.passed}


@router.get("/issues")
def list_issues(employee_id: Optional[int] = None, status: Optional[str] = None, db: Session = Depends(get_db), user: User = Depends(current_user)) -> list[dict[str, Any]]:
    ensure_feature_permission(db, user, "development", "view")
    query = select(DevelopmentIssue)
    if not _can_manage_training(user):
        employee = db.execute(select(DevelopmentEmployee).where(DevelopmentEmployee.user_id == user.id)).scalar_one_or_none()
        employee_id = employee.id if employee else -1
    if employee_id:
        query = query.where(DevelopmentIssue.employee_id == employee_id)
    if status:
        query = query.where(DevelopmentIssue.status == status)
    rows = db.execute(query.order_by(DevelopmentIssue.repeat_count.desc(), DevelopmentIssue.last_seen_on.desc(), DevelopmentIssue.id.desc())).scalars().all()
    return [_issue_payload(db, row) for row in rows]


def _fingerprint(competency_code: str, description: str) -> str:
    normalized = re.sub(r"\s+", "", f"{competency_code}|{description}").lower()
    return sha256(normalized.encode("utf-8")).hexdigest()


@router.post("/issues")
def create_issue(payload: DevelopmentIssueIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    _require(db, DevelopmentEmployee, payload.employee_id, "人员档案")
    fingerprint = _fingerprint(payload.competency_code, payload.description)
    existing = db.execute(select(DevelopmentIssue).where(DevelopmentIssue.employee_id == payload.employee_id, DevelopmentIssue.fingerprint == fingerprint).order_by(DevelopmentIssue.id.desc())).scalars().first()
    seen_on = payload.first_seen_on or date.today()
    if existing:
        existing.repeat_count += 1
        existing.last_seen_on = seen_on
        existing.status = "open"
        existing.source = payload.source or existing.source
        existing.improvement_requirement = payload.improvement_requirement or existing.improvement_requirement
        existing.verification_method = payload.verification_method or existing.verification_method
        row = existing
    else:
        row = DevelopmentIssue(**payload.model_dump(), fingerprint=fingerprint, repeat_count=1, first_seen_on=seen_on, last_seen_on=seen_on)
        db.add(row)
    db.commit()
    return _issue_payload(db, row)


@router.patch("/issues/{issue_id}")
def patch_issue(issue_id: int, payload: DevelopmentIssuePatchIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    row = _require(db, DevelopmentIssue, issue_id, "培养问题")
    for key, value in payload.model_dump(exclude_unset=True).items():
        setattr(row, key, value)
    if row.status == "verified":
        row.verified_at = datetime.utcnow()
    db.commit()
    return _issue_payload(db, row)


@router.post("/assessments")
def save_assessment(payload: DevelopmentMonthlyAssessmentIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    month = _require(db, MonthlyPlan, payload.monthly_plan_id, "月度计划")
    row = db.execute(select(MonthlyAssessment).where(MonthlyAssessment.monthly_plan_id == month.id)).scalar_one_or_none()
    if row is None:
        row = MonthlyAssessment(monthly_plan_id=month.id)
        db.add(row)
    values = payload.model_dump(exclude={"monthly_plan_id", "competency_changes"})
    for key, value in values.items():
        setattr(row, key, value)
    row.competency_changes_json = json.dumps(payload.competency_changes, ensure_ascii=False)
    open_issues = db.execute(select(DevelopmentIssue).join(DevelopmentPlan, DevelopmentPlan.employee_id == DevelopmentIssue.employee_id).where(DevelopmentPlan.id == month.development_plan_id, DevelopmentIssue.status.in_(["open", "pending_verification"]))).scalars().all()
    row.unresolved_issues_json = json.dumps([_issue_payload(db, issue) for issue in open_issues], ensure_ascii=False, default=str)
    row.total_score = round(payload.completion_score * 0.15 + payload.exercise_score * 0.25 + payload.unfamiliar_case_score * 0.25 + payload.project_performance_score * 0.25 + payload.initiative_score * 0.10, 1)
    row.assessed_at = datetime.utcnow()
    db.commit()
    return _assessment_payload(row)


def _draft_topics(db: Session, employee: DevelopmentEmployee, template_code: str = "") -> tuple[list[str], str]:
    if template_code and template_code in TEMPLATES:
        return list(TEMPLATES[template_code]["topics"]), TEMPLATES[template_code]["name"]
    latest, _ = _latest_competencies(db, employee.id)
    competency_topic = {
        "financial_audit": "业务财务映射", "itgc": "ITGC", "itac": "ITAC设计",
        "interface_ipe": "Interface/IPE", "data_analysis": "数据需求",
        "system_foundation": "系统技术基础", "communication": "PBC与客户沟通",
        "project_management": "项目交付", "ai_application": "综合案例",
    }
    weakest = sorted(latest.items(), key=lambda item: item[1])[:3]
    topics = [competency_topic[code] for code, _ in weakest if code in competency_topic]
    issues = db.execute(select(DevelopmentIssue).where(DevelopmentIssue.employee_id == employee.id, DevelopmentIssue.status.in_(["open", "pending_verification"])).order_by(DevelopmentIssue.repeat_count.desc(), DevelopmentIssue.id.desc())).scalars().all()
    for issue in issues:
        topic = competency_topic.get(issue.competency_code)
        if topic and topic not in topics:
            topics.insert(0, topic)
    route_topic = "综合案例" if "项目经理" not in employee.direction else "项目交付"
    topics.append(route_topic)
    defaults = ["ITGC", "ITAC设计", "Interface/IPE", "综合案例"]
    for topic in defaults:
        if len(topics) >= 4:
            break
        if topic not in topics:
            topics.append(topic)
    return topics[:4], "实际表现与半年成长路线"


def _create_draft(db: Session, employee: DevelopmentEmployee, payload: DraftPlanIn) -> MonthlyPlan:
    plan = db.execute(select(DevelopmentPlan).where(DevelopmentPlan.employee_id == employee.id, DevelopmentPlan.status.in_(["draft", "active"])).order_by(DevelopmentPlan.id.desc())).scalars().first()
    if plan is None:
        plan = DevelopmentPlan(
            employee_id=employee.id, title=f"{datetime.now().year}年度培养计划",
            plan_type="annual", half_year_goal=employee.half_year_goal,
            year_goal=employee.year_goal, status="draft", source_type="ai_assisted",
        )
        db.add(plan)
        db.flush()
    existing_months = db.execute(select(MonthlyPlan).where(MonthlyPlan.development_plan_id == plan.id).order_by(MonthlyPlan.month_no)).scalars().all()
    month_no = (max((row.month_no for row in existing_months), default=0) + 1)
    topics, source_label = _draft_topics(db, employee, payload.template_code)
    latest, _ = _latest_competencies(db, employee.id)
    low = sorted(latest.items(), key=lambda item: item[1])[:3]
    low_text = "、".join(f"{COMPETENCY_NAMES.get(code, code)} {score:g}分" for code, score in low) or "尚无能力评分"
    issues = db.execute(select(DevelopmentIssue).where(DevelopmentIssue.employee_id == employee.id, DevelopmentIssue.status.in_(["open", "pending_verification"])).order_by(DevelopmentIssue.repeat_count.desc())).scalars().all()
    issue_text = "；".join(f"{item.description} ×{item.repeat_count}" for item in issues[:3]) or "暂无未解决重复问题"
    basis = f"70%实际表现：{low_text}；{issue_text}。30%成长路线：{employee.half_year_goal or employee.direction or '夯实IT审计基本功'}。"
    rationale = f"基于{source_label}生成。优先安排{'、'.join(topics[:3])}，第4周用{topics[3]}验收；管理员确认后再发布。"
    month = MonthlyPlan(
        development_plan_id=plan.id, month_no=month_no,
        title=f"第{month_no}个月培养计划草案",
        month_goal=f"围绕{'、'.join(topics)}形成可验证能力和交付物。",
        generation_basis=basis, rationale=rationale,
        expected_project=payload.expected_project, status="draft", generated_by="ai_assisted_rules",
    )
    db.add(month)
    db.flush()
    for index, topic in enumerate(topics, start=1):
        spec = TOPIC_LIBRARY[topic]
        db.add(WeeklyTask(monthly_plan_id=month.id, week_no=index, topic=topic, learning_content=spec["learning"], exercise_case=spec["exercise"], deliverable=spec["deliverable"], acceptance_criteria=spec["acceptance"]))
    db.flush()
    return month


@router.post("/employees/{employee_id}/generate-plan")
def generate_plan(employee_id: int, payload: DraftPlanIn, db: Session = Depends(get_db), user: User = Depends(current_user)) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    employee = _require(db, DevelopmentEmployee, employee_id, "人员档案")
    month = _create_draft(db, employee, payload)
    db.commit()
    return _monthly_payload(db, month)


def _questionnaire_rows(workbook_bytes: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header_index = next((index for index, row in enumerate(rows) if row and str(row[0] or "").strip() == "模块"), None)
    if header_index is None:
        raise HTTPException(status_code=422, detail="未识别到问卷表头（模块/编号/问题/题型/你的答案）")
    output = []
    for row in rows[header_index + 1:]:
        if not row or not str(row[1] or "").strip():
            continue
        output.append({
            "module": str(row[0] or "").strip(), "code": str(row[1] or "").strip(),
            "question": str(row[2] or "").strip(), "type": str(row[3] or "").strip(),
            "options": str(row[4] or "").strip(), "answer": row[5],
            "note": str(row[6] or "").strip() if len(row) > 6 else "",
        })
    return output


def _derive_scores(rows: list[dict[str, Any]]) -> dict[str, tuple[float, str]]:
    buckets: dict[str, list[tuple[float, str]]] = {code: [] for code in COMPETENCY_NAMES}
    for row in rows:
        question = row["question"]
        if any(word in question for word in ["WorkBuddy的熟练", "AI工具的熟练", "AI应用能力"]):
            match = re.match(r"\s*([A-E])", str(row["answer"] or "").upper())
            if match:
                buckets["ai_application"].append((float(ord(match.group(1)) - 64), question))
        if "评分" not in row["type"] or not isinstance(row["answer"], (int, float)):
            continue
        score = float(row["answer"])
        codes = []
        if any(word in question for word in ["会计", "财务报表", "财审", "业务流程", "审计证据"]): codes.append("financial_audit")
        if any(word in question for word in ["用户权限", "职责分离", "程序变更", "系统运维", "日志", "备份", "ITGC"]): codes.append("itgc")
        if "ITAC" in question or "自动控制" in question: codes.append("itac")
        if any(word in question for word in ["接口", "Interface", "IPE", "系统报表", "关键报表"]): codes.append("interface_ipe")
        if any(word in question for word in ["Excel", "SQL", "Python", "数据分析", "数据库表结构"]): codes.append("data_analysis")
        if any(word in question for word in ["应用系统", "操作系统", "数据库", "技术"]): codes.append("system_foundation")
        if any(word in question for word in ["客户", "访谈", "资料需求", "汇报"]): codes.append("communication")
        if any(word in question for word in ["多任务", "项目推进", "项目管理"]): codes.append("project_management")
        for code in set(codes): buckets[code].append((score, question))
    output = {}
    for code, items in buckets.items():
        if items:
            output[code] = (round(sum(score for score, _ in items) / len(items), 1), "；".join(question for _, question in items))
    return output


@router.post("/questionnaire-import")
async def import_questionnaire(
    code: str = Form(...), name: str = Form(...), current_role: str = Form(""),
    direction: str = Form(""), pasted_text: str = Form(""),
    file: Optional[UploadFile] = File(None),
    db: Session = Depends(get_db), user: User = Depends(current_user),
) -> dict[str, Any]:
    ensure_feature_permission(db, user, "development", "edit")
    _require_training_management(user)
    rows: list[dict[str, Any]] = []
    source_reference = "粘贴调研结果"
    if file and file.filename:
        if not file.filename.lower().endswith(".xlsx"):
            raise HTTPException(status_code=422, detail="当前仅支持 .xlsx 调研问卷")
        rows = _questionnaire_rows(await file.read())
        source_reference = file.filename
    elif pasted_text.strip():
        try:
            parsed = json.loads(pasted_text)
            rows = parsed if isinstance(parsed, list) else [{"answer": parsed, "question": "粘贴调研结果", "type": "开放题", "code": "PASTE", "module": "导入", "options": "", "note": ""}]
        except json.JSONDecodeError:
            rows = [{"answer": pasted_text.strip(), "question": "粘贴调研结果", "type": "开放题", "code": "PASTE", "module": "导入", "options": "", "note": ""}]
    else:
        raise HTTPException(status_code=422, detail="请上传问卷或粘贴调研结果")
    normalized_code = code.strip().upper()
    employee = db.execute(select(DevelopmentEmployee).where(DevelopmentEmployee.code == normalized_code)).scalar_one_or_none()
    if employee is None:
        employee = DevelopmentEmployee(code=normalized_code, name=name.strip(), current_role=current_role, direction=direction)
        db.add(employee)
        db.flush()
    else:
        employee.name = name.strip() or employee.name
        employee.current_role = current_role or employee.current_role
        employee.direction = direction or employee.direction
    employee.questionnaire_json = json.dumps({"answers": rows}, ensure_ascii=False, default=str)
    employee.source_reference = source_reference
    scores = _derive_scores(rows)
    dimensions = {row.code: row for row in db.execute(select(CompetencyDimension)).scalars().all()}
    for competency_code, (score, note) in scores.items():
        dimension = dimensions.get(competency_code)
        if dimension:
            db.add(EmployeeCompetency(employee_id=employee.id, competency_id=dimension.id, score=score, assessed_on=date.today(), source="questionnaire", note=note, assessor_user_id=user.id))
    db.commit()
    return _employee_payload(db, employee, include_detail=True)
