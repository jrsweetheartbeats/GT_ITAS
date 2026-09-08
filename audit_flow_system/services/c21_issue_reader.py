from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Optional

from docx import Document
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.utils import obj_dict
from ..models import Client, Project, Workpaper


ISSUE_HEADERS = {
    "issue_no": ("编号", "序号", "问题编号", "发现编号"),
    "title": ("问题", "问题标题", "发现", "发现事项", "事项"),
    "description": ("问题描述", "发现描述", "具体问题", "缺陷描述", "情况描述", "内容"),
    "severity": ("严重程度", "风险等级", "重要程度", "等级"),
    "control_code": ("控制点", "控制编号", "控制程序", "索引", "底稿索引"),
    "recommendation": ("建议", "整改建议", "管理层建议", "改进建议"),
    "owner": ("责任人", "负责人", "整改责任人", "部门"),
    "status": ("状态", "整改状态", "处理状态", "是否解决"),
}


def normalize_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def compact_value(value: Any, limit: int = 1200) -> str:
    text = str(value or "").strip()
    return text if len(text) <= limit else text[:limit] + "..."


def find_c21_workpaper(workpapers: list[Workpaper]) -> Optional[Workpaper]:
    for wp in workpapers:
        code = normalize_text(wp.code)
        name = normalize_text(wp.name)
        if "c21-1" in code or "c21-1" in name or "问题发现" in name or "审计发现" in name:
            return wp
    return None


def header_mapping(values: list[Any]) -> dict[str, int]:
    normalized = [normalize_text(value) for value in values]
    mapping: dict[str, int] = {}
    for field, aliases in ISSUE_HEADERS.items():
        for idx, text in enumerate(normalized):
            if not text:
                continue
            if any(normalize_text(alias) in text for alias in aliases):
                mapping[field] = idx
                break
    return mapping


def detect_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    best_row = -1
    best_mapping: dict[str, int] = {}
    for idx, row in enumerate(rows[:80]):
        mapping = header_mapping(row)
        score = len(mapping)
        has_issue_field = any(field in mapping for field in ("title", "description"))
        if score > len(best_mapping) and has_issue_field:
            best_row = idx
            best_mapping = mapping
    return best_row, best_mapping


def issue_from_row(row: list[Any], mapping: dict[str, int], row_no: int, source: str) -> Optional[dict[str, Any]]:
    raw_values = [compact_value(value, 300) for value in row]
    if not any(raw_values):
        return None
    item = {
        field: compact_value(row[col]) if col < len(row) else ""
        for field, col in mapping.items()
    }
    issue_text = " ".join(item.get(field, "") for field in ("title", "description", "control_code")).strip()
    if not issue_text:
        return None
    item.update(
        {
            "source": source,
            "row_no": row_no,
            "raw_values": raw_values,
        }
    )
    return item


def issues_from_rows(rows: list[list[Any]], source: str) -> tuple[list[dict[str, Any]], str]:
    header_idx, mapping = detect_header(rows)
    if header_idx < 0:
        return [], "未识别问题发现清单表头"
    issues: list[dict[str, Any]] = []
    for offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        issue = issue_from_row(row, mapping, offset, source)
        if issue is not None:
            issues.append(issue)
    if not issues:
        return [], "已识别表头，但未读取到问题记录"
    return issues, ""


def read_excel_issues(path: Path) -> tuple[list[dict[str, Any]], str]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        all_issues: list[dict[str, Any]] = []
        errors: list[str] = []
        for ws in wb.worksheets:
            rows = [list(row) for row in ws.iter_rows(values_only=True)]
            issues, error = issues_from_rows(rows, ws.title)
            if issues:
                all_issues.extend(issues)
            elif error:
                errors.append(f"{ws.title}: {error}")
        if all_issues:
            return all_issues, ""
        return [], "；".join(errors) or "未读取到问题记录"
    finally:
        wb.close()


def read_docx_issues(path: Path) -> tuple[list[dict[str, Any]], str]:
    doc = Document(path)
    all_issues: list[dict[str, Any]] = []
    errors: list[str] = []
    for idx, table in enumerate(doc.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        issues, error = issues_from_rows(rows, f"table-{idx}")
        if issues:
            all_issues.extend(issues)
        elif error:
            errors.append(f"table-{idx}: {error}")
    if all_issues:
        return all_issues, ""
    paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    if paragraphs:
        return [], "Word 文件未识别到结构化问题表格"
    return [], "Word 文件为空或无可读取内容"


def read_c21_issue_file(path: Path) -> tuple[list[dict[str, Any]], str]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_excel_issues(path)
    if suffix == ".docx":
        return read_docx_issues(path)
    return [], f"不支持的 C21-1 文件类型：{suffix or '无扩展名'}"


def project_c21_issues(project: Project, workpapers: list[Workpaper]) -> dict[str, Any]:
    wp = find_c21_workpaper(workpapers)
    result: dict[str, Any] = {
        "project_id": project.id,
        "project_name": project.name,
        "workpaper_id": wp.id if wp else None,
        "workpaper_code": wp.code if wp else "",
        "workpaper_name": wp.name if wp else "",
        "file_path": wp.file_path if wp else "",
        "status": "missing",
        "error": "",
        "issues": [],
    }
    if wp is None:
        result["error"] = "项目未登记 C21-1 问题发现清单底稿"
        return result
    if not wp.file_path:
        result["status"] = "error"
        result["error"] = "C21-1 底稿未维护文件路径"
        return result
    path = Path(wp.file_path).expanduser()
    if not path.exists() or not path.is_file():
        result["status"] = "error"
        result["error"] = f"C21-1 底稿文件不存在：{path}"
        return result
    try:
        issues, error = read_c21_issue_file(path)
    except Exception as exc:
        result["status"] = "error"
        result["error"] = f"C21-1 底稿读取失败：{exc}"
        return result
    result["issues"] = issues
    result["status"] = "ok" if not error else "error"
    result["error"] = error
    return result


def client_c21_issues(db: Session, client: Client, projects: Optional[list[Project]] = None) -> dict[str, Any]:
    if projects is None:
        projects = db.execute(select(Project).where(Project.client_id == client.id).order_by(Project.id.desc())).scalars().all()
    project_results: list[dict[str, Any]] = []
    all_issues: list[dict[str, Any]] = []
    for project in projects:
        workpapers = db.execute(
            select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.stage, Workpaper.code)
        ).scalars().all()
        project_result = project_c21_issues(project, workpapers)
        project_results.append(project_result)
        for issue in project_result["issues"]:
            all_issues.append(
                {
                    **issue,
                    "project_id": project.id,
                    "project_name": project.name,
                    "workpaper_id": project_result["workpaper_id"],
                    "workpaper_code": project_result["workpaper_code"],
                }
            )
    return {
        **obj_dict(client),
        "project_count": len(projects),
        "issue_count": len(all_issues),
        "issues": all_issues,
        "projects": project_results,
    }
