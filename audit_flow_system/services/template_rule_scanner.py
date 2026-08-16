from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MAX_SCAN_ROWS = 180
MAX_SCAN_COLS = 80
MAX_FIELD_RULES_PER_WORKPAPER = 80
MAX_TABLE_RULES_PER_WORKPAPER = 35
MAX_PROCEDURES_PER_WORKPAPER = 60

UTILITY_SHEET_KEYWORDS = {"gt_custom", "acerno_cache", "cache"}

LABEL_KEYWORDS = [
    "客户名称",
    "被审计单位",
    "单位名称",
    "项目编号",
    "OA项目编号",
    "IMS项目编号",
    "审计期间",
    "会计期间",
    "财务报表截止日",
    "截止日",
    "编制人",
    "复核人",
    "日期",
    "索引号",
    "页次",
    "第一签字合伙人",
    "二签",
    "合伙人",
    "项目负责人",
    "项目负责经理",
    "项目现场负责人",
    "IT团队负责人",
    "IT合伙人",
    "IT复核人",
    "数据来源",
    "数据获取方式",
    "审计目的",
    "审计过程",
    "审计结论",
    "测试目的",
    "测试过程",
    "测试结果",
    "客户反馈",
]

GENERIC_NO_COLON_LABELS = {"项目", "日期"}

HEADER_KEYWORDS = [
    "序号",
    "程序",
    "控制",
    "编号",
    "目的",
    "测试",
    "过程",
    "结果",
    "结论",
    "样本",
    "总体",
    "索引",
    "执行人",
    "复核人",
    "问题",
    "缺陷",
    "风险",
    "影响",
    "涉及应用程序",
    "业务流程",
    "财务报表项目",
    "认定",
    "备注",
]

PROCEDURE_KEYWORDS = [
    "设计有效性",
    "执行有效性",
    "标准程序",
    "审计程序",
    "测试程序",
    "测试目的",
    "测试过程",
    "计划测试过程",
    "控制测试",
    "穿行测试",
    "检查",
    "访谈",
    "观察",
    "重新执行",
    "抽样",
    "样本",
    "总体",
    "结论",
]

CONTROL_CODE_RE = re.compile(r"\b(?:C22\.)?(?:SA|PE|PM|NS)-?\d+[A-Za-z]?(?:\.\d+)?\b", re.I)


def compact(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def clean_text(value: Any, limit: int = 300) -> str:
    text = re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()
    return text[:limit] + "..." if len(text) > limit else text


def non_empty(values: list[Any]) -> list[str]:
    return [clean_text(value) for value in values if clean_text(value)]


def looks_like_utility_sheet(title: str) -> bool:
    normalized = compact(title)
    return any(keyword in normalized for keyword in UTILITY_SHEET_KEYWORDS)


def label_name(text: str) -> str:
    value = clean_text(text, 120)
    if "：" in value:
        return value.split("：", 1)[0].strip() + "："
    if ":" in value:
        return value.split(":", 1)[0].strip() + ":"
    for keyword in LABEL_KEYWORDS:
        if keyword in value:
            return keyword
    return value


def embedded_value(text: str) -> str:
    value = clean_text(text, 300)
    for delimiter in ("：", ":"):
        if delimiter in value:
            return value.split(delimiter, 1)[1].strip()
    return ""


def is_candidate_label(text: str) -> bool:
    value = clean_text(text, 160)
    normalized = compact(value)
    if not normalized or len(normalized) < 2 or len(normalized) > 80:
        return False
    if value.startswith("="):
        return False
    if re.match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:\s+\d{1,2}:?\d{0,2}:?\d{0,2})?$", value):
        return False
    if "：" in value or ":" in value:
        left = label_name(value)
        if compact(left).isdigit():
            return False
        return 1 < len(compact(left)) <= 40
    return any(
        compact(keyword) in normalized
        for keyword in LABEL_KEYWORDS
        if keyword not in GENERIC_NO_COLON_LABELS and len(compact(keyword)) >= 3
    )


def source_guess(label: str, context: str = "") -> str:
    label_text = compact(label)
    text = compact(label + context)
    if any(keyword in label_text for keyword in ["编制人", "复核人", "项目负责人", "经理", "现场负责人", "it团队负责人", "it合伙人"]):
        return "project_members"
    if any(keyword in label_text for keyword in ["编制日期", "复核日期", "日期", "索引号", "页次"]):
        return "workpaper_metadata"
    if any(keyword in label_text for keyword in ["客户名称", "被审计单位", "单位名称", "行业", "现场地址"]):
        return "client_profile"
    if any(keyword in label_text for keyword in ["oa项目编号", "ims项目编号", "项目编号", "审计期间", "会计期间", "截止日"]):
        return "project_context"
    if any(keyword in text for keyword in ["数据来源", "获取方式", "样本", "总体", "测试过程", "测试结果", "客户反馈"]):
        return "attachments_and_test_evidence"
    if any(keyword in text for keyword in ["审计结论", "问题", "缺陷", "风险", "影响"]):
        return "workpaper_conclusion_and_review"
    return "attachments_or_manual_judgment"


def confidence_for(locator_method: str, old_value: str, context: str) -> str:
    if locator_method in {"same_row_right", "table_header"}:
        return "high"
    if old_value or any(keyword in compact(context) for keyword in ["编制人", "复核人", "日期", "索引号"]):
        return "medium"
    return "low"


def control_codes(text: str) -> list[str]:
    result: list[str] = []
    for match in CONTROL_CODE_RE.findall(text or ""):
        code = match.upper().replace("C22.", "")
        if code not in result:
            result.append(code)
    return result


def procedure_type(text: str) -> str:
    value = compact(text)
    if "设计有效性" in value:
        return "design_effectiveness"
    if "执行有效性" in value or "运行有效性" in value:
        return "operating_effectiveness"
    if "穿行" in value:
        return "walkthrough"
    if "结论" in value:
        return "conclusion"
    return "procedure"


def row_header_score(values: list[str]) -> int:
    joined = compact(" ".join(values))
    score = 0
    for keyword in HEADER_KEYWORDS:
        if compact(keyword) in joined:
            score += 1
    return score


def is_header_row(values: list[str], joined: str) -> bool:
    if len(values) < 3 or row_header_score(values) < 2:
        return False
    normalized = compact(joined)
    if CONTROL_CODE_RE.search(joined) and "控制编号" not in joined and "信息处理控制索引号" not in joined:
        return False
    if re.search(r"\bITGC#?\d+\b", joined, flags=re.I):
        return False
    if re.search(r"\b(?:B23-15|C26)-\d+\b", joined, flags=re.I) and "索引号" not in joined:
        return False
    long_cells = sum(1 for value in values if len(value) > 90)
    return long_cells <= max(1, len(values) // 4)


def excel_cell_ref(row_idx: int, col_idx: int) -> str:
    return f"{get_column_letter(col_idx)}{row_idx}"


def target_cell_for_label(row_values: list[Any], row_idx: int, col_idx: int, text: str) -> tuple[str, str, str]:
    inline_value = embedded_value(text)
    if inline_value:
        return excel_cell_ref(row_idx, col_idx), "same_cell_after_label", inline_value
    for next_col in range(col_idx + 1, min(len(row_values), col_idx + 8) + 1):
        value = clean_text(row_values[next_col - 1], 180)
        if value and compact(value) not in {compact(label_name(text)), compact(text)}:
            return excel_cell_ref(row_idx, next_col), "same_row_right", value
    return excel_cell_ref(row_idx, col_idx + 1), "same_row_right", ""


def append_limited(items: list[dict[str, Any]], item: dict[str, Any], limit: int) -> None:
    if len(items) < limit:
        items.append(item)


def scan_excel_workpaper(path: Path, workpaper: dict[str, Any]) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    fields: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    procedures: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []
    try:
        for ws in wb.worksheets:
            sheet_summary = {
                "sheet": ws.title,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "utility": looks_like_utility_sheet(ws.title),
                "field_count": 0,
                "table_count": 0,
                "procedure_count": 0,
            }
            max_col = min(ws.max_column or MAX_SCAN_COLS, MAX_SCAN_COLS)
            for row_idx, row in enumerate(ws.iter_rows(max_row=MAX_SCAN_ROWS, max_col=max_col, values_only=True), start=1):
                values = list(row)
                texts = [clean_text(value, 220) for value in values]
                row_texts = non_empty(values)
                if not row_texts:
                    continue
                joined = " | ".join(row_texts)
                if not sheet_summary["utility"]:
                    for col_idx, text in enumerate(texts, start=1):
                        if not is_candidate_label(text):
                            continue
                        label = label_name(text)
                        target_cell, write_mode, old_value = target_cell_for_label(values, row_idx, col_idx, text)
                        append_limited(
                            fields,
                            {
                                "rule_type": "field",
                                "workpaper_code": workpaper.get("code", ""),
                                "workpaper_name": workpaper.get("name", ""),
                                "sheet": ws.title,
                                "label": label,
                                "field_key": compact(label),
                                "label_cell": excel_cell_ref(row_idx, col_idx),
                                "target_cell": target_cell,
                                "locator_method": "label_match",
                                "write_mode": write_mode,
                                "current_value": old_value,
                                "source_guess": source_guess(label, joined),
                                "confidence": confidence_for(write_mode, old_value, joined),
                                "reason": "模板标签识别",
                            },
                            MAX_FIELD_RULES_PER_WORKPAPER,
                        )
                        sheet_summary["field_count"] += 1
                header_values = row_texts[:]
                if not sheet_summary["utility"] and is_header_row(header_values, joined):
                    append_limited(
                        tables,
                        {
                            "rule_type": "table",
                            "workpaper_code": workpaper.get("code", ""),
                            "workpaper_name": workpaper.get("name", ""),
                            "sheet": ws.title,
                            "header_row": row_idx,
                            "headers": header_values[:18],
                            "locator_method": "header_match",
                            "write_mode": "append_or_update_rows",
                            "source_guess": source_guess(" ".join(header_values), joined),
                            "confidence": "high",
                            "reason": "模板表头识别",
                        },
                        MAX_TABLE_RULES_PER_WORKPAPER,
                    )
                    sheet_summary["table_count"] += 1
                if any(compact(keyword) in compact(joined) for keyword in PROCEDURE_KEYWORDS):
                    append_limited(
                        procedures,
                        {
                            "rule_type": "procedure",
                            "workpaper_code": workpaper.get("code", ""),
                            "workpaper_name": workpaper.get("name", ""),
                            "sheet": ws.title,
                            "row": row_idx,
                            "procedure_type": procedure_type(joined),
                            "control_codes": control_codes(joined),
                            "text": clean_text(joined, 500),
                            "source_guess": source_guess(joined, joined),
                            "confidence": "medium",
                            "reason": "模板设计/测试程序识别",
                        },
                        MAX_PROCEDURES_PER_WORKPAPER,
                    )
                    sheet_summary["procedure_count"] += 1
            sheets.append(sheet_summary)
        return {"sheets": sheets, "fields": fields, "tables": tables, "procedures": procedures}
    finally:
        wb.close()


def scan_docx_workpaper(path: Path, workpaper: dict[str, Any]) -> dict[str, Any]:
    doc = Document(path)
    fields: list[dict[str, Any]] = []
    tables: list[dict[str, Any]] = []
    procedures: list[dict[str, Any]] = []
    for idx, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean_text(paragraph.text, 500)
        if not text:
            continue
        if is_candidate_label(text):
            label = label_name(text)
            append_limited(
                fields,
                {
                    "rule_type": "field",
                    "workpaper_code": workpaper.get("code", ""),
                    "workpaper_name": workpaper.get("name", ""),
                    "section": f"paragraph:{idx}",
                    "label": label,
                    "field_key": compact(label),
                    "target_cell": f"paragraph:{idx}",
                    "locator_method": "document_label_match",
                    "write_mode": "replace_paragraph_or_inline_value",
                    "current_value": embedded_value(text),
                    "source_guess": source_guess(label, text),
                    "confidence": "medium",
                    "reason": "Word段落标签识别",
                },
                MAX_FIELD_RULES_PER_WORKPAPER,
            )
        if any(compact(keyword) in compact(text) for keyword in PROCEDURE_KEYWORDS):
            append_limited(
                procedures,
                {
                    "rule_type": "procedure",
                    "workpaper_code": workpaper.get("code", ""),
                    "workpaper_name": workpaper.get("name", ""),
                    "section": f"paragraph:{idx}",
                    "procedure_type": procedure_type(text),
                    "control_codes": control_codes(text),
                    "text": text,
                    "source_guess": source_guess(text, text),
                    "confidence": "medium",
                    "reason": "Word设计/测试程序识别",
                },
                MAX_PROCEDURES_PER_WORKPAPER,
            )
    for table_idx, table in enumerate(doc.tables, start=1):
        for row_idx, row in enumerate(table.rows[:25], start=1):
            values = [clean_text(cell.text, 220) for cell in row.cells]
            headers = non_empty(values)
            if is_header_row(headers, " | ".join(headers)):
                append_limited(
                    tables,
                    {
                        "rule_type": "table",
                        "workpaper_code": workpaper.get("code", ""),
                        "workpaper_name": workpaper.get("name", ""),
                        "section": f"table:{table_idx}:row:{row_idx}",
                        "headers": headers[:18],
                        "locator_method": "document_table_header_match",
                        "write_mode": "append_or_update_rows",
                        "source_guess": source_guess(" ".join(headers), " ".join(headers)),
                        "confidence": "high",
                        "reason": "Word表格表头识别",
                    },
                    MAX_TABLE_RULES_PER_WORKPAPER,
                )
            for col_idx, text in enumerate(values, start=1):
                if not is_candidate_label(text):
                    continue
                label = label_name(text)
                append_limited(
                    fields,
                    {
                        "rule_type": "field",
                        "workpaper_code": workpaper.get("code", ""),
                        "workpaper_name": workpaper.get("name", ""),
                        "section": f"table:{table_idx}:r{row_idx}:c{col_idx}",
                        "label": label,
                        "field_key": compact(label),
                        "target_cell": f"table:{table_idx}:r{row_idx}:c{col_idx}",
                        "locator_method": "document_table_label_match",
                        "write_mode": "same_or_adjacent_cell",
                        "current_value": embedded_value(text),
                        "source_guess": source_guess(label, text),
                        "confidence": "medium",
                        "reason": "Word表格标签识别",
                    },
                    MAX_FIELD_RULES_PER_WORKPAPER,
                )
    return {
        "sheets": [],
        "fields": fields,
        "tables": tables,
        "procedures": procedures,
    }


def rule_candidates(fields: list[dict[str, Any]], tables: list[dict[str, Any]], procedures: list[dict[str, Any]]) -> list[dict[str, Any]]:
    candidates: list[dict[str, Any]] = []
    for item in fields:
        candidates.append(
            {
                "rule_type": "field",
                "label": item.get("label", ""),
                "location": item.get("target_cell") or item.get("section") or "",
                "locator": item.get("label_cell") or item.get("section") or item.get("label", ""),
                "write_mode": item.get("write_mode", ""),
                "source_guess": item.get("source_guess", ""),
                "confidence": item.get("confidence", ""),
                "reason": item.get("reason", ""),
            }
        )
    for item in tables:
        candidates.append(
            {
                "rule_type": "table",
                "label": " / ".join(item.get("headers", [])[:5]),
                "location": item.get("sheet") or item.get("section", ""),
                "locator": "headers:" + " | ".join(item.get("headers", [])[:8]),
                "write_mode": item.get("write_mode", ""),
                "source_guess": item.get("source_guess", ""),
                "confidence": item.get("confidence", ""),
                "reason": item.get("reason", ""),
            }
        )
    for item in procedures:
        candidates.append(
            {
                "rule_type": "procedure",
                "label": item.get("procedure_type", ""),
                "location": item.get("sheet") or item.get("section", ""),
                "locator": "control:" + "、".join(item.get("control_codes", [])) if item.get("control_codes") else item.get("location", ""),
                "write_mode": "generate_content_from_procedure_and_evidence",
                "source_guess": item.get("source_guess", ""),
                "confidence": item.get("confidence", ""),
                "reason": item.get("reason", ""),
                "procedure_text": item.get("text", ""),
            }
        )
    return candidates[:140]


def scan_one_workpaper(workpaper: dict[str, Any]) -> dict[str, Any]:
    path = Path(str(workpaper.get("file_path") or "")).expanduser()
    result = {
        "id": workpaper.get("id"),
        "code": workpaper.get("code", ""),
        "name": workpaper.get("name", ""),
        "stage": workpaper.get("stage", ""),
        "file_path": str(path) if str(path) != "." else "",
        "file_type": path.suffix.lower().lstrip("."),
        "file_exists": path.exists() and path.is_file(),
        "error": "",
        "sheets": [],
        "fields": [],
        "tables": [],
        "procedures": [],
        "candidate_rules": [],
        "summary": {},
    }
    if not result["file_path"]:
        result["error"] = "底稿未维护文件路径"
        return result
    if not result["file_exists"]:
        result["error"] = "底稿文件不存在"
        return result
    try:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            scanned = scan_excel_workpaper(path, workpaper)
        elif path.suffix.lower() == ".docx":
            scanned = scan_docx_workpaper(path, workpaper)
        else:
            result["error"] = f"暂不支持模板规则识别的文件类型：{path.suffix or '无扩展名'}"
            return result
    except Exception as exc:
        result["error"] = f"模板规则识别失败：{exc}"
        return result
    result.update(scanned)
    result["candidate_rules"] = rule_candidates(result["fields"], result["tables"], result["procedures"])
    result["summary"] = {
        "field_count": len(result["fields"]),
        "table_count": len(result["tables"]),
        "procedure_count": len(result["procedures"]),
        "candidate_rule_count": len(result["candidate_rules"]),
    }
    return result


def scan_workpaper_templates(workpapers: list[dict[str, Any]]) -> dict[str, Any]:
    scanned: list[dict[str, Any]] = []
    for workpaper in workpapers:
        scanned.append(scan_one_workpaper(workpaper))
    by_code: dict[str, dict[str, Any]] = {}
    for item in scanned:
        code = str(item.get("code") or "")
        summary = by_code.setdefault(
            code,
            {
                "code": code,
                "name": item.get("name", ""),
                "workpaper_count": 0,
                "field_count": 0,
                "table_count": 0,
                "procedure_count": 0,
                "candidate_rule_count": 0,
                "errors": [],
            },
        )
        summary["workpaper_count"] += 1
        item_summary = item.get("summary") or {}
        summary["field_count"] += int(item_summary.get("field_count") or 0)
        summary["table_count"] += int(item_summary.get("table_count") or 0)
        summary["procedure_count"] += int(item_summary.get("procedure_count") or 0)
        summary["candidate_rule_count"] += int(item_summary.get("candidate_rule_count") or 0)
        if item.get("error"):
            summary["errors"].append(item["error"])
    return {
        "summary": {
            "workpaper_count": len(scanned),
            "scanned_workpaper_count": sum(1 for item in scanned if not item.get("error")),
            "error_count": sum(1 for item in scanned if item.get("error")),
            "field_count": sum((item.get("summary") or {}).get("field_count", 0) for item in scanned),
            "table_count": sum((item.get("summary") or {}).get("table_count", 0) for item in scanned),
            "procedure_count": sum((item.get("summary") or {}).get("procedure_count", 0) for item in scanned),
            "candidate_rule_count": sum((item.get("summary") or {}).get("candidate_rule_count", 0) for item in scanned),
        },
        "by_code": list(by_code.values()),
        "workpapers": scanned,
    }
