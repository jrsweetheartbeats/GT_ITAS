from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
import hashlib
import json
import re
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.db import SessionLocal
from ..core.utils import list_dict
from ..models import Attachment, Project, RuleManualCorrection, RuleManualCorrectionImport, User, Workpaper
from .autofill_rule_inspector import inspect_project_autofill_rules
from .rule_exporter import MANUAL_CORRECTION_HEADER


FORMAL_SHEET = "正式规则"
CANDIDATE_SHEET = "模板识别候选"
VALID_STATUSES = {"draft", "reviewed", "approved", "rejected", "applied"}
IMPORT_STATUSES = {"pending", "running", "completed", "failed"}


def clean_text(value: Any, limit: int | None = None) -> str:
    text = re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()
    if limit and len(text) > limit:
        return text[:limit]
    return text


def normalize_key_part(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def stable_hash(*parts: Any) -> str:
    normalized = "\u001f".join(normalize_key_part(part) for part in parts)
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def formal_correction_key(rule_id: Any, workpaper_code: Any, target_field: Any, locator: Any) -> str:
    return stable_hash("formal", rule_id, workpaper_code, target_field, locator)


def candidate_locator_key(
    workpaper_code: Any,
    candidate_type: Any,
    label_or_procedure: Any,
    location: Any,
    locator: Any,
) -> str:
    return stable_hash("candidate", workpaper_code, candidate_type, label_or_procedure, location, locator)


def _json_text(payload: Any) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True)


def _parse_json_text(value: str) -> Any:
    text = str(value or "").strip()
    if not text or text[0] not in "[{":
        return None
    return json.loads(text)


def normalize_manual_correction(value: str) -> tuple[str, str]:
    text = clean_text(value)
    payload: dict[str, Any] = {"manual_correction": text}
    conflict = ""
    if not text:
        return _json_text(payload), conflict
    try:
        parsed = _parse_json_text(text)
    except json.JSONDecodeError as exc:
        conflict = f"修正内容疑似 JSON 但无法解析：{exc.msg}"
        parsed = None
    if parsed is not None:
        payload["parsed"] = parsed
    return _json_text(payload), conflict


def iter_manual_rows(wb) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for sheet_name, rule_kind in [(FORMAL_SHEET, "formal"), (CANDIDATE_SHEET, "candidate")]:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        row_iter = ws.iter_rows(values_only=True)
        headers = [clean_text(value) for value in next(row_iter, [])]
        if MANUAL_CORRECTION_HEADER not in headers:
            continue
        for row_index, values in enumerate(row_iter, start=2):
            item = {
                headers[index]: values[index] if index < len(values) else None
                for index in range(len(headers))
                if headers[index]
            }
            correction = clean_text(item.get(MANUAL_CORRECTION_HEADER))
            if not correction:
                continue
            item["_sheet"] = sheet_name
            item["_row_no"] = row_index
            item["_rule_kind"] = rule_kind
            item["_manual_correction"] = correction
            rows.append(item)
    return rows


def formal_target_key(rule: dict[str, Any], target: dict[str, Any]) -> str:
    locator = target.get("locator_display") or target.get("locator") or target.get("target_cell") or ""
    return formal_correction_key(rule.get("rule_id") or rule.get("id"), rule.get("workpaper_code"), target.get("target_field") or target.get("field"), locator)


def build_formal_index(payload: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], set[str]]:
    index: dict[str, dict[str, Any]] = {}
    rule_ids: set[str] = set()
    for rule in payload.get("rules") or []:
        rule_id = str(rule.get("rule_id") or rule.get("id") or "")
        if rule_id:
            rule_ids.add(rule_id)
        for target in rule.get("targets") or [{}]:
            key = formal_target_key(rule, target)
            index[key] = {
                "rule": rule,
                "target": target,
                "payload": {
                    "rule_id": rule_id,
                    "rule_name": rule.get("name", ""),
                    "workpaper_code": rule.get("workpaper_code", ""),
                    "workpaper_name": rule.get("workpaper_name", ""),
                    "sheet_or_section": rule.get("sheet") or rule.get("section") or "",
                    "target_field": target.get("target_field") or target.get("field") or "",
                    "locator": target.get("locator_display") or target.get("locator") or target.get("target_cell") or "",
                    "source_type": rule.get("source_type", ""),
                },
            }
    return index, rule_ids


def build_candidate_index(payload: dict[str, Any]) -> tuple[dict[str, dict[str, Any]], set[str]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    workpaper_codes: set[str] = set()
    template_scan = payload.get("template_scan") or {}
    for workpaper in template_scan.get("workpapers") or []:
        if not isinstance(workpaper, dict):
            continue
        code = str(workpaper.get("code") or "")
        if code:
            workpaper_codes.add(code)
        for candidate in workpaper.get("candidate_rules") or []:
            if not isinstance(candidate, dict):
                continue
            label = candidate.get("label") or candidate.get("procedure_text") or ""
            key = candidate_locator_key(code, candidate.get("rule_type"), label, candidate.get("location"), candidate.get("locator"))
            grouped[key].append(
                {
                    "workpaper": workpaper,
                    "candidate": candidate,
                    "payload": {
                        "workpaper_code": code,
                        "workpaper_name": workpaper.get("name", ""),
                        "candidate_type": candidate.get("rule_type", ""),
                        "label_or_procedure": label,
                        "location": candidate.get("location", ""),
                        "locator": candidate.get("locator", ""),
                        "source_type": candidate.get("source_guess", ""),
                    },
                }
            )
    return {key: items[0] for key, items in grouped.items() if len(items) == 1}, {key for key, items in grouped.items() if len(items) > 1}


def _existing_correction(db: Session, project_id: int, rule_kind: str, correction_key: str) -> RuleManualCorrection | None:
    return db.execute(
        select(RuleManualCorrection).where(
            RuleManualCorrection.project_id == project_id,
            RuleManualCorrection.rule_kind == rule_kind,
            RuleManualCorrection.correction_key == correction_key,
        )
    ).scalar_one_or_none()


def _upsert_correction(
    db: Session,
    *,
    project_id: int,
    user: User,
    rule_kind: str,
    correction_key: str,
    candidate_key: str = "",
    rule_id: str = "",
    workpaper_code: str = "",
    workpaper_name: str = "",
    sheet_or_section: str = "",
    target_field: str = "",
    locator: str = "",
    source_type: str = "",
    original_payload: dict[str, Any] | None = None,
    manual_correction: str,
    conflict_message: str = "",
) -> str:
    normalized_payload, parse_conflict = normalize_manual_correction(manual_correction)
    conflict = "；".join(item for item in [conflict_message, parse_conflict] if item)
    status = "draft"
    existing = _existing_correction(db, project_id, rule_kind, correction_key)
    if existing:
        existing.workpaper_code = workpaper_code
        existing.workpaper_name = workpaper_name
        existing.rule_id = rule_id
        existing.candidate_key = candidate_key
        existing.sheet_or_section = sheet_or_section
        existing.target_field = target_field
        existing.locator = locator
        existing.source_type = source_type
        existing.original_payload = _json_text(original_payload or {})
        existing.manual_correction = manual_correction
        existing.normalized_payload = normalized_payload
        existing.status = status
        existing.conflict_message = conflict
        existing.updated_at = datetime.utcnow()
        return "conflict" if conflict else "updated"
    db.add(
        RuleManualCorrection(
            project_id=project_id,
            workpaper_code=workpaper_code,
            workpaper_name=workpaper_name,
            rule_kind=rule_kind,
            rule_id=rule_id,
            candidate_key=candidate_key,
            correction_key=correction_key,
            sheet_or_section=sheet_or_section,
            target_field=target_field,
            locator=locator,
            source_type=source_type,
            original_payload=_json_text(original_payload or {}),
            manual_correction=manual_correction,
            normalized_payload=normalized_payload,
            status=status,
            conflict_message=conflict,
            created_by_user_id=user.id,
        )
    )
    return "conflict" if conflict else "created"


def import_manual_corrections_from_excel(
    db: Session,
    *,
    project_id: int,
    user: User,
    content: bytes,
    inspection_payload: dict[str, Any],
) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        rows = iter_manual_rows(wb)
    finally:
        wb.close()

    formal_index, formal_rule_ids = build_formal_index(inspection_payload)
    candidate_index, duplicate_candidate_keys = build_candidate_index(inspection_payload)
    has_candidate_scan = bool((inspection_payload.get("template_scan") or {}).get("workpapers"))
    seen_upload_keys: Counter[str] = Counter()
    result = {
        "project_id": project_id,
        "processed": 0,
        "created": 0,
        "updated": 0,
        "conflict": 0,
        "skipped": 0,
        "errors": 0,
        "details": [],
    }

    for row in rows:
        result["processed"] += 1
        row_no = row["_row_no"]
        rule_kind = row["_rule_kind"]
        correction = row["_manual_correction"]
        try:
            if rule_kind == "formal":
                rule_id = clean_text(row.get("规则编号"))
                workpaper_code = clean_text(row.get("底稿编号"))
                target_field = clean_text(row.get("目标字段"))
                locator = clean_text(row.get("定位内容"))
                correction_key = formal_correction_key(rule_id, workpaper_code, target_field, locator)
                matched = formal_index.get(correction_key)
                conflict = ""
                if not matched:
                    conflict = "规则编号不存在" if rule_id not in formal_rule_ids else "规则编号存在但目标字段/定位内容未匹配"
                if seen_upload_keys[correction_key]:
                    conflict = "同一正式规则定位在导入文件中出现多条修正"
                seen_upload_keys[correction_key] += 1
                original = matched["payload"] if matched else dict(row)
                action = _upsert_correction(
                    db,
                    project_id=project_id,
                    user=user,
                    rule_kind="formal",
                    correction_key=correction_key,
                    rule_id=rule_id,
                    workpaper_code=workpaper_code,
                    workpaper_name=clean_text(row.get("底稿名称")),
                    sheet_or_section=clean_text(row.get("Sheet/章节")),
                    target_field=target_field,
                    locator=locator,
                    source_type=clean_text(row.get("填充来源")),
                    original_payload=original,
                    manual_correction=correction,
                    conflict_message=conflict,
                )
            else:
                workpaper_code = clean_text(row.get("底稿编号"))
                candidate_type = clean_text(row.get("候选类型"))
                label = clean_text(row.get("标签/程序"))
                location = clean_text(row.get("位置"))
                locator = clean_text(row.get("定位内容"))
                candidate_key = candidate_locator_key(workpaper_code, candidate_type, label, location, locator)
                correction_key = candidate_key
                matched = candidate_index.get(candidate_key)
                conflict = ""
                if candidate_key in duplicate_candidate_keys:
                    conflict = "同一候选定位在当前模板扫描中匹配多条候选规则"
                elif has_candidate_scan and not matched:
                    conflict = "候选规则定位不存在或当前项目未扫描到该候选"
                if seen_upload_keys[correction_key]:
                    conflict = "同一候选定位在导入文件中出现多条修正"
                seen_upload_keys[correction_key] += 1
                original = matched["payload"] if matched else dict(row)
                action = _upsert_correction(
                    db,
                    project_id=project_id,
                    user=user,
                    rule_kind="candidate",
                    correction_key=correction_key,
                    candidate_key=candidate_key,
                    workpaper_code=workpaper_code,
                    workpaper_name=clean_text(row.get("底稿名称")),
                    sheet_or_section=location,
                    target_field=label,
                    locator=locator,
                    source_type=clean_text(row.get("来源推断")),
                    original_payload=original,
                    manual_correction=correction,
                    conflict_message=conflict,
                )
            if action == "created":
                result["created"] += 1
            elif action == "updated":
                result["updated"] += 1
            else:
                result["conflict"] += 1
            result["details"].append({"row": row_no, "sheet": row["_sheet"], "action": action})
        except Exception as exc:  # pragma: no cover - defensive import error reporting
            result["errors"] += 1
            result["details"].append({"row": row_no, "sheet": row["_sheet"], "action": "error", "message": str(exc)})
    db.commit()
    return result


def import_batch_to_summary(item: RuleManualCorrectionImport) -> dict[str, Any]:
    return {
        "id": item.id,
        "project_id": item.project_id,
        "filename": item.filename,
        "status": item.status,
        "total_rows": item.total_rows,
        "processed_rows": item.processed_rows,
        "created_count": item.created_count,
        "updated_count": item.updated_count,
        "conflict_count": item.conflict_count,
        "skipped_count": item.skipped_count,
        "error_count": item.error_count,
        "error_message": item.error_message,
        "created_by_user_id": item.created_by_user_id,
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
        "started_at": item.started_at.isoformat() if item.started_at else None,
        "completed_at": item.completed_at.isoformat() if item.completed_at else None,
    }


def create_import_batch(
    db: Session,
    *,
    project_id: int,
    user: User,
    filename: str,
    file_path: str,
) -> RuleManualCorrectionImport:
    batch = RuleManualCorrectionImport(
        project_id=project_id,
        filename=clean_text(filename, 255),
        file_path=file_path,
        status="pending",
        result_payload="{}",
        created_by_user_id=user.id,
    )
    db.add(batch)
    db.commit()
    db.refresh(batch)
    return batch


def list_import_batches(db: Session, project_id: int, limit: int = 20) -> list[RuleManualCorrectionImport]:
    return db.execute(
        select(RuleManualCorrectionImport)
        .where(RuleManualCorrectionImport.project_id == project_id)
        .order_by(RuleManualCorrectionImport.created_at.desc(), RuleManualCorrectionImport.id.desc())
        .limit(limit)
    ).scalars().all()


def get_import_batch(db: Session, project_id: int, batch_id: int) -> RuleManualCorrectionImport | None:
    return db.execute(
        select(RuleManualCorrectionImport).where(
            RuleManualCorrectionImport.project_id == project_id,
            RuleManualCorrectionImport.id == batch_id,
        )
    ).scalar_one_or_none()


def run_import_batch(batch_id: int) -> None:
    with SessionLocal() as db:
        batch = db.get(RuleManualCorrectionImport, batch_id)
        if not batch:
            return
        batch.status = "running"
        batch.started_at = datetime.utcnow()
        db.commit()

        try:
            project = db.get(Project, batch.project_id)
            user = db.get(User, batch.created_by_user_id) if batch.created_by_user_id else None
            if not project:
                raise RuntimeError("项目不存在")
            if not user:
                raise RuntimeError("导入用户不存在")
            with open(batch.file_path, "rb") as handle:
                content = handle.read()
            wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
            try:
                batch.total_rows = len(iter_manual_rows(wb))
                db.commit()
            finally:
                wb.close()
            attachments = db.execute(
                select(Attachment).where(Attachment.project_id == project.id).order_by(Attachment.index_no)
            ).scalars().all()
            workpapers = db.execute(
                select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.code)
            ).scalars().all()
            payload = inspect_project_autofill_rules(list_dict(workpapers), list_dict(attachments))
            payload["project_id"] = project.id
            payload["project_name"] = project.name
            result = import_manual_corrections_from_excel(
                db,
                project_id=project.id,
                user=user,
                content=content,
                inspection_payload=payload,
            )
            batch = db.get(RuleManualCorrectionImport, batch_id)
            if not batch:
                return
            batch.status = "completed"
            batch.total_rows = int(result.get("processed") or 0)
            batch.processed_rows = int(result.get("processed") or 0)
            batch.created_count = int(result.get("created") or 0)
            batch.updated_count = int(result.get("updated") or 0)
            batch.conflict_count = int(result.get("conflict") or 0)
            batch.skipped_count = int(result.get("skipped") or 0)
            batch.error_count = int(result.get("errors") or 0)
            batch.error_message = ""
            batch.result_payload = _json_text(result)
            batch.completed_at = datetime.utcnow()
            db.commit()
        except Exception as exc:  # pragma: no cover - background defensive path
            db.rollback()
            batch = db.get(RuleManualCorrectionImport, batch_id)
            if not batch:
                return
            batch.status = "failed"
            batch.error_count = max(batch.error_count, 1)
            batch.error_message = clean_text(str(exc), 2000)
            batch.completed_at = datetime.utcnow()
            db.commit()


def update_correction_status(
    db: Session,
    *,
    project_id: int,
    correction_id: int,
    user: User,
    status: str,
) -> RuleManualCorrection | None:
    if status not in VALID_STATUSES:
        raise ValueError(f"Unsupported correction status: {status}")
    item = db.execute(
        select(RuleManualCorrection).where(
            RuleManualCorrection.project_id == project_id,
            RuleManualCorrection.id == correction_id,
        )
    ).scalar_one_or_none()
    if not item:
        return None
    item.status = status
    item.approved_by_user_id = user.id
    item.approved_at = datetime.utcnow() if status in {"approved", "rejected"} else None
    item.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(item)
    return item


def correction_to_summary(item: RuleManualCorrection) -> dict[str, Any]:
    return {
        "id": item.id,
        "rule_kind": item.rule_kind,
        "rule_id": item.rule_id,
        "candidate_key": item.candidate_key,
        "workpaper_code": item.workpaper_code,
        "workpaper_name": item.workpaper_name,
        "sheet_or_section": item.sheet_or_section,
        "target_field": item.target_field,
        "locator": item.locator,
        "source_type": item.source_type,
        "manual_correction": item.manual_correction,
        "status": item.status,
        "conflict_message": item.conflict_message,
        "created_by_user_id": item.created_by_user_id,
        "approved_by_user_id": item.approved_by_user_id,
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
        "approved_at": item.approved_at.isoformat() if item.approved_at else None,
        "original_summary": {
            "sheet_or_section": item.sheet_or_section,
            "target_field": item.target_field,
            "locator": clean_text(item.locator, 160),
            "source_type": item.source_type,
        },
    }


def list_project_manual_corrections(db: Session, project_id: int) -> list[RuleManualCorrection]:
    return db.execute(
        select(RuleManualCorrection)
        .where(RuleManualCorrection.project_id == project_id)
        .order_by(RuleManualCorrection.rule_kind, RuleManualCorrection.workpaper_code, RuleManualCorrection.id)
    ).scalars().all()


def summarize_corrections(corrections: list[RuleManualCorrection]) -> dict[str, Any]:
    by_kind = Counter(item.rule_kind for item in corrections)
    by_status = Counter(item.status for item in corrections)
    conflicts = sum(1 for item in corrections if item.conflict_message)
    return {
        "total": len(corrections),
        "formal": by_kind.get("formal", 0),
        "candidate": by_kind.get("candidate", 0),
        "conflict": conflicts,
        "by_status": dict(by_status),
    }


def apply_manual_corrections_to_payload(payload: dict[str, Any], corrections: list[RuleManualCorrection]) -> dict[str, Any]:
    formal_by_key = defaultdict(list)
    candidate_by_key = defaultdict(list)
    for item in corrections:
        if item.rule_kind == "formal":
            formal_by_key[item.correction_key].append(item)
        elif item.rule_kind == "candidate":
            candidate_by_key[item.candidate_key or item.correction_key].append(item)

    for rule in payload.get("rules") or []:
        rule_corrections: list[dict[str, Any]] = []
        for target in rule.get("targets") or []:
            key = formal_target_key(rule, target)
            summaries = [correction_to_summary(item) for item in formal_by_key.get(key, [])]
            if summaries:
                target["manual_corrections"] = summaries
                active = next((item for item in reversed(summaries) if item["status"] != "rejected"), None)
                if active:
                    target["manual_correction"] = active["manual_correction"]
                    target["manual_correction_status"] = active["status"]
                    target["manual_correction_conflict"] = active["conflict_message"]
                    target["manual_correction_effective"] = active["status"] == "approved"
                rule_corrections.extend(summaries)
        if rule_corrections:
            rule["manual_corrections"] = rule_corrections
            rule["manual_correction_count"] = len(rule_corrections)
            active_rule_correction = next(
                (item for item in reversed(rule_corrections) if item["status"] != "rejected"),
                None,
            )
            rule["manual_correction_status"] = active_rule_correction["status"] if active_rule_correction else "rejected"
            rule["manual_correction_effective"] = bool(
                active_rule_correction and active_rule_correction["status"] == "approved"
            )
            rule["manual_correction_conflict"] = "；".join(
                item["conflict_message"] for item in rule_corrections if item.get("conflict_message")
            )
        else:
            rule["manual_correction_count"] = 0

    template_scan = payload.get("template_scan") or {}
    for workpaper in template_scan.get("workpapers") or []:
        code = workpaper.get("code", "")
        for candidate in workpaper.get("candidate_rules") or []:
            label = candidate.get("label") or candidate.get("procedure_text") or ""
            key = candidate_locator_key(code, candidate.get("rule_type"), label, candidate.get("location"), candidate.get("locator"))
            summaries = [correction_to_summary(item) for item in candidate_by_key.get(key, [])]
            if summaries:
                candidate["manual_corrections"] = summaries
                active = next((item for item in reversed(summaries) if item["status"] != "rejected"), None)
                if active:
                    candidate["manual_correction"] = active["manual_correction"]
                    candidate["manual_correction_status"] = active["status"]
                    candidate["manual_correction_conflict"] = active["conflict_message"]
                    candidate["manual_correction_effective"] = active["status"] == "approved"

    grouped = {"formal": [], "candidate": []}
    for item in corrections:
        grouped.setdefault(item.rule_kind, []).append(correction_to_summary(item))
    payload["manual_corrections"] = {
        "summary": summarize_corrections(corrections),
        "groups": grouped,
    }
    return payload


def _plan_field_matches_correction(plan_field: Any, correction_field: str) -> bool:
    field = clean_text(plan_field)
    target = clean_text(correction_field)
    return bool(target) and (field == target or field.startswith(f"{target}."))


def _format_plan_message(existing: Any, addition: str) -> str:
    text = clean_text(existing)
    return "；".join(item for item in [text, addition] if item)


def apply_approved_manual_corrections_to_plan(
    plan: list[dict[str, Any]],
    corrections: list[RuleManualCorrection],
) -> list[dict[str, Any]]:
    approved: dict[tuple[str, str], list[RuleManualCorrection]] = defaultdict(list)
    for item in corrections:
        if item.rule_kind != "formal" or item.status != "approved":
            continue
        approved[(clean_text(item.rule_id), clean_text(item.target_field))].append(item)

    for item in plan:
        item.setdefault("value_source", "rule")
        rule_id = clean_text(item.get("rule_id"))
        field = clean_text(item.get("field"))
        matches: list[RuleManualCorrection] = []
        for (candidate_rule_id, correction_field), rows in approved.items():
            if candidate_rule_id == rule_id and _plan_field_matches_correction(field, correction_field):
                matches.extend(rows)
        if not matches:
            continue

        original_value = item.get("new_value")
        item["original_rule_value"] = original_value
        if len(matches) > 1:
            item["value_source"] = "manual_correction_conflict"
            item["manual_correction_ids"] = [row.id for row in matches]
            item["manual_correction_status"] = "approved"
            item["status"] = "blocked"
            item["message"] = _format_plan_message(
                item.get("message"),
                "多条已批准人工修正同时匹配该计划项，需人工处理后再生效",
            )
            continue

        correction = matches[0]
        item["manual_correction_id"] = correction.id
        item["manual_correction_status"] = correction.status
        item["manual_correction_field"] = correction.target_field
        item["manual_correction_locator"] = clean_text(correction.locator, 160)
        if correction.conflict_message:
            item["value_source"] = "manual_correction_conflict"
            item["status"] = "blocked"
            item["message"] = _format_plan_message(
                item.get("message"),
                f"已批准人工修正存在冲突，未覆盖原规则值：{correction.conflict_message}",
            )
            continue

        item["value_source"] = "manual_correction_approved"
        item["new_value"] = correction.manual_correction
        item["message"] = _format_plan_message(
            item.get("message"),
            f"使用已批准人工修正 #{correction.id} 覆盖自动规则候选值",
        )
    return plan
