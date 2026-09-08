from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


VERIFYABLE_STATUSES = {"changed", "unchanged"}
SKIPPED_STATUSES = {"planned", "skipped", "blocked"}


@dataclass
class WorkpaperIdentity:
    workpaper_id: int | None = None
    workpaper_code: str = ""
    workpaper_name: str = ""


def normalize_value(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip())


def display_value(value: Any, limit: int = 1000) -> str:
    text = str(value or "")
    return text if len(text) <= limit else text[:limit] + "..."


def values_match(expected: Any, actual: Any) -> bool:
    return normalize_value(expected) == normalize_value(actual)


def is_single_cell(cell: str) -> bool:
    if not cell or ":" in cell:
        return False
    try:
        range_boundaries(cell)
    except ValueError:
        return False
    return True


def table_text(table: Any) -> str:
    return "; ".join(" / ".join(cell.text.strip() for cell in row.cells) for row in table.rows[1:])


def table_header_text(table: Any) -> str:
    if not table.rows:
        return ""
    return normalize_value(" ".join(cell.text for cell in table.rows[0].cells))


TABLE_LOCATOR_ALIASES = {
    "信息处理控制": ["信息处理控制", "应用控制"],
    "涉及的信息系统": ["涉及的信息系统", "涉及系统", "应用系统"],
    "财务报表科目": ["财务报表科目", "财务报表项目"],
}


def locator_keyword_aliases(keyword: str) -> list[str]:
    aliases = TABLE_LOCATOR_ALIASES.get(keyword, [keyword])
    return [normalize_value(alias) for alias in aliases if alias]


def find_table_for_locator(doc: Any, locator: str) -> Any | None:
    raw = locator.removeprefix("table:")
    keywords = [item for item in re.split(r"[/|,，、\s]+", raw) if item]
    keyword_groups = [locator_keyword_aliases(item) for item in keywords]
    for table in doc.tables:
        header = table_header_text(table)
        if keyword_groups and all(any(keyword in header for keyword in group) for group in keyword_groups):
            return table
    if keyword_groups:
        first_group = keyword_groups[0]
        for table in doc.tables:
            if any(keyword in table_header_text(table) for keyword in first_group):
                return table
    return None


def read_excel_target(path: Path, sheet_name: str, cell: str) -> tuple[Any | None, str]:
    if not sheet_name:
        return None, "missing sheet name"
    if not is_single_cell(cell):
        return None, "only single-cell Excel readback is supported"
    wb = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    try:
        if sheet_name not in wb.sheetnames:
            return None, f"sheet not found: {sheet_name}"
        return wb[sheet_name][cell].value, ""
    finally:
        wb.close()


def parse_document_locator(cell: str, locator: str) -> str:
    text = cell.removeprefix("document:")
    if text:
        return text
    return locator


def read_docx_target(path: Path, cell: str, locator: str) -> tuple[Any | None, str]:
    target = parse_document_locator(cell, locator)
    doc = Document(path)
    paragraph_match = re.fullmatch(r"paragraph:(\d+)", target)
    if paragraph_match:
        index = int(paragraph_match.group(1))
        if index >= len(doc.paragraphs):
            return None, f"paragraph index out of range: {index}"
        return doc.paragraphs[index].text, ""
    cell_match = re.fullmatch(r"table:(\d+):r(\d+):c(\d+)", target)
    if cell_match:
        table_idx, row_idx, col_idx = (int(item) for item in cell_match.groups())
        if table_idx >= len(doc.tables):
            return None, f"table index out of range: {table_idx}"
        table = doc.tables[table_idx]
        if row_idx >= len(table.rows) or col_idx >= len(table.rows[row_idx].cells):
            return None, f"table cell out of range: r{row_idx}:c{col_idx}"
        return table.rows[row_idx].cells[col_idx].text, ""
    if target.startswith("table:"):
        table = find_table_for_locator(doc, target)
        if table is None:
            return None, f"table not found for locator: {target}"
        return table_text(table), ""
    return None, "document locator is not supported"


def target_type_for_item(item: dict[str, Any], path: Path) -> str:
    locator = str(item.get("locator") or "")
    cell = str(item.get("cell") or "")
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        if is_single_cell(cell):
            return "cell"
        if cell or locator.startswith("headers:"):
            return "locator"
        return "locator"
    if path.suffix.lower() == ".docx":
        target = parse_document_locator(cell, locator)
        if target.startswith("paragraph:"):
            return "paragraph"
        if target.startswith("table:"):
            return "table"
        return "locator"
    return "locator"


def verify_plan_item(
    item: dict[str, Any],
    *,
    identity: WorkpaperIdentity | None = None,
) -> dict[str, Any]:
    path = Path(str(item.get("workbook_path") or "")).expanduser()
    identity = identity or WorkpaperIdentity()
    target_type = target_type_for_item(item, path)
    expected = item.get("new_value")
    status = str(item.get("status") or "")
    sheet_or_section = str(item.get("sheet_name") or "")
    target_cell = str(item.get("cell") or "")
    locator = str(item.get("locator") or "")
    actual: Any | None = None
    verification_status = "skipped"
    message = ""

    if status in SKIPPED_STATUSES:
        verification_status = "skipped"
        message = f"write status is {status}; no readback attempted"
    elif not path.exists():
        verification_status = "failed"
        message = "workpaper copy does not exist"
    elif not target_cell and not locator:
        verification_status = "missing_locator"
        message = "missing readback locator"
    elif status not in VERIFYABLE_STATUSES:
        verification_status = "skipped"
        message = f"write status is {status}; no readback attempted"
    else:
        try:
            if path.suffix.lower() in {".xlsx", ".xlsm"}:
                actual, read_message = read_excel_target(path, sheet_or_section, target_cell)
            elif path.suffix.lower() == ".docx":
                actual, read_message = read_docx_target(path, target_cell, locator)
            else:
                actual, read_message = None, f"unsupported file type: {path.suffix}"
            if read_message:
                verification_status = "not_supported" if "supported" in read_message else "failed"
                message = read_message
            elif values_match(expected, actual):
                verification_status = "passed"
                message = "readback matched expected value"
            else:
                verification_status = "failed"
                message = "readback value did not match expected value"
        except Exception as exc:  # pragma: no cover - defensive report path
            verification_status = "failed"
            message = f"readback error: {exc}"

    return {
        "workpaper_id": identity.workpaper_id,
        "workpaper_code": identity.workpaper_code,
        "workpaper_name": identity.workpaper_name,
        "rule_id": item.get("rule_id", ""),
        "scope": item.get("scope", ""),
        "target_type": target_type,
        "sheet_or_section": sheet_or_section,
        "target_cell": target_cell,
        "locator": locator,
        "field": item.get("field", ""),
        "old_value": item.get("old_value"),
        "expected_value": expected,
        "actual_value": display_value(actual),
        "write_status": status,
        "write_message": item.get("message", ""),
        "verification_status": verification_status,
        "verification_message": message,
        "workbook_path": str(path),
    }


def summarize_verification(records: list[dict[str, Any]]) -> dict[str, int]:
    summary = {
        "write_items": len(records),
        "passed": 0,
        "failed": 0,
        "skipped": 0,
        "not_supported": 0,
        "missing_locator": 0,
    }
    for record in records:
        status = str(record.get("verification_status") or "")
        if status in summary:
            summary[status] += 1
    return summary


def verify_plan_items(
    plan_items: list[dict[str, Any]],
    *,
    workpaper_by_path: dict[str, WorkpaperIdentity] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, int]]:
    workpaper_by_path = workpaper_by_path or {}
    records: list[dict[str, Any]] = []
    for item in plan_items:
        path = str(Path(str(item.get("workbook_path") or "")).expanduser())
        records.append(verify_plan_item(item, identity=workpaper_by_path.get(path)))
    return records, summarize_verification(records)
