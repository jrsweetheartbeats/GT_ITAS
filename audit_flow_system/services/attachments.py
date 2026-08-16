from __future__ import annotations

from pathlib import Path
import re
from typing import Any, Optional

from docx import Document
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.utils import get_or_404
from ..models import Attachment, Project, Workpaper
from .privacy import RedactionTerm, redact_text


ATTACHMENT_EXTENSIONS = {".xlsx", ".xlsm", ".xls", ".docx", ".doc", ".pdf", ".png", ".jpg", ".jpeg", ".txt", ".csv"}
AUTO_SCAN_EXTENSIONS = ATTACHMENT_EXTENSIONS - {".png", ".jpg", ".jpeg"}
SKIP_DIR_NAMES = {".git", ".svn", "__pycache__", "node_modules", ".venv", "venv", ".idea", ".vscode"}


def safe_project_code(project: Project) -> str:
    base = project.code or project.entity_name or f"P{project.id}"
    return re.sub(r"\s+", "-", base.strip())[:80] or f"P{project.id}"


def sanitize_workpaper_code(code: str) -> str:
    cleaned = re.sub(r"[^0-9A-Za-z.\-_]+", "", code.strip())
    return cleaned or "ATT"


def compact_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).lower()


def next_attachment_index(db: Session, project_id: int, workpaper_code: str = "") -> str:
    project = get_or_404(db, Project, project_id, "项目")
    base = sanitize_workpaper_code(workpaper_code) if workpaper_code else safe_project_code(project)
    rows = db.execute(select(Attachment.index_no).where(Attachment.project_id == project_id)).scalars().all()
    max_no = 0
    prefix = f"{base}-"
    for index_no in rows:
        if index_no.startswith(prefix):
            match = re.search(r"-(\d+)$", index_no)
            if match:
                max_no = max(max_no, int(match.group(1)))
    return f"{base}-{max_no + 1}"


def file_type_for_path(path: Path) -> str:
    return path.suffix.lower().lstrip(".")


def guess_workpaper_for_file(path: Path, workpapers: list[Workpaper]) -> Optional[Workpaper]:
    text = compact_text(" ".join(path.parts[-4:]))
    best: Optional[Workpaper] = None
    best_score = 0
    for wp in workpapers:
        score = 0
        code = compact_text(wp.code)
        name = compact_text(wp.name)
        if code and code in text:
            score += 10
        if name and name in text:
            score += 6
        for token in re.split(r"[\s._\-]+", wp.code):
            token_norm = compact_text(token)
            if len(token_norm) >= 3 and token_norm in text:
                score += 2
        if score > best_score:
            best = wp
            best_score = score
    return best if best_score else None


def scan_project_attachment_files(project: Project, subdir: str = "") -> list[Path]:
    if not project.project_root:
        raise HTTPException(status_code=400, detail="项目未维护根目录")
    root = Path(project.project_root).expanduser()
    if subdir:
        root = root / subdir
    if not root.exists() or not root.is_dir():
        raise HTTPException(status_code=400, detail=f"项目目录不存在：{root}")
    files: list[Path] = []
    for path in root.rglob("*"):
        if not path.is_file():
            continue
        if path.name.startswith("~$") or path.name.startswith("."):
            continue
        if any(part in SKIP_DIR_NAMES or part.startswith(".") for part in path.parts):
            continue
        if path.suffix.lower() not in AUTO_SCAN_EXTENSIONS:
            continue
        files.append(path)
    return sorted(files)


def extract_workpaper_text(path: Path, *, redact: bool = False, redaction_terms: list[RedactionTerm] | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        wb = load_workbook(path, read_only=True, data_only=True, keep_vba=suffix == ".xlsm")
        try:
            chunks: list[str] = []
            for ws in wb.worksheets:
                chunks.append(ws.title)
                for row in ws.iter_rows(values_only=True):
                    chunks.extend(str(value) for value in row if value not in (None, ""))
            text = "\n".join(chunks)
            return redact_text(text, redaction_terms)[0] if redact else text
        finally:
            wb.close()
    if suffix == ".docx":
        doc = Document(path)
        chunks = [para.text for para in doc.paragraphs if para.text]
        for table in doc.tables:
            for row in table.rows:
                chunks.extend(cell.text for cell in row.cells if cell.text)
        text = "\n".join(chunks)
        return redact_text(text, redaction_terms)[0] if redact else text
    if suffix in {".txt", ".csv"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return redact_text(text, redaction_terms)[0] if redact else text
    return ""


def candidate_reference_tokens(text: str) -> set[str]:
    tokens = set()
    for match in re.findall(r"[A-Za-z0-9][A-Za-z0-9_.-]{1,60}-\d{1,4}", text):
        tokens.add(match.strip(".,;:，。；：、()（）[]【】<>《》"))
    for match in re.findall(r"[<《]([^<>《》]{2,80})[>》]", text):
        cleaned = match.strip()
        if cleaned:
            tokens.add(cleaned)
    return tokens


def append_reference(existing: str, workpaper_code: str) -> str:
    parts = [part.strip() for part in re.split(r"[、,;\n]+", existing or "") if part.strip()]
    if workpaper_code and workpaper_code not in parts:
        parts.append(workpaper_code)
    return "、".join(parts)


def likely_attachment_reference(token: str) -> bool:
    if not token or len(token) > 80:
        return False
    return bool(re.search(r"\d", token) and re.search(r"[A-Za-z_.-]", token))
