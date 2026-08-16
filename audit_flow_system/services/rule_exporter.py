from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


MANUAL_CORRECTION_HEADER = "人工修正"


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _join_text(values: Any) -> str:
    if not isinstance(values, list):
        return _cell_text(values)
    return "；".join(_cell_text(item) for item in values if _cell_text(item))


def _manual_correction_text(item: dict[str, Any]) -> str:
    direct = _cell_text(item.get("manual_correction"))
    if direct:
        return direct
    corrections = item.get("manual_corrections")
    if isinstance(corrections, list) and corrections:
        return _cell_text(corrections[-1].get("manual_correction"))
    return ""


def _setup_sheet(ws, headers: list[str]) -> None:
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _finish_sheet(ws, widths: dict[str, int] | None = None) -> None:
    widths = widths or {}
    for index, column_cells in enumerate(ws.columns, start=1):
        letter = get_column_letter(index)
        header = str(ws.cell(row=1, column=index).value or "")
        max_len = min(
            60,
            max((len(str(cell.value or "")) for cell in column_cells), default=0) + 2,
        )
        ws.column_dimensions[letter].width = widths.get(header, max(12, max_len))
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _target_rows(rule: dict[str, Any]) -> list[dict[str, Any]]:
    targets = rule.get("targets")
    if isinstance(targets, list) and targets:
        return [target for target in targets if isinstance(target, dict)] or [{}]
    return [{}]


def _formal_rule_rows(payload: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for rule in payload.get("rules") or []:
        if not isinstance(rule, dict):
            continue
        for target in _target_rows(rule):
            rows.append(
                [
                    _cell_text(rule.get("rule_id") or rule.get("id")),
                    _cell_text(rule.get("name")),
                    _cell_text(rule.get("scope")),
                    _cell_text(rule.get("stage")),
                    _cell_text(rule.get("workpaper_code")),
                    _cell_text(rule.get("workpaper_name")),
                    _cell_text(rule.get("sheet") or rule.get("section")),
                    _cell_text(target.get("target_field") or target.get("field") or rule.get("target_field")),
                    _cell_text(target.get("target_cell") or target.get("cell") or rule.get("target_cell")),
                    _cell_text(target.get("locator_method")),
                    _cell_text(target.get("locator_display") or target.get("locator") or rule.get("locator")),
                    _cell_text(target.get("write_mode")),
                    _cell_text(rule.get("source_type")),
                    _join_text([item.get("name", "") for item in rule.get("evidence_requirements") or [] if isinstance(item, dict)]),
                    _cell_text(rule.get("matching_logic")),
                    _cell_text(rule.get("content_template")),
                    _cell_text(rule.get("definition_status") or rule.get("status")),
                    _cell_text(rule.get("project_status")),
                    _join_text(rule.get("definition_gaps")),
                    _join_text(rule.get("project_gaps")),
                    _manual_correction_text(target) or _manual_correction_text(rule),
                ]
            )
    return rows


def _candidate_rule_rows(payload: dict[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    template_scan = payload.get("template_scan") or {}
    for workpaper in template_scan.get("workpapers") or []:
        if not isinstance(workpaper, dict):
            continue
        for candidate in workpaper.get("candidate_rules") or []:
            if not isinstance(candidate, dict):
                continue
            rows.append(
                [
                    _cell_text(workpaper.get("code")),
                    _cell_text(workpaper.get("name")),
                    _cell_text(candidate.get("rule_type")),
                    _cell_text(candidate.get("label") or candidate.get("procedure_text")),
                    _cell_text(candidate.get("location")),
                    _cell_text(candidate.get("locator")),
                    _cell_text(candidate.get("write_mode")),
                    _cell_text(candidate.get("source_guess")),
                    _cell_text(candidate.get("confidence")),
                    _cell_text(candidate.get("reason")),
                    _cell_text(workpaper.get("file_path")),
                    _manual_correction_text(candidate),
                ]
            )
    return rows


def build_rule_inspection_export(payload: dict[str, Any]) -> bytes:
    wb = Workbook()

    summary = wb.active
    summary.title = "导出说明"
    _setup_sheet(summary, ["项目", "项目ID", "正式规则数", "候选规则数", "说明"])
    summary.append(
        [
            _cell_text(payload.get("project_name")),
            _cell_text(payload.get("project_id")),
            _cell_text(len(payload.get("rules") or [])),
            _cell_text((payload.get("template_scan") or {}).get("summary", {}).get("candidate_rule_count", 0)),
            "人工修正列用于人工补充定位、来源、匹配或适用范围修订意见；导入后保存为项目 overlay，不直接覆盖原规则或底稿。",
        ]
    )
    _finish_sheet(summary, {"说明": 72})

    formal = wb.create_sheet("正式规则")
    formal_headers = [
        "规则编号",
        "规则名称",
        "Scope",
        "阶段",
        "底稿编号",
        "底稿名称",
        "Sheet/章节",
        "目标字段",
        "目标单元格",
        "定位方式",
        "定位内容",
        "写入方式",
        "填充来源",
        "证据要求",
        "匹配逻辑",
        "内容模板",
        "定义状态",
        "项目状态",
        "定义缺口",
        "项目缺口",
        MANUAL_CORRECTION_HEADER,
    ]
    _setup_sheet(formal, formal_headers)
    for row in _formal_rule_rows(payload):
        formal.append(row)
    _finish_sheet(
        formal,
        {
            "规则编号": 22,
            "底稿名称": 26,
            "目标字段": 26,
            "定位内容": 34,
            "匹配逻辑": 42,
            "内容模板": 50,
            "定义缺口": 40,
            "项目缺口": 40,
            MANUAL_CORRECTION_HEADER: 32,
        },
    )

    candidates = wb.create_sheet("模板识别候选")
    candidate_headers = [
        "底稿编号",
        "底稿名称",
        "候选类型",
        "标签/程序",
        "位置",
        "定位内容",
        "写入方式",
        "来源推断",
        "置信度",
        "识别原因",
        "文件路径",
        MANUAL_CORRECTION_HEADER,
    ]
    _setup_sheet(candidates, candidate_headers)
    for row in _candidate_rule_rows(payload):
        candidates.append(row)
    _finish_sheet(
        candidates,
        {
            "底稿名称": 26,
            "标签/程序": 56,
            "定位内容": 32,
            "识别原因": 42,
            "文件路径": 56,
            MANUAL_CORRECTION_HEADER: 32,
        },
    )

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
