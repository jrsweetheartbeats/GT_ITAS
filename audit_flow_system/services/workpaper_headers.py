from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import json
from pathlib import Path
import re
import shutil
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..core.config import BASE_DIR
from ..models import Project, Workpaper
from .workpaper_metadata import clean_cell_text, parse_header_datetime


HEADER_UPDATE_ROOT = BASE_DIR / "tmp" / "workpaper_header_updates"
HEADER_BACKUP_ROOT = BASE_DIR / "tmp" / "workpaper_header_backups"
SUPPORTED_EXCEL = {".xlsx", ".xlsm"}
SUPPORTED_WORD = {".docx"}
HEADER_SCAN_ROWS = 12
HEADER_SCAN_SHEETS = 3

HEADER_FIELDS: dict[str, list[str]] = {
    "entity_name": ["被审计单位", "客户名称"],
    "preparer": ["编制人"],
    "prepared_date": ["编制日期"],
    "reviewer": ["复核人"],
    "reviewed_date": ["复核日期"],
    "audit_period": ["截止日", "审计期间"],
    "index_no": ["索引号"],
    "page_no": ["页次"],
}
LABEL_TO_FIELD = {re.sub(r"\s+", "", label): field for field, labels in HEADER_FIELDS.items() for label in labels}
LABEL_PATTERN = re.compile(
    r"(被\s*审计\s*单位|客户名称|编制日期|编制人|复核日期|复核人|截止日|审计期间|索引号|页\s*次)(?:\s*[（(][^）)]*[）)])?(?:\s*[：:]\s*|\s+|$)"
)
LABEL_ONLY_PATTERN = re.compile(
    r"^\s*(被\s*审计\s*单位|客户名称|编制日期|编制人|复核日期|复核人|截止日|审计期间|索引号|页\s*次)(?:\s*[（(][^）)]*[）)])?\s*[：:]?\s*$"
)


@dataclass(frozen=True)
class HeaderStyle:
    font_name: str = "宋体"
    font_size: int = 10
    bold: bool = False
    horizontal: str = "left"
    vertical: str = "center"


DEFAULT_HEADER_STYLE = HeaderStyle()


def normalize_label(value: str) -> str:
    return re.sub(r"\s+", "", value or "").replace("页次", "页次")


def label_field(label: str) -> str:
    return LABEL_TO_FIELD.get(normalize_label(label), "")


def normalize_date_text(value: Any) -> str:
    parsed = parse_header_datetime(value)
    return parsed.date().isoformat() if parsed else clean_cell_text(value)


def date_or_empty(value: Any) -> str:
    parsed = parse_header_datetime(value)
    return parsed.date().isoformat() if parsed else ""


def project_person_name(person: Any) -> str:
    return clean_cell_text(getattr(person, "display_name", "") or getattr(person, "username", ""))


def target_header_value(project: Project, workpaper: Workpaper, field: str, label: str, current_value: Any) -> str:
    current = clean_cell_text(current_value)
    if field == "entity_name":
        return clean_cell_text(project.entity_name)
    if field == "preparer":
        return project_person_name(workpaper.preparer) or project_person_name(project.field_leader) or project_person_name(project.project_leader)
    if field == "reviewer":
        return project_person_name(project.project_leader) or project_person_name(project.manager)
    if field == "prepared_date":
        return date_or_empty(current) or current
    if field == "reviewed_date":
        return date_or_empty(current) or current
    if field == "audit_period":
        label_text = normalize_label(label)
        if "审计期间" in label_text and project.audit_scope_start and project.audit_scope_end:
            return f"{project.audit_scope_start.isoformat()}至{project.audit_scope_end.isoformat()}"
        if project.audit_scope_end:
            return project.audit_scope_end.isoformat()
        if project.audit_year:
            return f"{project.audit_year}-12-31"
        return current
    if field == "index_no":
        return clean_cell_text(workpaper.code) or current
    if field == "page_no":
        return current
    return current


def split_labeled_cell(text: str) -> tuple[str, str, str]:
    matches = list(LABEL_PATTERN.finditer(clean_cell_text(text)))
    if len(matches) != 1:
        return "", "", ""
    match = matches[0]
    label = normalize_label(match.group(1))
    field = label_field(label)
    if not field:
        return "", "", ""
    value = text[match.end() :].strip(" \t\r\n：:")
    return field, label, value


def label_only(text: str) -> tuple[str, str]:
    match = LABEL_ONLY_PATTERN.match(clean_cell_text(text))
    if not match:
        return "", ""
    label = normalize_label(match.group(1))
    return label_field(label), label


def excel_cell_name(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def resolve_merged_cell(ws: Worksheet, row: int, col: int):
    cell = ws.cell(row=row, column=col)
    if not isinstance(cell, MergedCell):
        return cell
    coord = cell.coordinate
    for merged_range in ws.merged_cells.ranges:
        if coord in merged_range:
            return ws.cell(row=merged_range.min_row, column=merged_range.min_col)
    return cell


def apply_excel_style(cell: Any, style: HeaderStyle = DEFAULT_HEADER_STYLE) -> None:
    if isinstance(cell, MergedCell):
        return
    cell.font = Font(name=style.font_name, size=style.font_size, bold=style.bold)
    cell.alignment = Alignment(horizontal=style.horizontal, vertical=style.vertical, wrap_text=False)


def nearby_value_cell(ws: Worksheet, row_idx: int, label_col: int) -> tuple[int, str]:
    fallback = min(label_col + 1, ws.max_column)
    for col_idx in range(label_col + 1, min(ws.max_column, label_col + 8) + 1):
        text = clean_cell_text(ws.cell(row=row_idx, column=col_idx).value)
        if not text:
            continue
        if LABEL_PATTERN.search(text) or label_only(text)[0]:
            break
        return col_idx, text
    return fallback, clean_cell_text(ws.cell(row=row_idx, column=fallback).value)


def field_status(current: str, target: str) -> str:
    if not target:
        return "missing_target"
    return "ok" if clean_cell_text(current) == clean_cell_text(target) else "needs_update"


def is_probable_workpaper_index(value: str, expected_code: str) -> bool:
    text = clean_cell_text(value)
    if not text:
        return True
    if expected_code and expected_code.upper() in text.upper():
        return True
    return bool(re.fullmatch(r"[A-Z]{1,4}\d{1,3}[A-Z]?(?:[-.][0-9A-Z]+)*", text.upper()))


def should_include_header_item(field: str, current: str, workpaper: Workpaper) -> bool:
    if field == "index_no":
        return is_probable_workpaper_index(current, workpaper.code or "")
    if field == "page_no":
        text = clean_cell_text(current)
        return not text or bool(re.fullmatch(r"(第)?\d+(页)?", text))
    return True


def excel_field_item(
    *,
    workpaper: Workpaper,
    path: Path,
    ws: Worksheet,
    field: str,
    label: str,
    current: str,
    target: str,
    label_cell: str,
    value_cell: str,
    mode: str,
) -> dict[str, Any]:
    status = field_status(current, target)
    return {
        "workpaper_id": workpaper.id,
        "workpaper_code": workpaper.code,
        "workpaper_name": workpaper.name,
        "file_path": str(path),
        "file_type": path.suffix.lower(),
        "field": field,
        "label": label,
        "sheet_name": ws.title,
        "label_cell": label_cell,
        "value_cell": value_cell,
        "mode": mode,
        "current_value": current,
        "target_value": target,
        "status": status,
        "message": "将统一字段值和表头字体" if status != "missing_target" else "缺少可写入的目标值",
    }


def scan_excel_headers(project: Project, workpaper: Workpaper, path: Path) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=False, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        items: list[dict[str, Any]] = []
        seen: set[tuple[str, str, str]] = set()
        for ws in wb.worksheets[:HEADER_SCAN_SHEETS]:
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, HEADER_SCAN_ROWS)):
                for cell in row:
                    text = clean_cell_text(cell.value)
                    if not text:
                        continue
                    if len(list(LABEL_PATTERN.finditer(text))) > 1:
                        continue
                    field, label = label_only(text)
                    if field:
                        value_col, current = nearby_value_cell(ws, cell.row, cell.column)
                        value_coord = excel_cell_name(cell.row, value_col)
                        target = target_header_value(project, workpaper, field, label, current)
                        if not should_include_header_item(field, current, workpaper):
                            continue
                        key = (ws.title, field, value_coord)
                        if key in seen:
                            continue
                        items.append(
                            excel_field_item(
                                workpaper=workpaper,
                                path=path,
                                ws=ws,
                                field=field,
                                label=label,
                                current=current,
                                target=target,
                                label_cell=cell.coordinate,
                                value_cell=value_coord,
                                mode="split_cell",
                            )
                        )
                        seen.add(key)
                        continue
                    field, label, embedded_value = split_labeled_cell(text)
                    if field:
                        target = target_header_value(project, workpaper, field, label, embedded_value)
                        if not should_include_header_item(field, embedded_value, workpaper):
                            continue
                        key = (ws.title, field, cell.coordinate)
                        if key not in seen:
                            items.append(
                                excel_field_item(
                                    workpaper=workpaper,
                                    path=path,
                                    ws=ws,
                                    field=field,
                                    label=label,
                                    current=embedded_value,
                                    target=target,
                                    label_cell=cell.coordinate,
                                    value_cell=cell.coordinate,
                                    mode="same_cell",
                                )
                            )
                            seen.add(key)
                        continue
        return items
    finally:
        wb.close()


def set_docx_font(paragraph: Any, style: HeaderStyle = DEFAULT_HEADER_STYLE) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if style.horizontal == "left" else WD_ALIGN_PARAGRAPH.CENTER
    for run in paragraph.runs:
        run.font.name = style.font_name
        run.font.size = Pt(style.font_size)
        run.font.bold = style.bold
        if run._element.rPr is not None:
            run._element.rPr.rFonts.set(qn("w:eastAsia"), style.font_name)


def set_docx_text(container: Any, text: str, style: HeaderStyle = DEFAULT_HEADER_STYLE) -> None:
    container.text = text
    if hasattr(container, "runs"):
        set_docx_font(container, style)
        return
    paragraphs = getattr(container, "paragraphs", [])
    for paragraph in paragraphs:
        set_docx_font(paragraph, style)


def docx_field_item(
    *,
    workpaper: Workpaper,
    path: Path,
    field: str,
    label: str,
    current: str,
    target: str,
    label_cell: str,
    value_cell: str,
    mode: str,
) -> dict[str, Any]:
    status = field_status(current, target)
    return {
        "workpaper_id": workpaper.id,
        "workpaper_code": workpaper.code,
        "workpaper_name": workpaper.name,
        "file_path": str(path),
        "file_type": path.suffix.lower(),
        "field": field,
        "label": label,
        "sheet_name": "Word",
        "label_cell": label_cell,
        "value_cell": value_cell,
        "mode": mode,
        "current_value": current,
        "target_value": target,
        "status": status,
        "message": "将统一字段值和表头字体" if status != "missing_target" else "缺少可写入的目标值",
    }


def scan_docx_headers(project: Project, workpaper: Workpaper, path: Path) -> list[dict[str, Any]]:
    doc = Document(path)
    items: list[dict[str, Any]] = []
    for index, paragraph in enumerate(doc.paragraphs[:HEADER_SCAN_ROWS]):
        text = clean_cell_text(paragraph.text)
        if len(list(LABEL_PATTERN.finditer(text))) > 1:
            continue
        field, label, current = split_labeled_cell(text)
        if not field or not current:
            continue
        target = target_header_value(project, workpaper, field, label, current)
        if not should_include_header_item(field, current, workpaper):
            continue
        locator = f"paragraph:{index}"
        items.append(
            docx_field_item(
                workpaper=workpaper,
                path=path,
                field=field,
                label=label,
                current=current,
                target=target,
                label_cell=locator,
                value_cell=locator,
                mode="same_paragraph",
            )
        )
    for table_idx, table in enumerate(doc.tables[:3]):
        for row_idx, row in enumerate(table.rows[:HEADER_SCAN_ROWS]):
            cells = row.cells
            for col_idx, cell in enumerate(cells):
                text = clean_cell_text(cell.text)
                if not text or len(list(LABEL_PATTERN.finditer(text))) > 1:
                    continue
                locator = f"table:{table_idx}:r{row_idx}:c{col_idx}"
                field, label = label_only(text)
                if field:
                    value_idx = min(col_idx + 1, len(cells) - 1)
                    current = clean_cell_text(cells[value_idx].text)
                    target = target_header_value(project, workpaper, field, label, current)
                    if not should_include_header_item(field, current, workpaper):
                        continue
                    items.append(
                        docx_field_item(
                            workpaper=workpaper,
                            path=path,
                            field=field,
                            label=label,
                            current=current,
                            target=target,
                            label_cell=locator,
                            value_cell=f"table:{table_idx}:r{row_idx}:c{value_idx}",
                            mode="split_cell",
                        )
                    )
                    continue
                field, label, current = split_labeled_cell(text)
                if not field:
                    continue
                target = target_header_value(project, workpaper, field, label, current)
                if not should_include_header_item(field, current, workpaper):
                    continue
                items.append(
                    docx_field_item(
                        workpaper=workpaper,
                        path=path,
                        field=field,
                        label=label,
                        current=current,
                        target=target,
                        label_cell=locator,
                        value_cell=locator,
                        mode="same_cell",
                    )
                )
    return items


def scan_workpaper_headers(project: Project, workpaper: Workpaper, path_override: str | Path | None = None) -> dict[str, Any]:
    path = Path(path_override or workpaper.file_path or "").expanduser()
    result = {
        "workpaper_id": workpaper.id,
        "workpaper_code": workpaper.code,
        "workpaper_name": workpaper.name,
        "file_path": str(path),
        "file_type": path.suffix.lower(),
        "supported": False,
        "status": "unsupported",
        "fields": [],
        "field_count": 0,
        "needs_update_count": 0,
        "missing_target_count": 0,
        "message": "",
    }
    if not path.exists() or not path.is_file():
        result["status"] = "missing_file"
        result["message"] = "底稿文件不存在"
        return result
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXCEL | SUPPORTED_WORD:
        result["message"] = "暂不支持该文件类型"
        return result
    result["supported"] = True
    try:
        fields = scan_excel_headers(project, workpaper, path) if suffix in SUPPORTED_EXCEL else scan_docx_headers(project, workpaper, path)
    except Exception as exc:
        result["status"] = "error"
        result["message"] = str(exc)
        return result
    result["fields"] = fields
    result["field_count"] = len(fields)
    result["needs_update_count"] = sum(1 for item in fields if item.get("status") == "needs_update")
    result["missing_target_count"] = sum(1 for item in fields if item.get("status") == "missing_target")
    result["status"] = "ok" if fields else "no_header_found"
    result["message"] = "" if fields else "未在首页范围识别到标准表头字段"
    return result


def scan_project_headers(project: Project, workpapers: list[Workpaper]) -> dict[str, Any]:
    rows = [scan_workpaper_headers(project, workpaper) for workpaper in workpapers]
    field_rows = [field for row in rows for field in row.get("fields", [])]
    return {
        "project_id": project.id,
        "project_name": project.name,
        "mode": "scan",
        "style": DEFAULT_HEADER_STYLE.__dict__,
        "workpaper_count": len(rows),
        "supported_count": sum(1 for row in rows if row.get("supported")),
        "field_count": len(field_rows),
        "needs_update_count": sum(1 for field in field_rows if field.get("status") == "needs_update"),
        "missing_target_count": sum(1 for field in field_rows if field.get("status") == "missing_target"),
        "workpapers": rows,
        "fields": field_rows,
    }


def safe_copy_name(workpaper: Workpaper, source: Path) -> str:
    name = re.sub(r"[^0-9A-Za-z_.\-\u4e00-\u9fff]+", "_", source.name).strip("._") or f"workpaper_{workpaper.id}{source.suffix}"
    return f"{workpaper.id}_{name}"


def apply_excel_headers(path: Path, fields: list[dict[str, Any]], style: HeaderStyle = DEFAULT_HEADER_STYLE) -> int:
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    changed = 0
    try:
        for item in fields:
            if item.get("status") == "missing_target":
                continue
            ws = wb[item["sheet_name"]]
            target = clean_cell_text(item.get("target_value"))
            label = clean_cell_text(item.get("label"))
            if item.get("mode") == "same_cell":
                row = int(re.sub(r"[^0-9]", "", item["value_cell"]))
                col_letters = re.sub(r"[^A-Z]", "", item["value_cell"])
                col = sum((ord(char) - 64) * (26 ** idx) for idx, char in enumerate(reversed(col_letters)))
                cell = resolve_merged_cell(ws, row, col)
                new_value = f"{label}：{target}" if target else f"{label}："
                if clean_cell_text(cell.value) != new_value:
                    cell.value = new_value
                    changed += 1
                apply_excel_style(cell, style)
                continue
            label_cell = ws[item["label_cell"]]
            value_cell = ws[item["value_cell"]]
            label_cell = resolve_merged_cell(ws, label_cell.row, label_cell.column)
            value_cell = resolve_merged_cell(ws, value_cell.row, value_cell.column)
            label_text = f"{label}："
            if clean_cell_text(label_cell.value) != label_text:
                label_cell.value = label_text
                changed += 1
            if clean_cell_text(value_cell.value) != target:
                value_cell.value = target
                changed += 1
            apply_excel_style(label_cell, style)
            apply_excel_style(value_cell, style)
        wb.save(path)
    finally:
        wb.close()
    return changed


def parse_docx_locator(locator: str) -> tuple[str, int, int, int]:
    paragraph = re.fullmatch(r"paragraph:(\d+)", locator)
    if paragraph:
        return "paragraph", int(paragraph.group(1)), -1, -1
    table_cell = re.fullmatch(r"table:(\d+):r(\d+):c(\d+)", locator)
    if table_cell:
        table_idx, row_idx, col_idx = [int(item) for item in table_cell.groups()]
        return "table", table_idx, row_idx, col_idx
    raise ValueError(f"unsupported docx locator: {locator}")


def docx_target(doc: Any, locator: str) -> Any:
    kind, first, row_idx, col_idx = parse_docx_locator(locator)
    if kind == "paragraph":
        return doc.paragraphs[first]
    return doc.tables[first].rows[row_idx].cells[col_idx]


def apply_docx_headers(path: Path, fields: list[dict[str, Any]], style: HeaderStyle = DEFAULT_HEADER_STYLE) -> int:
    doc = Document(path)
    changed = 0
    for item in fields:
        if item.get("status") == "missing_target":
            continue
        target_value = clean_cell_text(item.get("target_value"))
        label = clean_cell_text(item.get("label"))
        if item.get("mode") in {"same_cell", "same_paragraph"}:
            target = docx_target(doc, item["value_cell"])
            new_text = f"{label}：{target_value}" if target_value else f"{label}："
            current_text = clean_cell_text(getattr(target, "text", ""))
            if current_text != new_text:
                changed += 1
            set_docx_text(target, new_text, style)
            continue
        label_target = docx_target(doc, item["label_cell"])
        value_target = docx_target(doc, item["value_cell"])
        label_text = f"{label}："
        if clean_cell_text(getattr(label_target, "text", "")) != label_text:
            changed += 1
        if clean_cell_text(getattr(value_target, "text", "")) != target_value:
            changed += 1
        set_docx_text(label_target, label_text, style)
        set_docx_text(value_target, target_value, style)
    doc.save(path)
    return changed


def verify_safe_path(path: Path, root: Path) -> bool:
    try:
        path.resolve().relative_to(root.resolve())
        return True
    except ValueError:
        return False


def apply_project_headers_to_test_copies(project: Project, workpapers: list[Workpaper]) -> dict[str, Any]:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = HEADER_UPDATE_ROOT / f"project_{project.id}_{timestamp}"
    workpaper_dir = output_dir / "workpapers"
    workpaper_dir.mkdir(parents=True, exist_ok=True)
    scan = scan_project_headers(project, workpapers)
    result_rows: list[dict[str, Any]] = []
    for row in scan["workpapers"]:
        source = Path(row["file_path"]).expanduser()
        result = dict(row)
        result["source_path"] = str(source)
        result["test_copy_path"] = ""
        result["write_status"] = "skipped"
        result["write_message"] = row.get("message", "")
        result["changed_cells"] = 0
        if not row.get("supported") or not row.get("fields") or not source.exists():
            result_rows.append(result)
            continue
        target = workpaper_dir / safe_copy_name(Workpaper(id=row["workpaper_id"], code=row["workpaper_code"], name=row["workpaper_name"], file_path=str(source)), source)
        shutil.copy2(source, target)
        if not verify_safe_path(target, HEADER_UPDATE_ROOT):
            result["write_status"] = "blocked"
            result["write_message"] = "测试副本路径越界"
            result_rows.append(result)
            continue
        try:
            if target.suffix.lower() in SUPPORTED_EXCEL:
                changed = apply_excel_headers(target, row["fields"])
            elif target.suffix.lower() in SUPPORTED_WORD:
                changed = apply_docx_headers(target, row["fields"])
            else:
                changed = 0
            result["test_copy_path"] = str(target)
            result["changed_cells"] = changed
            result["write_status"] = "updated"
            result["write_message"] = "已写入测试副本"
        except Exception as exc:
            result["test_copy_path"] = str(target)
            result["write_status"] = "error"
            result["write_message"] = str(exc)
        result_rows.append(result)
    report = {
        **scan,
        "mode": "test_copy",
        "output_dir": str(output_dir),
        "report_json": str(output_dir / "workpaper_header_update_report.json"),
        "all_write_targets_safe": all(
            not row.get("test_copy_path") or verify_safe_path(Path(row["test_copy_path"]), HEADER_UPDATE_ROOT)
            for row in result_rows
        ),
        "updated_count": sum(1 for row in result_rows if row.get("write_status") == "updated"),
        "error_count": sum(1 for row in result_rows if row.get("write_status") == "error"),
        "workpapers": result_rows,
        "fields": [field for row in result_rows for field in row.get("fields", [])],
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    Path(report["report_json"]).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return report


def load_workpaper_extracted_fields(workpaper: Workpaper) -> dict[str, Any]:
    try:
        data = json.loads(workpaper.extracted_fields_json or "{}")
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def update_workpaper_header_cache(workpaper: Workpaper, fields: list[dict[str, Any]]) -> None:
    data = load_workpaper_extracted_fields(workpaper)
    changed = False
    for field in fields:
        if field.get("status") == "missing_target":
            continue
        key = str(field.get("field") or "")
        value = clean_cell_text(field.get("target_value"))
        if not key or not value:
            continue
        if clean_cell_text(data.get(key)) != value:
            data[key] = value
            changed = True
    if changed:
        data["metadata_source"] = "workpaper_header_real_write"
        workpaper.extracted_fields_json = json.dumps(data, ensure_ascii=False)


def row_has_writable_fields(row: dict[str, Any]) -> bool:
    return any(field.get("status") != "missing_target" for field in row.get("fields", []))


def apply_project_headers_to_real_files(project: Project, workpapers: list[Workpaper]) -> dict[str, Any]:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_dir = HEADER_BACKUP_ROOT / f"project_{project.id}_{timestamp}"
    backup_dir = output_dir / "workpapers"
    backup_dir.mkdir(parents=True, exist_ok=True)
    scan = scan_project_headers(project, workpapers)
    workpapers_by_id = {int(workpaper.id): workpaper for workpaper in workpapers}
    result_rows: list[dict[str, Any]] = []
    for row in scan["workpapers"]:
        source = Path(row["file_path"]).expanduser()
        workpaper = workpapers_by_id.get(int(row["workpaper_id"]))
        result = dict(row)
        result["source_path"] = str(source)
        result["backup_path"] = ""
        result["write_status"] = "skipped"
        result["write_message"] = row.get("message", "")
        result["changed_cells"] = 0
        if not row.get("supported") or not row.get("fields") or not source.exists() or workpaper is None:
            result_rows.append(result)
            continue
        if not row_has_writable_fields(row):
            result["write_message"] = "无可写入字段；缺少主数据目标值"
            result_rows.append(result)
            continue
        backup = backup_dir / safe_copy_name(workpaper, source)
        try:
            shutil.copy2(source, backup)
            result["backup_path"] = str(backup)
            if source.suffix.lower() in SUPPORTED_EXCEL:
                changed = apply_excel_headers(source, row["fields"])
            elif source.suffix.lower() in SUPPORTED_WORD:
                changed = apply_docx_headers(source, row["fields"])
            else:
                changed = 0
            update_workpaper_header_cache(workpaper, row["fields"])
            result["changed_cells"] = changed
            result["write_status"] = "updated"
            result["write_message"] = "已写入真实底稿；原文件已备份"
        except Exception as exc:
            if backup.exists() and source.exists():
                try:
                    shutil.copy2(backup, source)
                    result["write_message"] = f"写入失败，已从备份恢复：{exc}"
                except Exception as restore_exc:
                    result["write_message"] = f"写入失败，且备份恢复失败：{exc}; restore={restore_exc}"
            else:
                result["write_message"] = f"写入失败：{exc}"
            result["write_status"] = "error"
        result_rows.append(result)
    report = {
        **scan,
        "mode": "real_write",
        "output_dir": str(output_dir),
        "backup_dir": str(backup_dir),
        "report_json": str(output_dir / "workpaper_header_real_write_report.json"),
        "all_write_targets_backed_up": all(
            row.get("write_status") != "updated" or bool(row.get("backup_path"))
            for row in result_rows
        ),
        "updated_count": sum(1 for row in result_rows if row.get("write_status") == "updated"),
        "error_count": sum(1 for row in result_rows if row.get("write_status") == "error"),
        "skipped_count": sum(1 for row in result_rows if row.get("write_status") == "skipped"),
        "workpapers": result_rows,
        "fields": [field for row in result_rows for field in row.get("fields", [])],
    }
    output_dir.mkdir(parents=True, exist_ok=True)
    Path(report["report_json"]).write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return report
