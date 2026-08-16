from __future__ import annotations

from datetime import date, datetime
import json
from pathlib import Path
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook

from ..models import Workpaper


LABEL_TO_FIELD = {
    "被审计单位": "entity_name",
    "客户名称": "entity_name",
    "编制人": "preparer",
    "编制日期": "prepared_date",
    "复核人": "reviewer",
    "复核日期": "reviewed_date",
    "截止日": "audit_period",
    "审计期间": "audit_period",
    "索引号": "index_no",
    "页次": "page_no",
}

LABEL_PATTERN = re.compile(
    r"(被\s*审计\s*单位|客户名称|编制日期|编制人|复核日期|复核人|截止日|审计期间|索引号|页\s*次)\s*[：:]?\s*"
)
LABEL_ONLY_PATTERN = re.compile(
    r"^\s*(被\s*审计\s*单位|客户名称|编制日期|编制人|复核日期|复核人|截止日|审计期间|索引号|页\s*次)\s*[：:]?\s*$"
)


def normalize_label(value: str) -> str:
    return re.sub(r"\s+", "", value or "")


def clean_cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_header_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = clean_cell_text(value)
    if not text:
        return None
    match = re.search(r"(20\d{2})[./\-年](\d{1,2})[./\-月](\d{1,2})", text)
    if not match:
        return None
    try:
        return datetime(int(match.group(1)), int(match.group(2)), int(match.group(3)))
    except ValueError:
        return None


def parse_labeled_text(text: str) -> dict[str, str]:
    text = clean_cell_text(text)
    if not text:
        return {}
    matches = list(LABEL_PATTERN.finditer(text))
    if not matches:
        return {}

    fields: dict[str, str] = {}
    for index, match in enumerate(matches):
        field = LABEL_TO_FIELD.get(normalize_label(match.group(1)))
        if not field:
            continue
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        value = text[match.end() : end].strip(" \t\r\n：:")
        if value:
            fields.setdefault(field, value)
    return fields


def label_only_field(text: str) -> str:
    match = LABEL_ONLY_PATTERN.match(clean_cell_text(text))
    if not match:
        return ""
    return LABEL_TO_FIELD.get(normalize_label(match.group(1)), "")


def first_right_value(cells: list[tuple[int, str]], start_index: int) -> str:
    for _column, value in cells[start_index + 1 :]:
        if not value:
            continue
        if parse_labeled_text(value) or label_only_field(value):
            return ""
        return value
    return ""


def merge_row_fields(fields: dict[str, str], cells: list[tuple[int, str]]) -> None:
    for index, (_column, text) in enumerate(cells):
        for field, value in parse_labeled_text(text).items():
            fields.setdefault(field, value)

        field = label_only_field(text)
        if field and field not in fields:
            value = first_right_value(cells, index)
            if value:
                fields[field] = value


def extract_excel_metadata(path: Path, max_rows: int = 12, max_sheets: int = 3) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        fields: dict[str, str] = {}
        for ws in wb.worksheets[:max_sheets]:
            for row in ws.iter_rows(min_row=1, max_row=max_rows):
                cells = [
                    (cell.column, clean_cell_text(cell.value))
                    for cell in row
                    if clean_cell_text(cell.value)
                ]
                if cells:
                    merge_row_fields(fields, cells)
            if fields.get("preparer") and fields.get("prepared_date"):
                break
        return fields
    finally:
        wb.close()


def extract_docx_metadata(path: Path, max_rows: int = 12) -> dict[str, str]:
    doc = Document(path)
    fields: dict[str, str] = {}
    for paragraph in doc.paragraphs[:max_rows]:
        merge_row_fields(fields, [(1, clean_cell_text(paragraph.text))])
    for table in doc.tables[:3]:
        for row in table.rows[:max_rows]:
            cells = [(index, clean_cell_text(cell.text)) for index, cell in enumerate(row.cells, start=1)]
            cells = [(index, text) for index, text in cells if text]
            if cells:
                merge_row_fields(fields, cells)
    return fields


def extract_workpaper_metadata(file_path: str | Path) -> dict[str, str]:
    if not file_path:
        return {}
    path = Path(file_path).expanduser()
    if not path.exists() or not path.is_file():
        return {}
    try:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            return extract_excel_metadata(path)
        if path.suffix.lower() == ".docx":
            return extract_docx_metadata(path)
    except Exception:
        return {}
    return {}


def load_extracted_fields(workpaper: Workpaper) -> dict[str, Any]:
    try:
        data = json.loads(workpaper.extracted_fields_json or "{}")
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def merge_workpaper_metadata(base: dict[str, Any], file_path: str | Path) -> dict[str, Any]:
    data = dict(base)
    metadata = extract_workpaper_metadata(file_path)
    if metadata:
        for key, value in metadata.items():
            data.setdefault(key, value)
        data.setdefault("metadata_source", "workpaper_header")
    return data


def workpaper_preparer_text(workpaper: Workpaper) -> str:
    fields = load_extracted_fields(workpaper)
    for key in ("preparer", "prepared_by", "prepared_by_name", "compiler"):
        value = clean_cell_text(fields.get(key))
        if value:
            return value
    return clean_cell_text(extract_workpaper_metadata(workpaper.file_path).get("preparer"))


def workpaper_header_fields(workpaper: Workpaper) -> dict[str, str]:
    fields = {
        key: clean_cell_text(value)
        for key, value in load_extracted_fields(workpaper).items()
        if key in {"entity_name", "preparer", "prepared_date", "reviewer", "reviewed_date", "audit_period", "index_no", "page_no"}
        and clean_cell_text(value)
    }
    for key, value in extract_workpaper_metadata(workpaper.file_path).items():
        fields.setdefault(key, value)
    return fields
