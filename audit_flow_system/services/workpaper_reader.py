from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..core.utils import obj_dict
from ..models import Attachment, Workpaper
from .workpaper_catalog import (
    BASIC_WORKPAPER_BY_CODE,
    c22_test_point_code,
    normalize_workpaper_code,
)


STAGE_ORDER = ("planning", "execution", "delivery", "reporting")
STAGE_LABELS = {
    "planning": "1.项目准备",
    "execution": "2.项目实施",
    "delivery": "3.项目交付",
    "reporting": "4.项目报告",
    "completion": "3.项目交付",
}


def stage_label(stage: str) -> str:
    return STAGE_LABELS.get(stage or "", stage or "未分组")


def workpaper_node(workpaper: Workpaper) -> dict[str, Any]:
    return {
        "id": f"workpaper-{workpaper.id}",
        "type": "workpaper",
        "label": f"{workpaper.code} {workpaper.name}".strip(),
        "workpaper_id": workpaper.id,
        "code": workpaper.code,
        "name": workpaper.name,
        "stage": workpaper.stage,
        "status": workpaper.status,
        "file_path": workpaper.file_path,
        "children": [],
    }


def attachment_node(attachment: Attachment) -> dict[str, Any]:
    label = " ".join(part for part in [attachment.index_no, attachment.title] if part).strip()
    return {
        "id": f"attachment-{attachment.id}",
        "type": "attachment",
        "label": label or attachment.file_path or f"附件 {attachment.id}",
        "attachment_id": attachment.id,
        "workpaper_id": attachment.workpaper_id,
        "index_no": attachment.index_no,
        "name": attachment.title,
        "file_type": attachment.file_type,
        "status": attachment.status,
        "file_path": attachment.file_path,
        "referenced_in": attachment.referenced_in,
        "children": [],
    }


def attachment_sheet_name(attachment: Attachment) -> str:
    reference = (attachment.referenced_in or "").strip()
    candidates = [
        r"\[([^\]]+)\]",
        r"([^!\s,，;；]+)!\$?[A-Z]{1,3}\$?\d+",
        r"(?:sheet|工作表|页签|表单)\s*[:：]\s*([^,，;；\s]+)",
    ]
    for pattern in candidates:
        match = re.search(pattern, reference, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()[:80] or "未识别Sheet"
    return "未识别Sheet"


def safe_node_part(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff_-]+", "-", value or "").strip("-")[:80] or "sheet"


def attachment_import_relative_path(attachment: Attachment, project_root: str = "") -> list[str]:
    path = Path(attachment.file_path or "")
    try:
        relative = path.relative_to(Path(project_root)) if project_root else path
    except ValueError:
        relative = path
    parts = list(relative.parts)
    try:
        marker = parts.index("批量导入")
    except ValueError:
        return []
    return [part for part in parts[marker + 1 :] if part not in {"", ".", ".."}]


def attachment_import_tree_node(parent_id: str, attachments: list[Attachment], project_root: str) -> dict[str, Any] | None:
    root: dict[str, Any] = {"folders": {}, "files": []}
    for attachment in attachments:
        parts = attachment_import_relative_path(attachment, project_root)
        if not parts:
            continue
        cursor = root
        for part in parts[:-1]:
            cursor = cursor["folders"].setdefault(part, {"folders": {}, "files": []})
        cursor["files"].append(attachment)

    def render_folder(name: str, value: dict[str, Any], path_parts: list[str]) -> dict[str, Any]:
        key = safe_node_part("-".join(path_parts + [name]))
        children = [
            render_folder(child_name, child_value, path_parts + [name])
            for child_name, child_value in sorted(value["folders"].items())
        ]
        children.extend(attachment_node(item) for item in sorted(value["files"], key=lambda row: (row.title, row.id)))
        return {
            "id": f"{parent_id}-import-{key}",
            "type": "attachment_path_folder",
            "label": name,
            "status": str(len(children)),
            "children": children,
        }

    if not root["folders"] and not root["files"]:
        return None
    children = [render_folder(name, value, []) for name, value in sorted(root["folders"].items())]
    children.extend(attachment_node(item) for item in sorted(root["files"], key=lambda row: (row.title, row.id)))
    return {
        "id": f"{parent_id}-batch-import",
        "type": "attachment_path_folder",
        "label": "批量导入文件",
        "status": str(len(children)),
        "children": children,
    }


def attachment_folder_node(parent_id: str, attachments: list[Attachment], project_root: str = "", label: str = "附件文件夹") -> dict[str, Any]:
    by_sheet: dict[str, list[Attachment]] = {}
    imported = [attachment for attachment in attachments if attachment_import_relative_path(attachment, project_root)]
    for attachment in attachments:
        if attachment in imported:
            continue
        by_sheet.setdefault(attachment_sheet_name(attachment), []).append(attachment)
    sheet_nodes = [
        {
            "id": f"{parent_id}-attachments-sheet-{safe_node_part(sheet)}",
            "type": "attachment_sheet_folder",
            "label": sheet,
            "sheet": sheet,
            "status": str(len(rows)),
            "children": [attachment_node(attachment) for attachment in rows],
        }
        for sheet, rows in sorted(by_sheet.items(), key=lambda item: (item[0] == "未识别Sheet", item[0]))
    ]
    children: list[dict[str, Any]] = []
    import_tree = attachment_import_tree_node(parent_id, imported, project_root)
    if import_tree is not None:
        children.append(import_tree)
    children.extend(sheet_nodes)
    return {
        "id": f"{parent_id}-attachments",
        "type": "attachment_folder",
        "label": label,
        "status": str(len(attachments)),
        "children": children,
    }


def normalize_code(code: str) -> str:
    return normalize_workpaper_code(code)


def preview_recognition_note(workpaper: Workpaper) -> str:
    code = normalize_code(workpaper.code)
    if tree_stage_for_workpaper(workpaper) == "planning" or code.startswith(("A", "B")) or code in {"C21", "C21-1"}:
        return "已按底稿结构展示可识别内容。"
    return "该底稿暂不识别关键字段，仅展示可读取的文件内容。"


def tree_stage_for_workpaper(workpaper: Workpaper) -> str:
    code = normalize_code(workpaper.code)
    path = workpaper.file_path or ""
    basic_spec = BASIC_WORKPAPER_BY_CODE.get(code)
    if basic_spec:
        return basic_spec["stage"]
    if code.startswith(("A27", "A14")):
        return "delivery"
    if "3.项目交付" in path or "项目交付" in path or "项目总结" in path:
        return "delivery"
    if "4.项目报告" in path or "项目报告" in path:
        return "reporting"
    if "1.项目准备" in path or "项目准备" in path or "项目计划" in path:
        return "planning"
    return workpaper.stage or "unknown"


def nest_c22_test_points(groups: dict[str, dict[str, Any]], test_point_nodes: list[dict[str, Any]]) -> None:
    if not test_point_nodes:
        return
    execution = groups["execution"]
    c22_node = next(
        (child for child in execution["children"] if normalize_code(child.get("code", "")) == "C22"),
        None,
    )
    for node in test_point_nodes:
        node["type"] = "test_point"
        node["test_point"] = c22_test_point_code(node.get("code"))
        node["label"] = f"{node['test_point']} {node.get('name') or '测试点'}".strip()
    folder = {
        "id": "c22-test-points",
        "type": "test_point_folder",
        "label": "C22 测试点（历史登记）" if c22_node else "C22 测试点（C22 底稿未上传）",
        "stage": "execution",
        "status": str(len(test_point_nodes)),
        "children": test_point_nodes,
    }
    if c22_node:
        c22_node["children"].append(folder)
    else:
        execution["children"].append(folder)


def build_workpaper_tree(
    workpapers: list[Workpaper],
    attachments: list[Attachment] | None = None,
    project_root: str = "",
) -> list[dict[str, Any]]:
    attachments = attachments or []
    workpaper_ids = {workpaper.id for workpaper in workpapers}
    attachments_by_workpaper: dict[int, list[Attachment]] = {}
    loose_attachments: list[Attachment] = []
    for attachment in attachments:
        if attachment.workpaper_id and attachment.workpaper_id in workpaper_ids:
            attachments_by_workpaper.setdefault(attachment.workpaper_id, []).append(attachment)
        else:
            loose_attachments.append(attachment)

    groups: dict[str, dict[str, Any]] = {
        key: {
            "id": f"stage-{key}",
            "type": "stage",
            "label": stage_label(key),
            "stage": key,
            "children": [],
        }
        for key in STAGE_ORDER
    }
    test_point_nodes: list[dict[str, Any]] = []
    for wp in workpapers:
        key = tree_stage_for_workpaper(wp)
        group = groups.setdefault(
            key,
            {
                "id": f"stage-{key}",
                "type": "stage",
                "label": stage_label(key),
                "stage": key,
                "children": [],
            },
        )
        node = workpaper_node(wp)
        node["stage"] = key
        linked_attachments = attachments_by_workpaper.get(wp.id, [])
        if linked_attachments:
            node["children"].append(attachment_folder_node(node["id"], linked_attachments, project_root))
        if c22_test_point_code(wp.code):
            test_point_nodes.append(node)
        else:
            group["children"].append(node)
    if loose_attachments:
        groups["execution"]["children"].append(
            attachment_folder_node("stage-execution-loose", loose_attachments, project_root, "未关联资料")
        )
    nest_c22_test_points(groups, test_point_nodes)
    known = [groups[key] for key in STAGE_ORDER]
    unknown = [groups[key] for key in sorted(groups) if key not in STAGE_ORDER]
    return known + unknown


def preview_lines_from_excel(path: Path, max_lines: int) -> tuple[list[str], list[str], list[dict[str, Any]]]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        lines: list[str] = []
        sheets: list[str] = []
        sheet_sections: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            sheets.append(ws.title)
            if len(lines) >= max_lines:
                continue
            lines.append(f"[{ws.title}]")
            rows: list[dict[str, Any]] = []
            for row in ws.iter_rows():
                cells: list[dict[str, Any]] = []
                for cell in row:
                    if cell.value in (None, ""):
                        continue
                    value = str(cell.value).strip()
                    if not value:
                        continue
                    cells.append(
                        {
                            "cell": f"{get_column_letter(cell.column)}{cell.row}",
                            "location": f"{ws.title}!{get_column_letter(cell.column)}{cell.row}",
                            "value": value,
                        }
                    )
                if not cells:
                    continue
                rows.append({"row": row[0].row if row else 0, "cells": cells})
                lines.append(" | ".join(f"{cell['cell']} {cell['value']}" for cell in cells))
                if len(lines) >= max_lines:
                    break
            sheet_sections.append({"sheet": ws.title, "rows": rows})
        return lines, sheets, sheet_sections
    finally:
        wb.close()


def preview_lines_from_docx(path: Path, max_lines: int) -> tuple[list[str], list[dict[str, Any]]]:
    doc = Document(path)
    lines: list[str] = []
    rows: list[dict[str, Any]] = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        lines.append(f"段落{index} {text}")
        rows.append({"row": index, "cells": [{"cell": f"段落{index}", "location": f"正文!段落{index}", "value": text}]})
        if len(lines) >= max_lines:
            return lines[:max_lines], [{"sheet": "正文", "rows": rows[:max_lines]}]
    for table in doc.tables:
        for row_index, row in enumerate(table.rows, start=1):
            cells = []
            for col_index, cell in enumerate(row.cells, start=1):
                value = cell.text.strip()
                if value:
                    cells.append({"cell": f"表格{row_index}-{col_index}", "location": f"正文!表格{row_index}-{col_index}", "value": value})
            if cells:
                rows.append({"row": row_index, "cells": cells})
                lines.append(" | ".join(f"{cell['cell']} {cell['value']}" for cell in cells))
            if len(lines) >= max_lines:
                return lines[:max_lines], [{"sheet": "正文", "rows": rows[:max_lines]}]
    return lines[:max_lines], [{"sheet": "正文", "rows": rows[:max_lines]}]


def preview_lines_from_text(path: Path, max_lines: int) -> list[str]:
    return path.read_text(encoding="utf-8", errors="ignore").splitlines()[:max_lines]


def read_workpaper_preview(workpaper: Workpaper, max_lines: int = 120) -> dict[str, Any]:
    data = obj_dict(workpaper)
    data.update(
        {
            "file_exists": False,
            "file_type": "",
            "preview_text": "",
            "lines": [],
            "sheets": [],
            "sheet_sections": [],
            "recognition_note": "",
            "error": "",
        }
    )
    if not workpaper.file_path:
        data["error"] = "底稿未维护文件路径"
        return data
    path = Path(workpaper.file_path).expanduser()
    data["file_path"] = str(path)
    data["file_type"] = path.suffix.lower().lstrip(".")
    data["recognition_note"] = preview_recognition_note(workpaper)
    if not path.exists() or not path.is_file():
        data["error"] = f"底稿文件不存在：{path}"
        return data
    data["file_exists"] = True
    try:
        if path.suffix.lower() in {".xlsx", ".xlsm"}:
            lines, sheets, sheet_sections = preview_lines_from_excel(path, max_lines)
            data["lines"] = lines
            data["sheets"] = sheets
            data["sheet_sections"] = sheet_sections
        elif path.suffix.lower() == ".docx":
            lines, sheet_sections = preview_lines_from_docx(path, max_lines)
            data["lines"] = lines
            data["sheet_sections"] = sheet_sections
        elif path.suffix.lower() in {".txt", ".csv"}:
            data["lines"] = preview_lines_from_text(path, max_lines)
        else:
            data["error"] = f"暂不支持预览该文件类型：{path.suffix or '无扩展名'}"
            return data
    except Exception as exc:
        data["error"] = f"底稿预览读取失败：{exc}"
        return data
    data["preview_text"] = "\n".join(data["lines"])
    return data
