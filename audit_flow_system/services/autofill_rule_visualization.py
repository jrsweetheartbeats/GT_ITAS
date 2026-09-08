from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..core.config import BASE_DIR
from ..core.utils import list_dict, obj_dict
from ..models import Attachment, AutofillPlanItem, AutofillRun, EnterpriseContact, Project, ProjectMember, RuleManualCorrection, Workpaper
from .autofill_rule_actions import action_index_for_project, action_key_for_row
from .autofill_rule_inspector import inspect_project_autofill_rules
from .rule_manual_corrections import apply_manual_corrections_to_payload, correction_to_summary


SOURCE_LABELS = {
    "project_context": "项目主数据",
    "attachments": "附件证据",
    "attachments_and_project_context": "附件证据/项目主数据",
    "client": "客户信息",
    "members": "项目成员",
    "partner": "财审签字人",
    "template": "默认模板",
    "manual_correction_approved": "人工修正",
    "manual_correction_conflict": "人工修正冲突",
    "rule": "规则生成",
}

VERIFICATION_STATUSES = ["passed", "failed", "skipped", "not_supported", "missing_locator", "not_verified"]
TARGET_B22A_CONTROLS = ["SA-3", "SA-11", "SA-14", "PE-6", "PE-7"]


def clean_text(value: Any, limit: int | None = None) -> str:
    text = re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()
    if limit and len(text) > limit:
        return text[:limit] + "..."
    return text


def norm(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def workpaper_category(code: str, scope: str = "") -> str:
    text = str(code or scope or "").upper()
    if text.startswith("A"):
        return "A"
    if text.startswith("B"):
        return "B"
    if text.startswith("C"):
        return "C"
    return ""


def source_label(value_source: str, source_type: str) -> str:
    if value_source in SOURCE_LABELS:
        return SOURCE_LABELS[value_source]
    return SOURCE_LABELS.get(source_type, source_type or "规则生成")


def latest_verification_report(project_id: int) -> dict[str, Any] | None:
    root = BASE_DIR / "tmp" / "autofill_verification"
    if not root.exists():
        return None
    reports: list[tuple[int, float, dict[str, Any]]] = []
    candidates = root.glob(f"project_{project_id}_*/autofill_verification_report.json")
    for path in candidates:
        try:
            data = {**json.loads(path.read_text(encoding="utf-8")), "_report_path": str(path)}
            reports.append((int((data.get("summary") or {}).get("write_items") or 0), path.stat().st_mtime, data))
        except (OSError, json.JSONDecodeError):
            continue
    if not reports:
        return None
    reports.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return reports[0][2]


def verification_keys(record: dict[str, Any]) -> list[tuple[str, str, str, str]]:
    rule_id = clean_text(record.get("rule_id"))
    field = clean_text(record.get("field"))
    locator = clean_text(record.get("locator"))
    cell = clean_text(record.get("target_cell"))
    code = clean_text(record.get("workpaper_code"))
    return [
        (rule_id, field, locator, cell),
        (rule_id, field, "", cell),
        (rule_id, field, locator, ""),
        (rule_id, field, code, cell),
        (rule_id, field, code, locator),
    ]


def build_verification_index(report: dict[str, Any] | None) -> dict[tuple[str, str, str, str], dict[str, Any]]:
    index: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    if not report:
        return index
    for record in report.get("verification_records") or []:
        if not isinstance(record, dict):
            continue
        for key in verification_keys(record):
            index.setdefault(key, record)
    return index


def find_verification(
    index: dict[tuple[str, str, str, str], dict[str, Any]],
    *,
    rule_id: str,
    field: str,
    locator: str,
    cell: str,
    workpaper_code: str,
) -> dict[str, Any] | None:
    keys = [
        (clean_text(rule_id), clean_text(field), clean_text(locator), clean_text(cell)),
        (clean_text(rule_id), clean_text(field), "", clean_text(cell)),
        (clean_text(rule_id), clean_text(field), clean_text(locator), ""),
        (clean_text(rule_id), clean_text(field), clean_text(workpaper_code), clean_text(cell)),
        (clean_text(rule_id), clean_text(field), clean_text(workpaper_code), clean_text(locator)),
    ]
    for key in keys:
        if key in index:
            return index[key]
    return None


def build_project_context(db: Session, project: Project) -> dict[str, Any]:
    context = obj_dict(project)
    contacts = db.execute(
        select(EnterpriseContact).where(EnterpriseContact.project_id == project.id).order_by(EnterpriseContact.id)
    ).scalars().all()
    members = db.execute(
        select(ProjectMember).where(ProjectMember.project_id == project.id).order_by(ProjectMember.id)
    ).scalars().all()
    context["contacts"] = list_dict(contacts)
    context["members"] = [
        {
            **obj_dict(member),
            "username": member.user.username if member.user else "",
            "display_name": member.user.display_name if member.user else "",
            "email": member.user.email if member.user else "",
        }
        for member in members
    ]
    return context


def latest_plan_items(db: Session, project_id: int) -> list[dict[str, Any]]:
    run = db.execute(
        select(AutofillRun)
        .where(AutofillRun.project_id == project_id, AutofillRun.apply.is_(False))
        .order_by(AutofillRun.id.desc())
        .limit(1)
    ).scalar_one_or_none()
    if not run:
        return []
    rows = db.execute(
        select(AutofillPlanItem)
        .where(AutofillPlanItem.run_id == run.id)
        .order_by(AutofillPlanItem.id)
    ).scalars().all()
    return list_dict(rows)


def plan_items_from_report(report: dict[str, Any] | None) -> list[dict[str, Any]]:
    if not report:
        return []
    items: list[dict[str, Any]] = []
    for record in report.get("verification_records") or []:
        if not isinstance(record, dict):
            continue
        items.append(
            {
                "rule_id": record.get("rule_id", ""),
                "scope": record.get("scope", ""),
                "workbook_path": record.get("workbook_path", ""),
                "sheet_name": record.get("sheet_or_section", ""),
                "field": record.get("field", ""),
                "locator": record.get("locator", ""),
                "cell": record.get("target_cell", ""),
                "old_value": record.get("old_value", ""),
                "new_value": record.get("expected_value", ""),
                "status": record.get("write_status", ""),
                "message": record.get("write_message", ""),
                "value_source": "rule",
                "original_rule_value": "",
            }
        )
    return items


def rule_map(payload: dict[str, Any]) -> dict[str, dict[str, Any]]:
    return {str(rule.get("rule_id") or rule.get("id") or ""): rule for rule in payload.get("rules") or []}


def target_matches_plan(target: dict[str, Any], item: dict[str, Any]) -> bool:
    target_field = clean_text(target.get("target_field") or target.get("field"))
    item_field = clean_text(item.get("field"))
    if target_field and item_field and (item_field == target_field or item_field.startswith(f"{target_field}.")):
        return True
    target_locator = norm(target.get("locator_display") or target.get("locator") or target.get("target_cell") or target.get("cell"))
    item_locator = norm(item.get("locator") or item.get("cell"))
    return bool(target_locator and item_locator and (target_locator in item_locator or item_locator in target_locator))


def best_target(rule: dict[str, Any], item: dict[str, Any]) -> dict[str, Any]:
    targets = [target for target in rule.get("targets") or [] if isinstance(target, dict)]
    for target in targets:
        if target_matches_plan(target, item):
            return target
    return targets[0] if targets else {}


def manual_by_id(corrections: list[RuleManualCorrection]) -> dict[int, dict[str, Any]]:
    return {item.id: correction_to_summary(item) for item in corrections}


def latest_non_rejected_manual(rule: dict[str, Any], target: dict[str, Any]) -> dict[str, Any] | None:
    rows: list[dict[str, Any]] = []
    for item in target.get("manual_corrections") or []:
        if isinstance(item, dict):
            rows.append(item)
    for item in rule.get("manual_corrections") or []:
        if isinstance(item, dict):
            rows.append(item)
    for item in reversed(rows):
        if item.get("status") != "rejected":
            return item
    return None


def workpaper_by_path(workpapers: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {str(Path(str(item.get("file_path") or "")).expanduser()): item for item in workpapers if item.get("file_path")}


def workpaper_from_plan(item: dict[str, Any], path_index: dict[str, dict[str, Any]]) -> dict[str, Any]:
    path = str(Path(str(item.get("workbook_path") or "")).expanduser())
    return path_index.get(path, {})


def applicable_scope(project: Project, rule: dict[str, Any], item: dict[str, Any] | None = None) -> str:
    parts = [
        f"项目:{project.name}",
        f"底稿:{rule.get('workpaper_code') or ''}",
        f"scope:{rule.get('scope') or ''}",
    ]
    if item and item.get("field"):
        parts.append(f"字段:{item.get('field')}")
    if project.audit_year:
        parts.append(f"审计期间:{project.audit_year}")
    return "；".join(part for part in parts if not part.endswith(":"))


def row_from_plan(
    *,
    project: Project,
    item: dict[str, Any],
    rule: dict[str, Any],
    target: dict[str, Any],
    path_index: dict[str, dict[str, Any]],
    verification_index: dict[tuple[str, str, str, str], dict[str, Any]],
    manual_index: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    wp = workpaper_from_plan(item, path_index)
    rule_id = str(item.get("rule_id") or rule.get("rule_id") or "")
    field = clean_text(item.get("field") or target.get("target_field") or target.get("field"))
    locator = clean_text(item.get("locator") or target.get("locator_display") or target.get("locator"))
    cell = clean_text(item.get("cell") or target.get("target_cell") or target.get("cell"))
    workpaper_code = clean_text(wp.get("code") or rule.get("workpaper_code"))
    verification = find_verification(
        verification_index,
        rule_id=rule_id,
        field=field,
        locator=locator,
        cell=cell,
        workpaper_code=workpaper_code,
    )
    value_source = clean_text(item.get("value_source") or "rule")
    manual = manual_index.get(item.get("manual_correction_id") or -1) or latest_non_rejected_manual(rule, target)
    conflict = clean_text(item.get("message")) if value_source == "manual_correction_conflict" else ""
    manual_text = clean_text((manual or {}).get("manual_correction"))
    return {
        "project_id": project.id,
        "project_name": project.name,
        "category": workpaper_category(workpaper_code, rule.get("scope", "")),
        "workpaper_code": workpaper_code,
        "workpaper_name": clean_text(wp.get("name") or rule.get("workpaper_name") or rule.get("workpaper")),
        "rule_id": rule_id,
        "rule_name": clean_text(rule.get("name")),
        "field": field,
        "target_field": clean_text(target.get("target_field") or target.get("field") or field),
        "sheet": clean_text(item.get("sheet_name") or rule.get("sheet") or rule.get("section")),
        "locator": locator,
        "cell": cell,
        "paragraph": locator if locator.startswith("paragraph:") else "",
        "table": locator if locator.startswith("table:") else "",
        "locator_method": clean_text(target.get("locator_method") or ("fixed_cell" if cell else "plan_locator")),
        "fill_source": source_label(value_source, clean_text(rule.get("source_type"))),
        "source_type": clean_text(rule.get("source_type")),
        "match_logic": clean_text(rule.get("matching_logic")),
        "applicable_scope": applicable_scope(project, rule, item),
        "current_value": clean_text(item.get("old_value"), 500),
        "suggested_value": clean_text(item.get("new_value"), 500),
        "approved_value": clean_text(item.get("new_value"), 500) if value_source == "manual_correction_approved" else "",
        "original_rule_value": clean_text(item.get("original_rule_value"), 500),
        "manual_correction": manual_text,
        "manual_correction_status": clean_text((manual or {}).get("status")),
        "value_source": value_source,
        "write_status": clean_text(item.get("status")),
        "verification_status": clean_text((verification or {}).get("verification_status") or "not_verified"),
        "verification_message": clean_text((verification or {}).get("verification_message")),
        "conflict_message": conflict or clean_text((manual or {}).get("conflict_message")),
        "warning": clean_text("；".join(rule.get("project_gaps") or []) or item.get("message")),
    }


def row_from_rule_target(
    *,
    project: Project,
    rule: dict[str, Any],
    target: dict[str, Any],
) -> dict[str, Any]:
    manual = latest_non_rejected_manual(rule, target)
    workpaper_code = clean_text(rule.get("workpaper_code"))
    status = clean_text(rule.get("project_status") or rule.get("definition_status") or "not_checked")
    conflict = clean_text(rule.get("manual_correction_conflict") or (manual or {}).get("conflict_message"))
    return {
        "project_id": project.id,
        "project_name": project.name,
        "category": workpaper_category(workpaper_code, rule.get("scope", "")),
        "workpaper_code": workpaper_code,
        "workpaper_name": clean_text(rule.get("workpaper_name") or rule.get("workpaper")),
        "rule_id": clean_text(rule.get("rule_id") or rule.get("id")),
        "rule_name": clean_text(rule.get("name")),
        "field": clean_text(target.get("target_field") or target.get("field") or rule.get("target_field")),
        "target_field": clean_text(target.get("target_field") or target.get("field") or rule.get("target_field")),
        "sheet": clean_text(rule.get("sheet") or rule.get("section")),
        "locator": clean_text(target.get("locator_display") or target.get("locator")),
        "cell": clean_text(target.get("target_cell") or target.get("cell")),
        "paragraph": "",
        "table": "",
        "locator_method": clean_text(target.get("locator_method")),
        "fill_source": source_label("rule", clean_text(rule.get("source_type"))),
        "source_type": clean_text(rule.get("source_type")),
        "match_logic": clean_text(rule.get("matching_logic")),
        "applicable_scope": applicable_scope(project, rule),
        "current_value": "",
        "suggested_value": clean_text(rule.get("content_template"), 500),
        "approved_value": clean_text((manual or {}).get("manual_correction"), 500) if (manual or {}).get("status") == "approved" else "",
        "original_rule_value": "",
        "manual_correction": clean_text((manual or {}).get("manual_correction")),
        "manual_correction_status": clean_text((manual or {}).get("status")),
        "value_source": "manual_correction_approved" if (manual or {}).get("status") == "approved" else "rule",
        "write_status": "not_planned",
        "verification_status": "not_verified",
        "verification_message": "",
        "conflict_message": conflict,
        "warning": clean_text("；".join(rule.get("project_gaps") or rule.get("definition_gaps") or []) or status),
    }


def row_key(row: dict[str, Any]) -> tuple[str, str, str, str, str]:
    return (
        clean_text(row.get("rule_id")),
        clean_text(row.get("field") or row.get("target_field")),
        clean_text(row.get("workpaper_code")),
        clean_text(row.get("cell")),
        clean_text(row.get("locator")),
    )


def evidence_rule_details(payload: dict[str, Any], rows: list[dict[str, Any]], report: dict[str, Any] | None) -> dict[str, Any]:
    by_rule = {str(rule.get("rule_id") or ""): rule for rule in payload.get("rules") or []}
    c211_rule = by_rule.get("C-C21-1-CURRENT-FINDINGS", {})
    a27_rule = by_rule.get("A-A27-SUMMARY-MEMO", {})
    b44_rule = by_rule.get("B-B22A-4-4-1-ITGC-UNDERSTANDING", {})
    c211_status = Counter(row.get("verification_status") for row in rows if row.get("workpaper_code") == "C21-1")
    a27_status = Counter(row.get("verification_status") for row in rows if row.get("workpaper_code") == "A27")
    b44_warning = "；".join(b44_rule.get("project_gaps") or [])
    missing_codes: set[str] = set()
    match = re.search(r"缺少控制编号行：([^；]+)", b44_warning)
    if match:
        missing_codes = {clean_text(item) for item in re.split(r"[、,，\s]+", match.group(1)) if item}
    return {
        "C21-1": {
            "status": "需人工确认/证据不足" if c211_status.get("not_verified") or c211_status.get("skipped") else "已进入验证范围",
            "source": "C22设计/执行有效性测试结果、C26问题、C21-1现有发现和附件证据",
            "fields": [
                "缺陷编号", "类别", "控制类型", "涉及应用", "问题描述", "补偿性控制", "风险影响", "报表项目", "审计影响", "认定列",
            ],
            "logic": "仅在证据能支持缺陷事实时形成具体缺陷；当前 AUTO-ITGC 行对缺陷事实、责任范围、严重程度和认定影响保留人工确认。",
            "verification_counts": dict(c211_status),
            "warnings": c211_rule.get("project_gaps") or ["若 C22/C26 未提供异常结论，不编造具体缺陷。"],
        },
        "A27": {
            "status": "需人工确认" if a27_status.get("not_verified") or a27_status.get("skipped") else "已进入验证范围",
            "source": "项目主数据、项目成员、C22/C26测试结论、C21-1缺陷汇总和支持性底稿清单",
            "logic": "目的段和IT团队来自项目主数据/成员；ITAC/ITGC结论依赖C26/C22与C21-1；缺陷摘要依赖缺陷严重程度和人工确认状态。",
            "verification_counts": dict(a27_status),
            "warnings": a27_rule.get("project_gaps") or ["结论段不应脱离 C21-1 缺陷严重程度自动定稿。"],
        },
        "B22A": {
            "status": "模板行/别名映射检查",
            "controls": [
                {
                    "control_code": code,
                    "judgment": "模板无行或需要别名映射判断" if code in missing_codes else "未在最新项目检查中识别为缺失；仍需按底稿模板行和别名复核",
                    "action": "不硬插风险行；先确认模板行、控制编号别名和 C22 映射。",
                }
                for code in TARGET_B22A_CONTROLS
            ],
            "warnings": b44_rule.get("project_gaps") or ["未发现指定控制编号的模板缺行提示。"],
        },
        "verification_report": {
            "path": (report or {}).get("_report_path", ""),
            "generated_at": (report or {}).get("generated_at", ""),
            "summary": (report or {}).get("summary") or {},
        },
    }


def build_autofill_rule_visualization(
    db: Session,
    project: Project,
    corrections: list[RuleManualCorrection],
) -> dict[str, Any]:
    attachments = db.execute(
        select(Attachment).where(Attachment.project_id == project.id).order_by(Attachment.index_no)
    ).scalars().all()
    workpapers = db.execute(
        select(Workpaper).where(Workpaper.project_id == project.id).order_by(Workpaper.code)
    ).scalars().all()
    attachment_rows = list_dict(attachments)
    workpaper_rows = list_dict(workpapers)
    payload = inspect_project_autofill_rules(workpaper_rows, attachment_rows, include_template_scan=False)
    payload["project_id"] = project.id
    payload["project_name"] = project.name
    apply_manual_corrections_to_payload(payload, corrections)
    report = latest_verification_report(project.id)
    plan = plan_items_from_report(report) or latest_plan_items(db, project.id)
    verification_index = build_verification_index(report)
    rules = rule_map(payload)
    path_index = workpaper_by_path(workpaper_rows)
    manual_index = manual_by_id(corrections)
    rows: list[dict[str, Any]] = []
    planned_target_keys: set[tuple[str, str, str, str, str]] = set()

    for item in plan:
        rule = rules.get(str(item.get("rule_id") or ""), {})
        target = best_target(rule, item)
        row = row_from_plan(
            project=project,
            item=item,
            rule=rule,
            target=target,
            path_index=path_index,
            verification_index=verification_index,
            manual_index=manual_index,
        )
        rows.append(row)
        planned_target_keys.add(row_key(row))

    for rule in payload.get("rules") or []:
        targets = [target for target in rule.get("targets") or [{}] if isinstance(target, dict)]
        for target in targets or [{}]:
            row = row_from_rule_target(project=project, rule=rule, target=target)
            if row_key(row) not in planned_target_keys:
                rows.append(row)

    action_index = action_index_for_project(db, project.id)
    for row in rows:
        action_key = action_key_for_row(row)
        action = action_index.get(action_key)
        row["action_key"] = action_key
        row["action"] = action or None
        row["action_status"] = (action or {}).get("action_status", "")
        row["action_note"] = (action or {}).get("action_note", "")
        row["action_updated_at"] = (action or {}).get("updated_at", "")
        row["action_updated_by_user_id"] = (action or {}).get("updated_by_user_id")

    rows.sort(key=lambda item: (item.get("category", ""), item.get("workpaper_code", ""), item.get("rule_id", ""), item.get("field", "")))
    status_counts = Counter(row.get("verification_status") or "not_verified" for row in rows)
    value_source_counts = Counter(row.get("value_source") or "rule" for row in rows)
    action_counts = Counter(row.get("action_status") for row in rows if row.get("action_status"))
    conflict_count = sum(1 for row in rows if row.get("conflict_message"))
    summary = {
        "row_count": len(rows),
        "verification_status": {status: status_counts.get(status, 0) for status in VERIFICATION_STATUSES},
        "value_source": dict(value_source_counts),
        "action_status": dict(action_counts),
        "approved_manual_correction": value_source_counts.get("manual_correction_approved", 0),
        "manual_correction_conflict": value_source_counts.get("manual_correction_conflict", 0),
        "conflict_count": conflict_count,
    }
    return {
        "project_id": project.id,
        "project_name": project.name,
        "generated_at": datetime.utcnow().isoformat(),
        "summary": summary,
        "rows": rows,
        "evidence_rule_details": evidence_rule_details(payload, rows, report),
    }


VISUALIZATION_EXPORT_HEADERS = [
    ("category", "类别"),
    ("workpaper_code", "底稿编号"),
    ("workpaper_name", "底稿名称"),
    ("rule_id", "规则编号"),
    ("rule_name", "规则名称"),
    ("field", "字段"),
    ("target_field", "目标字段"),
    ("sheet", "Sheet/章节"),
    ("locator_method", "定位方式"),
    ("locator", "定位内容"),
    ("cell", "单元格"),
    ("fill_source", "填充来源"),
    ("match_logic", "匹配逻辑"),
    ("applicable_scope", "适用范围"),
    ("current_value", "当前值"),
    ("suggested_value", "建议值"),
    ("approved_value", "已批准值"),
    ("original_rule_value", "原规则值"),
    ("value_source", "值来源"),
    ("verification_status", "验证状态"),
    ("action_status", "团队处理状态"),
    ("action_note", "团队处理备注"),
    ("action_updated_at", "处理更新时间"),
    ("conflict_message", "冲突提示"),
    ("warning", "警告"),
    ("manual_correction", "人工修正"),
]


def finish_sheet(ws: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, column_cells in enumerate(ws.columns, start=1):
        letter = get_column_letter(index)
        max_len = min(60, max((len(str(cell.value or "")) for cell in column_cells), default=0) + 2)
        ws.column_dimensions[letter].width = max(12, max_len)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_rule_visualization_export(payload: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "规则可视化"
    ws.append([label for _, label in VISUALIZATION_EXPORT_HEADERS])
    for row in payload.get("rows") or []:
        ws.append([clean_text(row.get(key)) for key, _ in VISUALIZATION_EXPORT_HEADERS])
    finish_sheet(ws)

    summary = wb.create_sheet("汇总")
    summary.append(["项目", payload.get("project_name", "")])
    summary.append(["项目ID", payload.get("project_id", "")])
    summary.append(["生成时间", payload.get("generated_at", "")])
    for key, value in (payload.get("summary") or {}).items():
        summary.append([key, json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value])
    finish_sheet(summary)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
