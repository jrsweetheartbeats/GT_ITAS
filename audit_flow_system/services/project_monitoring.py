from __future__ import annotations

from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from ..models import Project, ReviewFinding, Workpaper
from .workpaper_metadata import clean_cell_text, extract_workpaper_metadata, load_extracted_fields


RESOLVED_REVIEW_VALUES = {"1", "true", "yes", "y", "是", "已解决", "解决", "已完成", "完成", "closed", "resolved"}
PLACEHOLDER_PERSON_VALUES = {
    "编制人",
    "复核人",
    "项目负责经理",
    "项目合伙人",
    "项目负责人",
    "负责人",
    "经理",
}
OPEN_FINDING_STATUSES = {
    "open",
    "assigned",
    "changes_requested",
    "responded",
    "retained",
    "revised",
    "returned",
    "blocked",
    "待处理",
    "已分派",
    "保留",
    "已修订",
    "退回",
    "阻塞",
    "待复核",
}
HIGH_SEVERITIES = {"high", "critical", "高", "重大"}
EXECUTION_PROJECT_STATUSES = {"in_progress", "review", "active", "running", "项目实施", "复核整改", "进行中", "正在执行"}
REVIEW_SHEET_NAMES = ("非C22_编辑页签", "C22_编辑页签")
REVIEW_FILE_TOKENS = ("IT审计检查表", "复核记录", "复核表")
SKIP_REVIEW_DIRS = {".codex_work", ".codex_tmp", ".git", "outputs", "output", "备份", "backup"}


def _normalized(value: Any) -> str:
    return clean_cell_text(value).strip().lower()


def _meaningful_person(value: Any) -> bool:
    text = clean_cell_text(value)
    if not text:
        return False
    compact = text.replace("（", "").replace("）", "").replace("(", "").replace(")", "").replace("：", "")
    if compact in PLACEHOLDER_PERSON_VALUES:
        return False
    if len(compact) > 30 or any(token in compact for token in ("认为必要", "审计业务", "如有")):
        return False
    return True


@lru_cache(maxsize=1024)
def _metadata_from_file(path_text: str, mtime_ns: int, size: int) -> dict[str, str]:
    del mtime_ns, size
    return extract_workpaper_metadata(path_text)


def _workpaper_metadata(workpaper: Workpaper, path: Path) -> dict[str, Any]:
    stored = load_extracted_fields(workpaper)
    header_keys = {"preparer", "prepared_date", "reviewer", "reviewed_date"}
    if any(clean_cell_text(stored.get(key)) for key in header_keys):
        return stored
    if not path.exists() or not path.is_file():
        return stored
    try:
        stat = path.stat()
        return {**stored, **_metadata_from_file(str(path), stat.st_mtime_ns, stat.st_size)}
    except OSError:
        return stored


def workpaper_monitor_item(workpaper: Workpaper) -> dict[str, Any]:
    path = Path(workpaper.file_path).expanduser() if workpaper.file_path else Path()
    file_exists = bool(workpaper.file_path) and path.exists() and path.is_file()
    metadata = _workpaper_metadata(workpaper, path)
    preparer = clean_cell_text(metadata.get("preparer"))
    prepared_date = clean_cell_text(metadata.get("prepared_date"))
    reviewer = clean_cell_text(metadata.get("reviewer"))
    reviewed_date = clean_cell_text(metadata.get("reviewed_date"))
    preparation_complete = _meaningful_person(preparer) and bool(prepared_date)
    preparation_partial = not preparation_complete and (_meaningful_person(preparer) or bool(prepared_date))
    review_complete = _meaningful_person(reviewer) and bool(reviewed_date)
    review_partial = not review_complete and (_meaningful_person(reviewer) or bool(reviewed_date))
    modified_at = ""
    if file_exists:
        try:
            modified_at = datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")
        except OSError:
            modified_at = ""
    return {
        "file_exists": file_exists,
        "file_modified_at": modified_at,
        "preparer": preparer,
        "prepared_date": prepared_date,
        "reviewer": reviewer,
        "reviewed_date": reviewed_date,
        "preparation_complete": preparation_complete,
        "preparation_partial": preparation_partial,
        "review_complete": review_complete,
        "review_partial": review_partial,
    }


def _review_file_candidates(project_root: str) -> list[Path]:
    if not project_root:
        return []
    root = Path(project_root).expanduser()
    if not root.exists() or not root.is_dir():
        return []
    candidates: list[Path] = []
    root_depth = len(root.parts)
    for path in root.rglob("*.xlsx"):
        relative_parts = path.relative_to(root).parts
        if len(path.parts) - root_depth > 3:
            continue
        if any(part in SKIP_REVIEW_DIRS or part.startswith(".") for part in relative_parts[:-1]):
            continue
        if path.name.startswith("~$") or "模板" in path.name:
            continue
        if any(token in path.name for token in REVIEW_FILE_TOKENS):
            candidates.append(path)
    return candidates


def _review_file_rank(path: Path) -> tuple[int, int]:
    reply_rank = 2 if "已回复" in path.name else 1
    try:
        modified = path.stat().st_mtime_ns
    except OSError:
        modified = 0
    return reply_rank, modified


def _is_review_resolved(value: Any) -> bool:
    return _normalized(value) in RESOLVED_REVIEW_VALUES


@lru_cache(maxsize=128)
def _read_review_workbook(path_text: str, mtime_ns: int, size: int) -> dict[str, Any]:
    del mtime_ns, size
    path = Path(path_text)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        issue_count = 0
        replied_count = 0
        resolved_count = 0
        source_counts: dict[str, int] = {}
        for sheet_name in REVIEW_SHEET_NAMES:
            if sheet_name not in workbook.sheetnames:
                continue
            worksheet = workbook[sheet_name]
            sheet_issue_count = 0
            for row in worksheet.iter_rows(min_row=2, min_col=1, max_col=13, values_only=True):
                issue = clean_cell_text(row[7])
                detail = clean_cell_text(row[8])
                if not issue and not detail:
                    continue
                sheet_issue_count += 1
                issue_count += 1
                if clean_cell_text(row[12]):
                    replied_count += 1
                if _is_review_resolved(row[11]):
                    resolved_count += 1
            source_counts["C22" if sheet_name.startswith("C22") else "非C22"] = sheet_issue_count
        return {
            "issue_count": issue_count,
            "replied_count": replied_count,
            "resolved_count": resolved_count,
            "source_counts": source_counts,
        }
    finally:
        workbook.close()


def latest_review_workbook_summary(project_root: str) -> dict[str, Any]:
    candidates = _review_file_candidates(project_root)
    if not candidates:
        return {
            "source": "database",
            "source_file": "",
            "source_name": "",
            "source_modified_at": "",
            "issue_count": 0,
            "replied_count": 0,
            "resolved_count": 0,
            "pending_confirmation_count": 0,
            "unreplied_count": 0,
            "reply_rate": 0,
            "closure_rate": 0,
            "source_counts": {},
        }
    selected = max(candidates, key=_review_file_rank)
    try:
        stat = selected.stat()
        summary = _read_review_workbook(str(selected), stat.st_mtime_ns, stat.st_size)
        modified_at = datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds")
    except (OSError, ValueError):
        return latest_review_workbook_summary("")
    issue_count = int(summary["issue_count"])
    replied_count = int(summary["replied_count"])
    resolved_count = int(summary["resolved_count"])
    return {
        "source": "review_workbook",
        "source_file": str(selected),
        "source_name": selected.name,
        "source_modified_at": modified_at,
        **summary,
        "pending_confirmation_count": max(0, replied_count - resolved_count),
        "unreplied_count": max(0, issue_count - replied_count),
        "reply_rate": round(replied_count / issue_count * 100) if issue_count else 100,
        "closure_rate": round(resolved_count / issue_count * 100) if issue_count else 100,
    }


def _database_review_summary(findings: Iterable[ReviewFinding]) -> dict[str, Any]:
    rows = list(findings)
    open_rows = [row for row in rows if _normalized(row.status) in OPEN_FINDING_STATUSES]
    replied_count = sum(1 for row in rows if clean_cell_text(row.review_comment))
    resolved_count = len(rows) - len(open_rows)
    issue_count = len(rows)
    return {
        "source": "database",
        "source_file": "",
        "source_name": "",
        "source_modified_at": "",
        "issue_count": issue_count,
        "replied_count": replied_count,
        "resolved_count": resolved_count,
        "pending_confirmation_count": max(0, replied_count - resolved_count),
        "unreplied_count": max(0, issue_count - replied_count),
        "reply_rate": round(replied_count / issue_count * 100) if issue_count else 100,
        "closure_rate": round(resolved_count / issue_count * 100) if issue_count else 100,
        "source_counts": {},
    }


def _latest_activity(workpaper_rows: list[dict[str, Any]], review: dict[str, Any]) -> str:
    candidates = [row["file_modified_at"] for row in workpaper_rows if row["file_modified_at"]]
    if review.get("source_modified_at"):
        candidates.append(str(review["source_modified_at"]))
    return max(candidates, default="")


def build_project_monitoring(
    project: Project,
    workpapers: Iterable[Workpaper],
    findings: Iterable[ReviewFinding],
) -> dict[str, Any]:
    workpaper_models = list(workpapers)
    finding_rows = list(findings)
    workpaper_rows = [workpaper_monitor_item(row) for row in workpaper_models]
    total = len(workpaper_rows)
    available = sum(1 for row in workpaper_rows if row["file_exists"])
    prepared = sum(1 for row in workpaper_rows if row["preparation_complete"])
    preparation_partial = sum(1 for row in workpaper_rows if row["preparation_partial"])
    reviewed = sum(1 for row in workpaper_rows if row["review_complete"])
    review_partial = sum(1 for row in workpaper_rows if row["review_partial"])
    workpaper_summary = {
        "total": total,
        "available_count": available,
        "missing_count": max(0, total - available),
        "prepared_count": prepared,
        "preparation_partial_count": preparation_partial,
        "unprepared_count": max(0, available - prepared - preparation_partial),
        "reviewed_count": reviewed,
        "review_partial_count": review_partial,
        "unreviewed_count": max(0, available - reviewed - review_partial),
        "availability_rate": round(available / total * 100) if total else 0,
        "preparation_rate": round(prepared / total * 100) if total else 0,
        "review_rate": round(reviewed / total * 100) if total else 0,
    }
    review = latest_review_workbook_summary(project.project_root)
    if review["source"] != "review_workbook":
        review = _database_review_summary(finding_rows)

    availability_rate = workpaper_summary["availability_rate"]
    preparation_rate = workpaper_summary["preparation_rate"]
    review_rate = workpaper_summary["review_rate"]
    closure_rate = review["closure_rate"]
    readiness_rate = round(
        availability_rate * 0.25
        + preparation_rate * 0.30
        + review_rate * 0.25
        + closure_rate * 0.20
    )
    open_high_count = sum(
        1
        for row in finding_rows
        if _normalized(row.status) in OPEN_FINDING_STATUSES and _normalized(row.severity) in HIGH_SEVERITIES
    )
    blockers: list[dict[str, str]] = []
    if total == 0:
        execution_started = _normalized(project.status) in EXECUTION_PROJECT_STATUSES
        blockers.append(
            {
                "level": "high" if execution_started else "medium",
                "code": "workpaper_not_registered",
                "message": "尚未登记项目底稿" if execution_started else "计划阶段尚未建立底稿台账",
            }
        )
    if workpaper_summary["missing_count"]:
        blockers.append(
            {
                "level": "high",
                "code": "workpaper_file_missing",
                "message": f"{workpaper_summary['missing_count']} 份底稿登记路径失联",
            }
        )
    signature_gap = max(0, available - prepared)
    if signature_gap:
        blockers.append(
            {
                "level": "medium",
                "code": "preparation_signature_incomplete",
                "message": f"{signature_gap} 份可用底稿尚未形成完整编制签名",
            }
        )
    review_gap = max(0, available - reviewed)
    if review_gap:
        blockers.append(
            {
                "level": "medium",
                "code": "review_signature_incomplete",
                "message": f"{review_gap} 份可用底稿尚未形成完整复核签名",
            }
        )
    if review["unreplied_count"]:
        blockers.append(
            {
                "level": "high",
                "code": "review_issue_unreplied",
                "message": f"{review['unreplied_count']} 条复核问题尚未回复",
            }
        )
    if review["pending_confirmation_count"]:
        blockers.append(
            {
                "level": "medium",
                "code": "review_pending_confirmation",
                "message": f"{review['pending_confirmation_count']} 条整改回复待复核确认",
            }
        )
    if open_high_count:
        blockers.append(
            {
                "level": "high",
                "code": "open_high_risk_finding",
                "message": f"{open_high_count} 条系统高风险问题尚未关闭",
            }
        )
    status = "green"
    if any(row["level"] == "high" for row in blockers):
        status = "red"
    elif blockers:
        status = "amber"
    blockers.sort(key=lambda row: 0 if row["level"] == "high" else 1)
    next_action = blockers[0]["message"] if blockers else "底稿、复核与整改均已形成闭环"
    return {
        "status": status,
        "readiness_rate": readiness_rate,
        "workpapers": workpaper_summary,
        "review": review,
        "open_high_risk_count": open_high_count,
        "blocker_count": len(blockers),
        "blockers": blockers,
        "next_action": next_action,
        "last_activity_at": _latest_activity(workpaper_rows, review),
        "readiness_formula": "文件可用25% + 编制签名30% + 复核签名25% + 问题关闭20%",
    }
