from __future__ import annotations

from datetime import date, timedelta
from functools import lru_cache
import json
from pathlib import Path
import re
from typing import Any

from docx import Document
from openpyxl import load_workbook

from ..models import Project, Workpaper
from .workpaper_metadata import clean_cell_text


DELIVERY_OFFSET_DAYS = 20
PLAN_MEMO_CODE_PREFIX = "B60-2-3"
PLAN_MEMO_NAME_KEYWORDS = ("IT审计计划备忘录", "计划备忘录")
EXIT_LABEL_KEYWORDS = (
    "进离场",
    "项目离场时间",
    "离场时间",
    "计划离场",
    "现场离场",
    "项目离场",
    "退场时间",
    "退场",
    "现场工作结束",
    "审计现场时间安排",
    "现场时间安排",
    "项目计划",
)
RANGE_END_LABELS = {"进离场", "审计现场时间安排", "现场时间安排", "项目计划"}

FULL_DATE_PATTERN = re.compile(
    r"(?P<year>20\d{2})\s*[年./\-]\s*(?P<month>\d{1,2})\s*[月./\-]\s*(?P<day>\d{1,2})\s*日?"
)
COMPACT_DATE_PATTERN = re.compile(r"(?<!\d)(?P<year>20\d{2})(?P<month>\d{2})(?P<day>\d{2})(?!\d)")
RANGE_END_PATTERN = re.compile(r"(?:至|到|~|～|-|—|－)\s*(?P<month>\d{1,2})\s*[月./\-]\s*(?P<day>\d{1,2})\s*日?")
NEXT_LABEL_KEYWORDS = ("项目交付", "交付时间", "预计交付", "报告日期", "归档", "项目进场", "进场时间")


def is_plan_memo_workpaper(workpaper: Workpaper) -> bool:
    code = str(workpaper.code or "").upper().replace(" ", "")
    name = str(workpaper.name or "")
    return code.startswith(PLAN_MEMO_CODE_PREFIX) or any(keyword in name for keyword in PLAN_MEMO_NAME_KEYWORDS)


def _safe_date(year: int, month: int, day: int) -> date | None:
    try:
        return date(year, month, day)
    except ValueError:
        return None


def _date_candidates(text: Any) -> list[tuple[int, date]]:
    source = clean_cell_text(text)
    if not source:
        return []
    candidates: list[tuple[int, date]] = []
    occupied: list[range] = []
    default_year: int | None = None
    for pattern in (FULL_DATE_PATTERN, COMPACT_DATE_PATTERN):
        for match in pattern.finditer(source):
            parsed = _safe_date(int(match.group("year")), int(match.group("month")), int(match.group("day")))
            if parsed:
                candidates.append((match.start(), parsed))
                occupied.append(range(match.start(), match.end()))
                default_year = parsed.year
    if default_year:
        for match in RANGE_END_PATTERN.finditer(source):
            if any(match.start() in span or match.end() - 1 in span for span in occupied):
                continue
            parsed = _safe_date(default_year, int(match.group("month")), int(match.group("day")))
            if parsed:
                candidates.append((match.start(), parsed))
    return sorted(candidates, key=lambda item: item[0])


def _last_date(text: Any) -> date | None:
    candidates = _date_candidates(text)
    return candidates[-1][1] if candidates else None


def _exit_date_from_fields(workpaper: Workpaper) -> date | None:
    try:
        fields = json.loads(workpaper.extracted_fields_json or "{}")
    except json.JSONDecodeError:
        fields = {}
    if not isinstance(fields, dict):
        return None
    for key, value in fields.items():
        key_text = str(key or "").lower()
        if any(keyword in key_text for keyword in ("离场", "退场", "exit", "departure", "leave")):
            parsed = _last_date(value)
            if parsed:
                return parsed
    for value in fields.values():
        value_text = clean_cell_text(value)
        if any(keyword in value_text for keyword in EXIT_LABEL_KEYWORDS):
            parsed = _last_date(value_text)
            if parsed:
                return parsed
    return None


def _line_priority(line: str) -> int:
    compact = re.sub(r"\s+", "", line)
    if "项目离场时间" in compact or "离场时间" in compact:
        return 30
    if "项目离场" in compact or "计划离场" in compact or "现场离场" in compact:
        return 25
    if "进离场" in compact:
        return 20
    if "退场" in compact or "现场工作结束" in compact:
        return 15
    if "审计现场时间安排" in compact or "现场时间安排" in compact or "项目计划" in compact:
        return 10
    return 0


def _exit_date_from_lines(lines: list[str]) -> date | None:
    matches: list[tuple[int, int, date]] = []
    for index, line in enumerate(lines):
        priority = _line_priority(line)
        if not priority:
            continue
        window = " ".join(item for item in lines[index : index + 2] if item)
        parsed = _date_from_exit_window(window) or _last_date(window)
        if parsed:
            matches.append((priority, index, parsed))
    if not matches:
        return None
    matches.sort(key=lambda item: (item[0], item[2], -item[1]), reverse=True)
    return matches[0][2]


def _date_from_exit_window(text: str) -> date | None:
    for keyword in EXIT_LABEL_KEYWORDS:
        start = text.find(keyword)
        if start < 0:
            continue
        segment = text[start:]
        for next_label in NEXT_LABEL_KEYWORDS:
            if next_label == keyword:
                continue
            stop = segment.find(next_label, len(keyword))
            if stop > 0:
                segment = segment[:stop]
        candidates = _date_candidates(segment)
        if not candidates:
            continue
        return candidates[-1][1] if keyword in RANGE_END_LABELS else candidates[0][1]
    return None


def _excel_lines(path: Path, max_rows: int = 220, max_sheets: int = 6) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True, keep_vba=path.suffix.lower() == ".xlsm")
    try:
        lines: list[str] = []
        for ws in wb.worksheets[:max_sheets]:
            for row in ws.iter_rows(max_row=max_rows):
                values = [clean_cell_text(cell.value) for cell in row if clean_cell_text(cell.value)]
                if values:
                    lines.append(" | ".join(values))
        return lines
    finally:
        wb.close()


def _docx_lines(path: Path, max_lines: int = 360) -> list[str]:
    doc = Document(path)
    lines: list[str] = [clean_cell_text(paragraph.text) for paragraph in doc.paragraphs if clean_cell_text(paragraph.text)]
    for table in doc.tables:
        for row in table.rows:
            values = [clean_cell_text(cell.text) for cell in row.cells if clean_cell_text(cell.text)]
            if values:
                lines.append(" | ".join(values))
            if len(lines) >= max_lines:
                return lines[:max_lines]
    return lines[:max_lines]


@lru_cache(maxsize=256)
def _exit_date_from_file(path_text: str, mtime_ns: int, size: int) -> date | None:
    path = Path(path_text)
    suffix = path.suffix.lower()
    try:
        if suffix in {".xlsx", ".xlsm"}:
            return _exit_date_from_lines(_excel_lines(path))
        if suffix == ".docx":
            return _exit_date_from_lines(_docx_lines(path))
    except Exception:
        return None
    return None


def _exit_date_from_workpaper(workpaper: Workpaper) -> date | None:
    parsed = _exit_date_from_fields(workpaper)
    if parsed:
        return parsed
    if not workpaper.file_path:
        return None
    path = Path(workpaper.file_path).expanduser()
    if not path.exists() or not path.is_file():
        return None
    stat = path.stat()
    return _exit_date_from_file(str(path), stat.st_mtime_ns, stat.st_size)


def delivery_info(project: Project, workpapers: list[Workpaper] | None = None) -> dict[str, Any]:
    candidates: list[tuple[date, Workpaper]] = []
    for workpaper in workpapers or []:
        if not is_plan_memo_workpaper(workpaper):
            continue
        parsed = _exit_date_from_workpaper(workpaper)
        if parsed:
            candidates.append((parsed, workpaper))
    if candidates:
        exit_date, workpaper = max(candidates, key=lambda item: item[0])
        delivery_date = exit_date + timedelta(days=DELIVERY_OFFSET_DAYS)
        return {
            "project_exit_date": exit_date,
            "delivery_date": delivery_date,
            "due_days": (delivery_date - date.today()).days,
            "delivery_source": f"{workpaper.code or PLAN_MEMO_CODE_PREFIX} IT审计计划备忘录：项目离场时间+{DELIVERY_OFFSET_DAYS}天",
            "delivery_source_workpaper_id": workpaper.id,
            "delivery_source_workpaper_code": workpaper.code,
            "delivery_source_workpaper_name": workpaper.name,
            "delivery_source_type": "plan_memo_exit_date",
        }

    fallback_exit = project.end_date
    if fallback_exit:
        delivery_date = fallback_exit + timedelta(days=DELIVERY_OFFSET_DAYS)
        return {
            "project_exit_date": fallback_exit,
            "delivery_date": delivery_date,
            "due_days": (delivery_date - date.today()).days,
            "delivery_source": f"项目结束日期+{DELIVERY_OFFSET_DAYS}天（未识别到计划备忘录离场时间）",
            "delivery_source_type": "project_end_date_fallback",
        }

    fallback_delivery = project.audit_scope_end
    if fallback_delivery:
        return {
            "project_exit_date": None,
            "delivery_date": fallback_delivery,
            "due_days": (fallback_delivery - date.today()).days,
            "delivery_source": "审计范围结束日兜底（未识别到计划备忘录离场时间）",
            "delivery_source_type": "audit_scope_end_fallback",
        }

    return {
        "project_exit_date": None,
        "delivery_date": None,
        "due_days": None,
        "delivery_source": "未识别到计划备忘录离场时间",
        "delivery_source_type": "missing",
    }
