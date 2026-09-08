# ====== [模块 1] 解析：B60-2-3 IT审计计划备忘录.docx ======
from docx import Document
import re
from datetime import datetime

DATE_CN_RE = re.compile(r"(\d{4})年(\d{1,2})月(\d{1,2})日")
DATE_SLASH_RANGE_RE = re.compile(
    r"期间[:：]\s*(\d{4})/(\d{1,2})/(\d{1,2})\s*[-~—至]\s*(\d{4})/(\d{1,2})/(\d{1,2})"
)

def _cn_date_to_iso(s: str) -> str:
    m = DATE_CN_RE.search(s)
    if not m:
        return ""
    y, mo, d = map(int, m.groups())
    return f"{y:04d}-{mo:02d}-{d:02d}"

def _extract_audit_period_from_title(title_line: str) -> str:
    m = DATE_SLASH_RANGE_RE.search(title_line)
    if not m:
        return ""
    y1, m1, d1, y2, m2, d2 = map(int, m.groups())
    return f"{y1:04d}-{m1:02d}-{d1:02d}~{y2:04d}-{m2:02d}-{d2:02d}"

def parse_it_audit_plan_memo(docx_path: str) -> dict:
    """
    专用于：IT审计计划备忘录.docx
    输出尽量结构化，供后续“勾稽一致性 + 覆盖性”检查使用。
    """
    doc = Document(docx_path)

    # --- 1) 段落整体文本 ---
    paras = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    full_text = "\n".join(paras)

    result = {
        "template": "IT审计计划备忘录(docx)",
        "file": docx_path,
        # 勾稽核心字段（该底稿是“标准源”）
        "entity_name": "",       # 九芝堂股份有限公司
        "cutoff_date": "",       # 2025-12-31
        "audit_period": "",      # 2025-01-01~2025-12-31
        # 计划会议/关键日期
        "memo_date": "",
        "schedule": {            # 项目关键时间
            "project_plan_start": "", "project_plan_end": "",
            "fieldwork_start": "", "fieldwork_end": "",
            "sign_date": "", "report_date": "", "archive_date": ""
        },
        # 参与人员
        "to": "",
        "from": "",
        "subject": "",
        "project_team_people_text": [],   # 文本段落里列出的项目组人员
        "it_specialists": [],             # 表格里的 IT 专业人员名单（强结构化）
        # 业务流程与系统、控制表、系统环境表
        "process_system_map": [],
        "ipc_controls": [],
        "system_env": []
    }

    # --- 2) 从标题行提取 audit_period ---
    if paras:
        result["audit_period"] = _extract_audit_period_from_title(paras[0])

    # --- 3) 从“目的段”提取企业名称 & 截止日（句子通常固定） ---
    # 例：本备忘录...九芝堂股份有限公司...截止于2025年12月31日
    m_company = re.search(r"关于(.+?)(?:（|\"|“)", full_text)
    if m_company:
        result["entity_name"] = m_company.group(1).strip()
    else:
        # 兜底：直接找 “股份有限公司/有限公司”
        m2 = re.search(r"([\u4e00-\u9fff]{2,40}(?:股份有限公司|有限公司|有限责任公司))", full_text)
        if m2:
            result["entity_name"] = m2.group(1)

    m_cutoff = re.search(r"截止于(\d{4}年\d{1,2}月\d{1,2}日)", full_text)
    if m_cutoff:
        result["cutoff_date"] = _cn_date_to_iso(m_cutoff.group(1))

    # --- 4) 解析表格 ---
    tables = doc.tables

    # Table 0：日期/致/发自/主题
    if len(tables) >= 1:
        t0 = tables[0]
        # 预期 4x2
        for r in range(len(t0.rows)):
            k = t0.cell(r, 0).text.strip().replace("：", "").replace(":", "")
            v = t0.cell(r, 1).text.strip()
            if k in ("日期",):
                result["memo_date"] = _cn_date_to_iso(v) or v
            elif k in ("致",):
                result["to"] = v
            elif k in ("发自",):
                result["from"] = v
            elif k in ("主题",):
                result["subject"] = v

    # Table 1：IT专业人员（姓名/职级）
    if len(tables) >= 2:
        t1 = tables[1]
        # 第一行是表头
        for r in range(1, len(t1.rows)):
            name = t1.cell(r, 0).text.strip()
            title = t1.cell(r, 1).text.strip()
            if name:
                result["it_specialists"].append({"name": name, "title": title})

    # Table 2：重大流程 vs 系统
    if len(tables) >= 3:
        t2 = tables[2]
        for r in range(1, len(t2.rows)):
            process = t2.cell(r, 0).text.strip()
            systems = t2.cell(r, 1).text.strip()
            if process:
                sys_list = [s.strip() for s in re.split(r"[、,，]\s*", systems) if s.strip()]
                result["process_system_map"].append({"process": process, "systems": sys_list})

    # Table 3：信息处理控制及数据（A~E）
    if len(tables) >= 4:
        t3 = tables[3]
        for r in range(1, len(t3.rows)):
            row = {
                "account": t3.cell(r, 0).text.strip(),
                "process": t3.cell(r, 1).text.strip(),
                "risk": t3.cell(r, 2).text.strip(),
                "control_desc": t3.cell(r, 3).text.strip(),
                "systems_raw": t3.cell(r, 4).text.strip(),
            }
            row["systems"] = [s.strip() for s in re.split(r"[、,，]\s*", row["systems_raw"]) if s.strip()]
            if any(row.values()):
                result["ipc_controls"].append(row)

    # Table 4：系统环境（应用/DB/OS/DC/网络）
    if len(tables) >= 5:
        t4 = tables[4]
        for r in range(1, len(t4.rows)):
            env = {
                "app": t4.cell(r, 0).text.strip(),
                "db": t4.cell(r, 1).text.strip(),
                "os": t4.cell(r, 2).text.strip(),
                "dc": t4.cell(r, 3).text.strip(),
                "network": t4.cell(r, 4).text.strip(),
            }
            if env["app"]:
                result["system_env"].append(env)

    # --- 5) 解析“时间安排”段落里的关键日期 ---
    def _find_bracket_date(line: str) -> str:
        # 形如：[2026年01月04日]
        m = re.search(r"\[(\d{4}年\d{1,2}月\d{1,2}日)\]", line)
        return _cn_date_to_iso(m.group(1)) if m else ""

    for line in paras:
        if line.startswith("项目计划："):
            dates = re.findall(r"\[(\d{4}年\d{1,2}月\d{1,2}日)\]", line)
            if len(dates) == 2:
                result["schedule"]["project_plan_start"] = _cn_date_to_iso(dates[0])
                result["schedule"]["project_plan_end"] = _cn_date_to_iso(dates[1])
        elif line.startswith("审计执行："):
            dates = re.findall(r"\[(\d{4}年\d{1,2}月\d{1,2}日)\]", line)
            if len(dates) == 2:
                result["schedule"]["fieldwork_start"] = _cn_date_to_iso(dates[0])
                result["schedule"]["fieldwork_end"] = _cn_date_to_iso(dates[1])
        elif line.startswith("签字日期："):
            result["schedule"]["sign_date"] = _cn_date_to_iso(line) or result["schedule"]["sign_date"]
        elif line.startswith("报告日期："):
            result["schedule"]["report_date"] = _cn_date_to_iso(line) or result["schedule"]["report_date"]
        elif line.startswith("预计归档日期："):
            result["schedule"]["archive_date"] = _cn_date_to_iso(line) or result["schedule"]["archive_date"]

    # --- 6) 提取“参与计划会议人员（文本段）” ---
    # 这块名单在段落里：一般是“姓名，职级”
    people = []
    for line in paras:
        if re.search(r"[，,]\s*(合伙人|总监|高级经理|经理|项目经理)", line):
            people.append(line)
    result["project_team_people_text"] = people

    return result


# ====== [总入口处的模板路由示例] ======
def parse_one_file_by_template(path: str) -> dict:
    """
    根据文件名/扩展名选择专用解析器。
    目前只加了 1 份：IT审计计划备忘录.docx
    """
    name = path.split("/")[-1]
    if name.endswith(".docx") and ("IT审计计划备忘录" in name or "计划备忘录" in name):
        return parse_it_audit_plan_memo(path)

    # 其他底稿：下一轮我们再逐个加
    return {"template": "unknown", "file": path}
